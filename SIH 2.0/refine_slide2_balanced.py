import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

pptx_path = "SIH 2.0/SIH2025-IDEA-Presentation-Format.pptx"
prs = Presentation(pptx_path)

# Target Slide 2 ONLY
slide2 = prs.slides[1]

# Palette - Editorial Minimalist & High Contrast
COLOR_NAVY_DARK = RGBColor(0x0B, 0x13, 0x2B)
COLOR_NAVY_TITLE = RGBColor(0x1F, 0x4E, 0x79)
COLOR_SLATE_DARK = RGBColor(0x1E, 0x29, 0x3B)
COLOR_SLATE_MUTED = RGBColor(0x47, 0x55, 0x69)
COLOR_SLATE_LIGHT = RGBColor(0x64, 0x74, 0x8B)

COLOR_TEAL_DARK = RGBColor(0x0D, 0x94, 0x88)
COLOR_TEAL_ACCENT = RGBColor(0x14, 0xB8, 0xA6)
COLOR_TEAL_TINT = RGBColor(0xF0, 0xFD, 0xFA)
COLOR_TEAL_BORDER = RGBColor(0x0D, 0x94, 0x88)

COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_OFFWHITE = RGBColor(0xF8, 0xF9, 0xFA)
COLOR_CARD_BORDER = RGBColor(0xCB, 0xD5, 0xE1)
COLOR_SUBTLE_BORDER = RGBColor(0xE2, 0xE8, 0xF0)

# 1. Clean slide2 back to only the 6 official base template shapes
base_names = {'Rectangle 8', 'Title 1', 'Slide Number Placeholder 5', 'Footer Placeholder 6', 'Oval 9', 'Picture 2'}
to_remove = [s for s in slide2.shapes if s.name not in base_names]
for s in to_remove:
    sp = s._element
    sp.getparent().remove(sp)

# 2. Update Oval 9 (Team Badge)
for s in slide2.shapes:
    if s.name == "Oval 9":
        s.left = Inches(0.80)
        s.top = Inches(0.30)
        s.width = Inches(1.70)
        s.height = Inches(0.55)
        s.fill.solid()
        s.fill.fore_color.rgb = COLOR_OFFWHITE
        s.line.color.rgb = COLOR_CARD_BORDER
        s.line.width = Pt(1)
        if s.has_text_frame:
            tf = s.text_frame
            tf.clear()
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            p = tf.paragraphs[0]
            p.text = "TeamChargers"
            p.alignment = PP_ALIGN.CENTER
            p.font.name = "Arial"
            p.font.size = Pt(11)
            p.font.bold = True
            p.font.color.rgb = COLOR_NAVY_DARK

# 3. Update Title 1 (Respecting buffer before top-right SIH logo)
for s in slide2.shapes:
    if s.name == "Title 1":
        s.left = Inches(2.65)
        s.top = Inches(0.16)
        s.width = Inches(7.80)
        s.height = Inches(0.95)
        if s.has_text_frame:
            tf = s.text_frame
            tf.word_wrap = True
            tf.clear()
            # Mandatory Template Heading: IDEA TITLE
            p1 = tf.paragraphs[0]
            p1.alignment = PP_ALIGN.CENTER
            r1 = p1.add_run()
            r1.text = "IDEA TITLE: "
            r1.font.name = "Garamond"
            r1.font.size = Pt(22)
            r1.font.bold = True
            r1.font.color.rgb = COLOR_NAVY_TITLE
            
            r2 = p1.add_run()
            r2.text = "SkillSaathi"
            r2.font.name = "Arial"
            r2.font.size = Pt(22)
            r2.font.bold = True
            r2.font.color.rgb = COLOR_NAVY_DARK
            p1.space_after = Pt(2)
            
            # Tagline
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.CENTER
            r3 = p2.add_run()
            r3.text = "Beyond Career Advice. Towards Family Consensus."
            r3.font.name = "Arial"
            r3.font.size = Pt(14)
            r3.font.italic = True
            r3.font.color.rgb = COLOR_TEAL_DARK

# ==============================================================================
# ONE-LINE DESCRIPTION (Sub-Header)
# ==============================================================================
tb_desc = slide2.shapes.add_textbox(Inches(0.80), Inches(1.22), Inches(11.733), Inches(0.40))
tb_desc.name = "Refined_OneLineDesc"
tf_desc = tb_desc.text_frame
tf_desc.word_wrap = True
tf_desc.margin_left = tf_desc.margin_right = tf_desc.margin_top = tf_desc.margin_bottom = 0
p_desc = tf_desc.paragraphs[0]
p_desc.alignment = PP_ALIGN.CENTER
r_desc = p_desc.add_run()
r_desc.text = "An AI-powered family counselling platform that bridges student aspirations and parental concerns through evidence-backed vocational guidance."
r_desc.font.name = "Arial"
r_desc.font.size = Pt(13)
r_desc.font.color.rgb = COLOR_SLATE_DARK

# ==============================================================================
# MAIN VISUAL: FAMILY DECISION INTELLIGENCE WORKFLOW (60% Visual Storytelling)
# Convergent Architecture:
# [Student Aspirations] ──↘
#                           ──→ [AI-Powered Family Counselling] ──→ [Evidence-Based Comparison] ──→ [Family Decision]
# [Parent Concerns]     ──↗
# Top: 1.78", Height: 2.85"
# ==============================================================================
wf_top = Inches(1.78)
wf_h = Inches(2.85)

# Outer workflow canvas panel
wf_canvas = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.80), wf_top, Inches(11.733), wf_h)
wf_canvas.name = "Refined_WF_Canvas"
wf_canvas.fill.solid()
wf_canvas.fill.fore_color.rgb = COLOR_WHITE
wf_canvas.line.color.rgb = COLOR_SUBTLE_BORDER
wf_canvas.line.width = Pt(1)

# Subtle top tag for the workflow
tb_wftag = slide2.shapes.add_textbox(Inches(1.00), wf_top + Inches(0.08), Inches(11.333), Inches(0.28))
tb_wftag.name = "Refined_WF_Tag"
tf_wft = tb_wftag.text_frame
tf_wft.margin_left = tf_wft.margin_right = tf_wft.margin_top = tf_wft.margin_bottom = 0
pwft = tf_wft.paragraphs[0]
rw1 = pwft.add_run()
rw1.text = "FAMILY DECISION INTELLIGENCE™ WORKFLOW"
rw1.font.name = "Arial"
rw1.font.size = Pt(8.5)
rw1.font.bold = True
rw1.font.color.rgb = COLOR_TEAL_DARK

# --- STAGE 1: DUAL INPUTS (Converging) ---
# Student Node
c_student = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.00), wf_top + Inches(0.42), Inches(2.40), Inches(1.05))
c_student.name = "Refined_Node_Student"
c_student.fill.solid()
c_student.fill.fore_color.rgb = COLOR_OFFWHITE
c_student.line.color.rgb = COLOR_CARD_BORDER
c_student.line.width = Pt(1)
tfs = c_student.text_frame
tfs.word_wrap = True
tfs.margin_left = tfs.margin_right = Inches(0.12)
tfs.margin_top = Inches(0.12)
tfs.margin_bottom = Inches(0.08)
ps1 = tfs.paragraphs[0]
ps1.text = "Student Aspirations"
ps1.font.name = "Arial"
ps1.font.size = Pt(13)
ps1.font.bold = True
ps1.font.color.rgb = COLOR_NAVY_DARK
ps1.space_after = Pt(2)
ps2 = tfs.add_paragraph()
ps2.text = "Interests, Aptitude & Career Passion"
ps2.font.name = "Arial"
ps2.font.size = Pt(9.5)
ps2.font.color.rgb = COLOR_SLATE_MUTED

# Parent Node
c_parent = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.00), wf_top + Inches(1.62), Inches(2.40), Inches(1.05))
c_parent.name = "Refined_Node_Parent"
c_parent.fill.solid()
c_parent.fill.fore_color.rgb = COLOR_OFFWHITE
c_parent.line.color.rgb = COLOR_CARD_BORDER
c_parent.line.width = Pt(1)
tfp = c_parent.text_frame
tfp.word_wrap = True
tfp.margin_left = tfp.margin_right = Inches(0.12)
tfp.margin_top = Inches(0.12)
tfp.margin_bottom = Inches(0.08)
pp1 = tfp.paragraphs[0]
pp1.text = "Parent Concerns"
pp1.font.name = "Arial"
pp1.font.size = Pt(13)
pp1.font.bold = True
pp1.font.color.rgb = COLOR_NAVY_DARK
pp1.space_after = Pt(2)
pp2 = tfp.add_paragraph()
pp2.text = "Financial Security, ROI & Social Stigma"
pp2.font.name = "Arial"
pp2.font.size = Pt(9.5)
pp2.font.color.rgb = COLOR_SLATE_MUTED

# Converging arrows / indicators from Student and Parent to Center
tb_arr_s = slide2.shapes.add_textbox(Inches(3.45), wf_top + Inches(0.72), Inches(0.40), Inches(0.40))
tb_arr_s.name = "Refined_Arr_Student"
tfa_s = tb_arr_s.text_frame
tfa_s.margin_left = tfa_s.margin_right = tfa_s.margin_top = tfa_s.margin_bottom = 0
pa_s = tfa_s.paragraphs[0]
pa_s.text = "↘"
pa_s.font.name = "Arial"
pa_s.font.size = Pt(18)
pa_s.font.bold = True
pa_s.font.color.rgb = COLOR_TEAL_ACCENT

tb_arr_p = slide2.shapes.add_textbox(Inches(3.45), wf_top + Inches(1.85), Inches(0.40), Inches(0.40))
tb_arr_p.name = "Refined_Arr_Parent"
tfa_p = tb_arr_p.text_frame
tfa_p.margin_left = tfa_p.margin_right = tfa_p.margin_top = tfa_p.margin_bottom = 0
pa_p = tfa_p.paragraphs[0]
pa_p.text = "↗"
pa_p.font.name = "Arial"
pa_p.font.size = Pt(18)
pa_p.font.bold = True
pa_p.font.color.rgb = COLOR_TEAL_ACCENT

# --- STAGE 2: AI-POWERED FAMILY COUNSELLING (Center Convergence Hub) ---
c_counsel = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(3.90), wf_top + Inches(0.42), Inches(2.65), Inches(2.25))
c_counsel.name = "Refined_Node_Counsel"
c_counsel.fill.solid()
c_counsel.fill.fore_color.rgb = COLOR_TEAL_TINT
c_counsel.line.color.rgb = COLOR_TEAL_BORDER
c_counsel.line.width = Pt(1.5)
tfc = c_counsel.text_frame
tfc.word_wrap = True
tfc.margin_left = tfc.margin_right = Inches(0.14)
tfc.margin_top = Inches(0.20)
tfc.margin_bottom = Inches(0.10)

pc0 = tfc.paragraphs[0]
pc0.text = "CONVERGENCE HUB"
pc0.font.name = "Arial"
pc0.font.size = Pt(8)
pc0.font.bold = True
pc0.font.color.rgb = COLOR_TEAL_DARK
pc0.space_after = Pt(4)

pc1 = tfc.add_paragraph()
pc1.text = "AI-Powered Family Counselling"
pc1.font.name = "Arial"
pc1.font.size = Pt(13.5)
pc1.font.bold = True
pc1.font.color.rgb = COLOR_NAVY_DARK
pc1.space_after = Pt(8)

pc2 = tfc.add_paragraph()
pc2.text = "• Bridges student passion with parental risk factors"
pc2.font.name = "Arial"
pc2.font.size = Pt(9.5)
pc2.font.color.rgb = COLOR_SLATE_DARK
pc2.space_after = Pt(4)

pc3 = tfc.add_paragraph()
pc3.text = "• Personalized multi-turn joint family guidance"
pc3.font.name = "Arial"
pc3.font.size = Pt(9.5)
pc3.font.color.rgb = COLOR_SLATE_DARK

# Arrow 2 to Stage 3
tb_arr2 = slide2.shapes.add_textbox(Inches(6.60), wf_top + Inches(1.30), Inches(0.35), Inches(0.40))
tb_arr2.name = "Refined_Arr2"
tfa2 = tb_arr2.text_frame
tfa2.margin_left = tfa2.margin_right = tfa2.margin_top = tfa2.margin_bottom = 0
pa2 = tfa2.paragraphs[0]
pa2.text = "→"
pa2.alignment = PP_ALIGN.CENTER
pa2.font.name = "Arial"
pa2.font.size = Pt(22)
pa2.font.bold = True
pa2.font.color.rgb = COLOR_TEAL_ACCENT

# --- STAGE 3: EVIDENCE-BASED CAREER COMPARISON ---
c_compare = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.98), wf_top + Inches(0.42), Inches(2.65), Inches(2.25))
c_compare.name = "Refined_Node_Compare"
c_compare.fill.solid()
c_compare.fill.fore_color.rgb = COLOR_OFFWHITE
c_compare.line.color.rgb = COLOR_CARD_BORDER
c_compare.line.width = Pt(1)
tfm = c_compare.text_frame
tfm.word_wrap = True
tfm.margin_left = tfm.margin_right = Inches(0.14)
tfm.margin_top = Inches(0.20)
tfm.margin_bottom = Inches(0.10)

pm0 = tfm.paragraphs[0]
pm0.text = "DATA GROUNDING"
pm0.font.name = "Arial"
pm0.font.size = Pt(8)
pm0.font.bold = True
pm0.font.color.rgb = COLOR_SLATE_LIGHT
pm0.space_after = Pt(4)

pm1 = tfm.add_paragraph()
pm1.text = "Evidence-Based Career Comparison"
pm1.font.name = "Arial"
pm1.font.size = Pt(13.5)
pm1.font.bold = True
pm1.font.color.rgb = COLOR_NAVY_DARK
pm1.space_after = Pt(8)

pm2 = tfm.add_paragraph()
pm2.text = "• Official NSQF/NCVET accredited pathways"
pm2.font.name = "Arial"
pm2.font.size = Pt(9.5)
pm2.font.color.rgb = COLOR_SLATE_MUTED
pm2.space_after = Pt(4)

pm3 = tfm.add_paragraph()
pm3.text = "• Verified industry earnings & placement ratios"
pm3.font.name = "Arial"
pm3.font.size = Pt(9.5)
pm3.font.color.rgb = COLOR_SLATE_MUTED

# Arrow 3 to Stage 4
tb_arr3 = slide2.shapes.add_textbox(Inches(9.68), wf_top + Inches(1.30), Inches(0.35), Inches(0.40))
tb_arr3.name = "Refined_Arr3"
tfa3 = tb_arr3.text_frame
tfa3.margin_left = tfa3.margin_right = tfa3.margin_top = tfa3.margin_bottom = 0
pa3 = tfa3.paragraphs[0]
pa3.text = "→"
pa3.alignment = PP_ALIGN.CENTER
pa3.font.name = "Arial"
pa3.font.size = Pt(22)
pa3.font.bold = True
pa3.font.color.rgb = COLOR_TEAL_ACCENT

# --- STAGE 4: FAMILY CONFIDENCE ASSESSMENT & DECISION ---
c_decision = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(10.05), wf_top + Inches(0.42), Inches(2.35), Inches(2.25))
c_decision.name = "Refined_Node_Decision"
c_decision.fill.solid()
c_decision.fill.fore_color.rgb = COLOR_OFFWHITE
c_decision.line.color.rgb = COLOR_CARD_BORDER
c_decision.line.width = Pt(1)
tfd = c_decision.text_frame
tfd.word_wrap = True
tfd.margin_left = tfd.margin_right = Inches(0.14)
tfd.margin_top = Inches(0.20)
tfd.margin_bottom = Inches(0.10)

pd0 = tfd.paragraphs[0]
pd0.text = "DECISION OUTCOME"
pd0.font.name = "Arial"
pd0.font.size = Pt(8)
pd0.font.bold = True
pd0.font.color.rgb = COLOR_SLATE_LIGHT
pd0.space_after = Pt(4)

pd1 = tfd.add_paragraph()
pd1.text = "Family Decision & Confidence"
pd1.font.name = "Arial"
pd1.font.size = Pt(13.5)
pd1.font.bold = True
pd1.font.color.rgb = COLOR_NAVY_DARK
pd1.space_after = Pt(8)

pd2 = tfd.add_paragraph()
pd2.text = "• Quantified family confidence tracking"
pd2.font.name = "Arial"
pd2.font.size = Pt(9.5)
pd2.font.color.rgb = COLOR_SLATE_MUTED
pd2.space_after = Pt(4)

pd3 = tfd.add_paragraph()
pd3.text = "• High-clarity, united vocational choice"
pd3.font.name = "Arial"
pd3.font.size = Pt(9.5)
pd3.font.color.rgb = COLOR_SLATE_MUTED

# ==============================================================================
# SECTION C: FOUR CORE INNOVATIONS (Bottom 4 Columns)
# Top: 4.88", Height: 1.70"
# ==============================================================================
b_card_w = Inches(2.72)
b_gap_w = Inches(0.28)
b_top = Inches(4.88)
b_h = Inches(1.70)
b_start_x = Inches(0.80)

innovations = [
    ("01 — Family Consensus Intelligence",
     "Identifies disagreements and tracks changes in parental confidence.",
     COLOR_TEAL_TINT, COLOR_TEAL_BORDER),
    ("02 — Evidence-Grounded AI",
     "Uses placement, earnings, and NSQF progression data.",
     COLOR_OFFWHITE, COLOR_CARD_BORDER),
    ("03 — Regional & Inclusive",
     "Provides local-language guidance with simplified explanations.",
     COLOR_OFFWHITE, COLOR_CARD_BORDER),
    ("04 — Human + Admin Support",
     "Escalates unresolved cases and identifies parental resistance patterns.",
     COLOR_OFFWHITE, COLOR_CARD_BORDER)
]

for idx, (head, sub, bg_c, brd_c) in enumerate(innovations):
    bx = b_start_x + idx * (b_card_w + b_gap_w)
    
    b_card = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, bx, b_top, b_card_w, b_h)
    b_card.name = f"Refined_InnovCard_{idx}"
    b_card.fill.solid()
    b_card.fill.fore_color.rgb = bg_c
    b_card.line.color.rgb = brd_c
    b_card.line.width = Pt(1)
    
    tfb = b_card.text_frame
    tfb.word_wrap = True
    tfb.margin_left = tfb.margin_right = Inches(0.14)
    tfb.margin_top = Inches(0.18)
    tfb.margin_bottom = Inches(0.10)
    
    pb0 = tfb.paragraphs[0]
    pb0.text = head
    pb0.font.name = "Arial"
    pb0.font.size = Pt(11)
    pb0.font.bold = True
    pb0.font.color.rgb = COLOR_NAVY_DARK
    pb0.space_after = Pt(6)
    
    pb1 = tfb.add_paragraph()
    pb1.text = sub
    pb1.font.name = "Arial"
    pb1.font.size = Pt(9.5)
    pb1.font.color.rgb = COLOR_SLATE_MUTED

# Save presentation
prs.save(pptx_path)
print("Successfully saved refined Slide 2 to:", pptx_path)
