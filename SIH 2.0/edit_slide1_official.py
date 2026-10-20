import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

pptx_path = "SIH 2.0/SIH2025-IDEA-Presentation-Format.pptx"
prs = Presentation(pptx_path)

slide1 = prs.slides[0]

# Palette
NAVY_PRIMARY = RGBColor(0x0B, 0x13, 0x2B)
NAVY_TITLE = RGBColor(0x1F, 0x4E, 0x79)
TEAL_ACCENT = RGBColor(0x0D, 0x94, 0x88)
TEXT_DARK = RGBColor(0x1E, 0x29, 0x3B)
TEXT_MUTED = RGBColor(0x47, 0x55, 0x69)

# 1. Update Title 7: "SMART INDIA HACKATHON 2026"
title_shape = None
for s in slide1.shapes:
    if s.name == "Title 7":
        title_shape = s
        break

if title_shape and title_shape.has_text_frame:
    title_shape.left = Inches(0.50)
    title_shape.top = Inches(0.35)
    title_shape.width = Inches(11.20)
    title_shape.height = Inches(0.75)
    
    tf = title_shape.text_frame
    tf.word_wrap = True
    tf.clear()
    p = tf.paragraphs[0]
    p.text = "SMART INDIA HACKATHON 2026"
    p.alignment = PP_ALIGN.CENTER
    r = p.runs[0]
    r.font.name = "Garamond"
    r.font.size = Pt(38)
    r.font.bold = True
    r.font.color.rgb = NAVY_TITLE

# 2. Update Subtitle 3: Project Name & Tagline
subtitle_shape = None
for s in slide1.shapes:
    if s.name == "Subtitle 3":
        subtitle_shape = s
        break

if subtitle_shape and subtitle_shape.has_text_frame:
    subtitle_shape.left = Inches(0.50)
    subtitle_shape.top = Inches(1.18)
    subtitle_shape.width = Inches(11.20)
    subtitle_shape.height = Inches(0.85)
    
    tf = subtitle_shape.text_frame
    tf.word_wrap = True
    tf.clear()
    
    # Project Name: SkillSaathi
    p1 = tf.paragraphs[0]
    p1.alignment = PP_ALIGN.CENTER
    r1 = p1.add_run()
    r1.text = "SkillSaathi"
    r1.font.name = "Arial"
    r1.font.size = Pt(26)
    r1.font.bold = True
    r1.font.color.rgb = NAVY_PRIMARY
    p1.space_after = Pt(2)
    
    # Tagline
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = "From Career Confusion to Family Consensus"
    r2.font.name = "Arial"
    r2.font.size = Pt(15)
    r2.font.bold = True
    r2.font.color.rgb = TEAL_ACCENT

# 3. Update TextBox 9: Mandatory Pointers
tb_shape = None
for s in slide1.shapes:
    if s.name == "TextBox 9":
        tb_shape = s
        break

if tb_shape and tb_shape.has_text_frame:
    tb_shape.left = Inches(0.80)
    tb_shape.top = Inches(2.25)
    tb_shape.width = Inches(6.00)
    tb_shape.height = Inches(4.70)
    
    tf = tb_shape.text_frame
    tf.word_wrap = True
    tf.clear()
    
    fields = [
        ("Problem Statement ID – ", "26241"),
        ("Problem Statement Title – ", "AI-Enabled Career Counselling and Family Decision-Support Platform for Vocational Education"),
        ("Theme – ", "Smart Education"),
        ("PS Category – ", "Software"),
        ("Team ID – ", "[To Be Added]"),
        ("Team Name (Registered on portal) – ", "TeamChargers")
    ]
    
    for i, (label, val) in enumerate(fields):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(13)
        
        # Bullet label
        r_lbl = p.add_run()
        r_lbl.text = "•  " + label
        r_lbl.font.name = "Arial"
        r_lbl.font.size = Pt(17)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = NAVY_PRIMARY
        
        # Value
        r_val = p.add_run()
        r_val.text = val
        r_val.font.name = "Arial"
        r_val.font.size = Pt(17)
        r_val.font.bold = (label.startswith("Team Name") or label.startswith("Problem Statement ID"))
        r_val.font.color.rgb = TEAL_ACCENT if label.startswith("Theme") else TEXT_DARK

# Save updated presentation
prs.save(pptx_path)
print("Saved official-compliant presentation to:", pptx_path)
