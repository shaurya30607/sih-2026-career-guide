import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def update_slides():
    ppt_path = "SIH 2.0/SIH2025-IDEA-Presentation-Format.pptx"
    prs = Presentation(ppt_path)
    
    # =========================================================================
    # 1. ENHANCE SLIDE 2 WITH VECTOR ILLUSTRATIONS
    # =========================================================================
    slide2 = prs.slides[1]
    
    # Coordinates of the 4 cards from apply_slide2_refined.py:
    # col_w = 5.72", card_h = 1.64"
    # c1_left = 0.80", c2_left = 6.81"
    # r1_top = 3.52", r2_top = 5.26"
    
    # We remove any previous illustration shapes if re-running
    icon_names = {
        "Icon_USP_Family", "Icon_USP_India", "Icon_USP_Scale", "Icon_USP_Heatmap",
        "Logo_React_S3", "Logo_TS_S3", "Logo_Python_S3", "Logo_FastAPI_S3",
        "Icon_Brain_S3", "Icon_DB_S3", "Icon_Analytics_S3", "Logo_Vite_S3"
    }
    for s in list(slide2.shapes):
        if s.name in icon_names:
            sp = s._element
            sp.getparent().remove(sp)
            
    # Place subtle icon badges at the top-right corner of each of the 4 cards:
    # Card 1: Family Convergence (Top-Left)
    pic1 = slide2.shapes.add_picture(
        "SIH 2.0/assets/family_convergence.svg.png",
        Inches(5.78), Inches(3.58), Inches(0.64), Inches(0.64)
    )
    pic1.name = "Icon_USP_Family"

    # Card 2: India Map Outline with Location Pin (Top-Right)
    pic2 = slide2.shapes.add_picture(
        "SIH 2.0/assets/india_map.svg.png",
        Inches(11.78), Inches(3.56), Inches(0.64), Inches(0.64)
    )
    pic2.name = "Icon_USP_India"

    # Card 3: Trade-off Balance Scale (Bottom-Left)
    pic3 = slide2.shapes.add_picture(
        "SIH 2.0/assets/tradeoff_scale.svg.png",
        Inches(5.78), Inches(5.30), Inches(0.64), Inches(0.64)
    )
    pic3.name = "Icon_USP_Scale"

    # Card 4: Regional Heatmap / Analytics Grid (Bottom-Right)
    pic4 = slide2.shapes.add_picture(
        "SIH 2.0/assets/resistance_heatmap.svg.png",
        Inches(11.78), Inches(5.30), Inches(0.64), Inches(0.64)
    )
    pic4.name = "Icon_USP_Heatmap"

    # =========================================================================
    # 2. ENHANCE SLIDE 3 WITH AUTHENTIC TECH LOGOS
    # =========================================================================
    slide3 = prs.slides[2]
    
    for s in list(slide3.shapes):
        if s.name in icon_names:
            sp = s._element
            sp.getparent().remove(sp)
            
    # Stage 1 Card: Left 0.80", Top 1.24", Width 2.62"
    # Place React and TypeScript logos at the top-right of Stage 1 card
    s1_rx = Inches(0.80 + 2.62 - 0.82)
    s1_ry = Inches(1.24 + 0.12)
    logo_react = slide3.shapes.add_picture(
        "SIH 2.0/assets/react.svg.png",
        s1_rx, s1_ry, Inches(0.34), Inches(0.34)
    )
    logo_react.name = "Logo_React_S3"

    logo_ts = slide3.shapes.add_picture(
        "SIH 2.0/assets/typescript.svg.png",
        s1_rx + Inches(0.38), s1_ry, Inches(0.32), Inches(0.32)
    )
    logo_ts.name = "Logo_TS_S3"

    # Stage 2 Card: Left 3.84", Top 1.24", Width 2.62"
    # Place Python and FastAPI logos at top-right of Stage 2 card
    s2_rx = Inches(3.84 + 2.62 - 0.82)
    s2_ry = Inches(1.24 + 0.12)
    logo_py = slide3.shapes.add_picture(
        "SIH 2.0/assets/python.svg.png",
        s2_rx, s2_ry, Inches(0.33), Inches(0.33)
    )
    logo_py.name = "Logo_Python_S3"

    logo_fastapi = slide3.shapes.add_picture(
        "SIH 2.0/assets/fastapi.svg.png",
        s2_rx + Inches(0.38), s2_ry + Inches(0.02), Inches(0.31), Inches(0.31)
    )
    logo_fastapi.name = "Logo_FastAPI_S3"

    # Stage 3 Card: Left 6.88", Top 1.24", Width 2.62"
    # Place AI Neural Network / Brain icon at top-right of Stage 3 card
    s3_rx = Inches(6.88 + 2.62 - 0.52)
    s3_ry = Inches(1.24 + 0.12)
    icon_brain = slide3.shapes.add_picture(
        "SIH 2.0/assets/ai_brain.svg.png",
        s3_rx, s3_ry, Inches(0.40), Inches(0.40)
    )
    icon_brain.name = "Icon_Brain_S3"

    # Stage 4 Card: Left 9.92", Top 1.24", Width 2.62"
    # Place Analytics / Dashboard icon at top-right of Stage 4 card
    s4_rx = Inches(9.92 + 2.62 - 0.52)
    s4_ry = Inches(1.24 + 0.12)
    icon_analytics = slide3.shapes.add_picture(
        "SIH 2.0/assets/analytics.svg.png",
        s4_rx, s4_ry, Inches(0.40), Inches(0.40)
    )
    icon_analytics.name = "Icon_Analytics_S3"

    # Data Grounding Engine Box (Underneath Stage 3: Left 6.40", Top 4.04", Width 3.58", Height 0.68")
    # Place Database Cylinder icon on the left inside the box
    icon_db = slide3.shapes.add_picture(
        "SIH 2.0/assets/database.svg.png",
        Inches(6.48), Inches(4.14), Inches(0.44), Inches(0.44)
    )
    icon_db.name = "Icon_DB_S3"

    # Bottom Technology Strip: Left 0.80", Top 4.96"
    # In Column 1 (Frontend, Left 0.92", Top 5.30", Width 2.78")
    # Place Vite logo at the top right of the Frontend column card
    logo_vite = slide3.shapes.add_picture(
        "SIH 2.0/assets/vite.svg.png",
        Inches(0.92 + 2.78 - 0.42), Inches(5.36), Inches(0.30), Inches(0.30)
    )
    logo_vite.name = "Logo_Vite_S3"

    prs.save(ppt_path)
    print("Successfully enhanced Slides 2 & 3 with visuals and logos in:", ppt_path)

if __name__ == "__main__":
    update_slides()
