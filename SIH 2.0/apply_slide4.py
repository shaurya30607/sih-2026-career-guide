import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def update_slide4():
    ppt_path = "SIH 2.0/SIH2025-IDEA-Presentation-Format.pptx"
    prs = Presentation(ppt_path)
    
    NAVY = RGBColor(11, 19, 43)          # #0B132B
    TEAL = RGBColor(13, 148, 136)        # #0D9488
    DARK_TEAL = RGBColor(15, 118, 110)   # #0F766E
    LIGHT_TEAL = RGBColor(240, 253, 250) # #F0FDFA
    BORDER_TEAL = RGBColor(45, 212, 191) # #2DD4BF
    WHITE = RGBColor(255, 255, 255)
    GRAY_TEXT = RGBColor(71, 85, 105)    # #475569
    SLATE_LIGHT = RGBColor(100, 116, 139)# #64748B
    CARD_BG = RGBColor(248, 250, 252)    # #F8FAFC
    BORDER_GRAY = RGBColor(203, 213, 225)# #CBD5E1
    BORDER_LIGHT = RGBColor(226, 232, 240)# #E2E8F0
    AMBER_TEXT = RGBColor(146, 64, 14)   # #92400E
    BLUE_BG = RGBColor(224, 242, 254)    # #E0F2FE
    BLUE_TEXT = RGBColor(7, 89, 133)     # #075985
    GREEN_TEXT = RGBColor(22, 101, 52)   # #166534
    ROSE_TEXT = RGBColor(159, 18, 57)    # #9F1239
    PURPLE_TEXT = RGBColor(79, 70, 229)  # #4F46E5

    # =========================================================================
    # REBUILD SLIDE 4 (Index 3) - STRICT SIH COMPLIANCE
    # =========================================================================
    slide4 = prs.slides[3]
    preserved_names = {"Rectangle 9", "Slide Number Placeholder 5", "Footer Placeholder 6", "Picture 2"}
    for s in list(slide4.shapes):
        if s.name not in preserved_names:
            sp = s._element
            sp.getparent().remove(sp)

    # 1. Team Name Badge (Left-aligned, standard placement)
    team_badge4 = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.80), Inches(0.32), Inches(1.65), Inches(0.42))
    team_badge4.fill.solid()
    team_badge4.fill.fore_color.rgb = WHITE
    team_badge4.line.color.rgb = BORDER_TEAL
    team_badge4.line.width = Pt(1.5)
    tf = team_badge4.text_frame
    tf.word_wrap = False
    tf.margin_top = Inches(0.06)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "TeamChargers"
    r.font.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = NAVY

    # 2. Official Slide Title - STRICTLY "FEASIBILITY AND VIABILITY"
    # No extra subtitles added to the title box, maintaining 100% compliance with SIH template heading
    title_box4 = slide4.shapes.add_textbox(Inches(2.65), Inches(0.28), Inches(7.50), Inches(0.45))
    tf = title_box4.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "FEASIBILITY AND VIABILITY"
    r.font.name = "Georgia"
    r.font.size = Pt(21)
    r.font.bold = True
    r.font.color.rgb = NAVY

    # 3. Sub-header Description (contains execution context without disturbing heading or touching SIH logo)
    desc_box4 = slide4.shapes.add_textbox(Inches(0.80), Inches(0.82), Inches(9.80), Inches(0.28))
    tf = desc_box4.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Execution Roadmap & Risk Mitigation — "
    r.font.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = DARK_TEAL
    r = p.add_run()
    r.text = "Demonstrating practical implementation through a 3-phase progression, grounded data controls, and modular architecture."
    r.font.size = Pt(10.5)
    r.font.color.rgb = GRAY_TEXT

    # =========================================================================
    # SECTION 1: 3-PHASE IMPLEMENTATION ROADMAP (Left 55% -> width = 6.45")
    # =========================================================================
    sec1_lbl = slide4.shapes.add_textbox(Inches(0.80), Inches(1.18), Inches(6.45), Inches(0.26))
    p = sec1_lbl.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = "SECTION 1: 3-PHASE IMPLEMENTATION ROADMAP"
    r.font.size = Pt(9.0)
    r.font.bold = True
    r.font.color.rgb = NAVY
    r = p.add_run()
    r.text = "   [ From Live Prototype to Phased Field Deployment ]"
    r.font.size = Pt(8.0)
    r.font.italic = True
    r.font.color.rgb = TEAL

    c_left = Inches(0.80)
    c_w = Inches(6.45)
    c_h = Inches(1.36)
    
    phases = [
        {
            "top": Inches(1.48),
            "num": "01",
            "badge_color": TEAL,
            "bg_color": LIGHT_TEAL,
            "border_color": BORDER_TEAL,
            "border_w": Pt(1.5),
            "title": "PHASE 01 — FOUNDATION",
            "tag": "[ ACTIVE IN CODE ]",
            "tag_color": GREEN_TEXT,
            "icon": "SIH 2.0/assets/phase1_code.svg.png",
            "icon_w": Inches(0.72),
            "icon_h": Inches(0.48),
            "bullets": [
                ("React + TypeScript Frontend: ", "Operational responsive UI, student & parent intake forms, trade explorer"),
                ("Student-Parent Interface: ", "Dual-persona guided workflow capturing individual aspirations & concerns"),
                ("Course Discovery Engine: ", "Structured vocational datasets & local ITI benchmark feeds")
            ]
        },
        {
            "top": Inches(3.06),
            "num": "02",
            "badge_color": RGBColor(2, 132, 199), # Sky Blue
            "bg_color": WHITE,
            "border_color": TEAL,
            "border_w": Pt(1.2),
            "title": "PHASE 02 — INTELLIGENCE",
            "tag": "[ PLANNED ARCHITECTURE ]",
            "tag_color": BLUE_TEXT,
            "icon": "SIH 2.0/assets/phase2_ai.svg.png",
            "icon_w": Inches(0.55),
            "icon_h": Inches(0.55),
            "bullets": [
                ("FastAPI Backend Integration (Planned): ", "REST API client, session state persistence & data contracts"),
                ("Evidence-Grounded AI Engine (Planned): ", "RAG counselling bound to NCVET qualification packs & trade criteria"),
                ("Consensus & Confidence Metrics (Planned): ", "Quantified family agreement scoring & objection tracking")
            ]
        },
        {
            "top": Inches(4.64),
            "num": "03",
            "badge_color": RGBColor(51, 65, 85), # Slate
            "bg_color": CARD_BG,
            "border_color": BORDER_GRAY,
            "border_w": Pt(1.0),
            "title": "PHASE 03 — DEPLOYMENT",
            "tag": "[ FUTURE PILOT ROLLOUT ]",
            "tag_color": GRAY_TEXT,
            "icon": "SIH 2.0/assets/phase3_deploy.svg.png",
            "icon_w": Inches(0.46),
            "icon_h": Inches(0.58),
            "bullets": [
                ("Vernacular Language Support (Planned): ", "English + Hindi/regional language testing for inclusive rural access"),
                ("Human Counsellor Escalation (Planned): ", "Automated triage routing high-resistance families to 1-on-1 counsellors"),
                ("Admin Heatmaps & Pilot Testing (Planned): ", "District-level aggregate hesitation insights for institutional planners")
            ]
        }
    ]

    for idx, p_info in enumerate(phases):
        # Card container
        card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_left, p_info["top"], c_w, c_h)
        card.fill.solid()
        card.fill.fore_color.rgb = p_info["bg_color"]
        card.line.color.rgb = p_info["border_color"]
        card.line.width = p_info["border_w"]

        # Milestone Badge (Wide Rounded Pill)
        m_badge = slide4.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            c_left + Inches(0.12),
            p_info["top"] + Inches(0.12),
            Inches(0.42),
            Inches(0.32)
        )
        m_badge.fill.solid()
        m_badge.fill.fore_color.rgb = p_info["badge_color"]
        m_badge.line.fill.background()
        tf_m = m_badge.text_frame
        tf_m.word_wrap = False
        tf_m.margin_left = 0
        tf_m.margin_right = 0
        tf_m.margin_top = Inches(0.04)
        tf_m.margin_bottom = 0
        p = tf_m.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = p_info["num"]
        r.font.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = WHITE

        # Header Text
        h_box = slide4.shapes.add_textbox(c_left + Inches(0.60), p_info["top"] + Inches(0.08), Inches(4.80), Inches(0.30))
        tf_h = h_box.text_frame
        tf_h.margin_left = 0
        tf_h.margin_top = 0
        p = tf_h.paragraphs[0]
        r = p.add_run()
        r.text = p_info["title"] + "  "
        r.font.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = NAVY
        r = p.add_run()
        r.text = p_info["tag"]
        r.font.bold = True
        r.font.size = Pt(7.5)
        r.font.color.rgb = p_info["tag_color"]

        # Vertically Centered Icon on the right
        if os.path.exists(p_info["icon"]):
            icon_top = p_info["top"] + (c_h - p_info["icon_h"]) / 2
            slide4.shapes.add_picture(
                p_info["icon"],
                c_left + Inches(5.62),
                icon_top,
                p_info["icon_w"],
                p_info["icon_h"]
            )

        # Bullets (constrained width to 5.0" to give ample margin before right icon)
        b_box = slide4.shapes.add_textbox(c_left + Inches(0.60), p_info["top"] + Inches(0.36), Inches(4.95), Inches(0.92))
        tf_b = b_box.text_frame
        tf_b.word_wrap = True
        tf_b.margin_left = 0
        tf_b.margin_top = 0
        for b_idx, (b_title, b_desc) in enumerate(p_info["bullets"]):
            p = tf_b.paragraphs[0] if b_idx == 0 else tf_b.add_paragraph()
            p.space_after = Pt(2)
            r = p.add_run()
            r.text = "• " + b_title
            r.font.bold = True
            r.font.size = Pt(7.8)
            r.font.color.rgb = NAVY
            r = p.add_run()
            r.text = b_desc
            r.font.size = Pt(7.5)
            r.font.color.rgb = GRAY_TEXT

    # Progression Connector Arrows
    arr1 = slide4.shapes.add_textbox(c_left + Inches(0.21), Inches(2.83), Inches(0.24), Inches(0.22))
    p = arr1.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "↓"
    r.font.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = TEAL

    arr2 = slide4.shapes.add_textbox(c_left + Inches(0.21), Inches(4.41), Inches(0.24), Inches(0.22))
    p = arr2.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "↓"
    r.font.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = TEAL

    # =========================================================================
    # SECTION 2: CHALLENGES & MITIGATION MATRIX (Right 45% -> width = 5.00")
    # =========================================================================
    m_left = Inches(7.53)
    m_w = Inches(5.00)
    
    sec2_lbl = slide4.shapes.add_textbox(m_left, Inches(1.18), m_w, Inches(0.26))
    p = sec2_lbl.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = "SECTION 2: CHALLENGES & MITIGATION MATRIX"
    r.font.size = Pt(9.0)
    r.font.bold = True
    r.font.color.rgb = NAVY
    r = p.add_run()
    r.text = "   [ Technical Safeguards & Risk Controls ]"
    r.font.size = Pt(8.0)
    r.font.italic = True
    r.font.color.rgb = TEAL

    r_h = Inches(0.80)
    r_gap = Inches(0.11)
    
    matrix = [
        {
            "top": Inches(1.48),
            "icon": "SIH 2.0/assets/risk_data.svg.png",
            "ch_title": "Limited Verified Data",
            "ch_sub": "Uncurated regional wage reports",
            "mit_title": "Evaluation → Official Data",
            "mit_sub": "Structured data transitioning to NCVET/NCS APIs",
            "accent": AMBER_TEXT
        },
        {
            "top": Inches(2.39),
            "icon": "SIH 2.0/assets/risk_shield.svg.png",
            "ch_title": "AI Hallucinations",
            "ch_sub": "Unverified claims or false guarantees",
            "mit_title": "Grounded RAG + Source Citations",
            "mit_sub": "Strict retrieval bounds with source references",
            "accent": ROSE_TEXT
        },
        {
            "top": Inches(3.30),
            "icon": "SIH 2.0/assets/risk_lang.svg.png",
            "ch_title": "Language Barriers",
            "ch_sub": "Vernacular hesitation among rural parents",
            "mit_title": "Regional UI + Simplified Layout",
            "mit_sub": "Hindi & regional testing with plain terms",
            "accent": BLUE_TEXT
        },
        {
            "top": Inches(4.21),
            "icon": "SIH 2.0/assets/risk_privacy.svg.png",
            "ch_title": "Data Privacy Concerns",
            "ch_sub": "Sensitive family financial ambitions",
            "mit_title": "Consent Intake + Minimal Retention",
            "mit_sub": "Role-based consent with zero session leakage",
            "accent": PURPLE_TEXT
        },
        {
            "top": Inches(5.12),
            "icon": "SIH 2.0/assets/risk_scale.svg.png",
            "ch_title": "System Scalability",
            "ch_sub": "Traffic surges during admission cycles",
            "mit_title": "Modular Microservices Architecture",
            "mit_sub": "Stateless backend design + phased regional rollout",
            "accent": DARK_TEAL
        }
    ]

    for row_info in matrix:
        # Card container
        rcard = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, m_left, row_info["top"], m_w, r_h)
        rcard.fill.solid()
        rcard.fill.fore_color.rgb = WHITE
        rcard.line.color.rgb = BORDER_LIGHT
        rcard.line.width = Pt(1.0)

        # Left accent pill
        pill = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, m_left + Inches(0.04), row_info["top"] + Inches(0.08), Inches(0.06), Inches(0.64))
        pill.fill.solid()
        pill.fill.fore_color.rgb = row_info["accent"]
        pill.line.fill.background()

        # Icon
        if os.path.exists(row_info["icon"]):
            slide4.shapes.add_picture(
                row_info["icon"],
                m_left + Inches(0.16),
                row_info["top"] + Inches(0.18),
                Inches(0.42),
                Inches(0.44)
            )

        # Challenge Column (Left)
        ch_box = slide4.shapes.add_textbox(m_left + Inches(0.64), row_info["top"] + Inches(0.09), Inches(1.85), Inches(0.62))
        tf_ch = ch_box.text_frame
        tf_ch.word_wrap = True
        tf_ch.margin_left = 0
        tf_ch.margin_top = 0
        p = tf_ch.paragraphs[0]
        r = p.add_run()
        r.text = row_info["ch_title"] + "\n"
        r.font.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = NAVY
        r = p.add_run()
        r.text = row_info["ch_sub"]
        r.font.size = Pt(7.0)
        r.font.color.rgb = GRAY_TEXT

        # Center Arrow
        arr = slide4.shapes.add_textbox(m_left + Inches(2.48), row_info["top"] + Inches(0.24), Inches(0.24), Inches(0.28))
        tf_arr = arr.text_frame
        p = tf_arr.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = "→"
        r.font.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = TEAL

        # Mitigation Column (Right)
        mit_box = slide4.shapes.add_textbox(m_left + Inches(2.74), row_info["top"] + Inches(0.09), Inches(2.18), Inches(0.62))
        tf_mit = mit_box.text_frame
        tf_mit.word_wrap = True
        tf_mit.margin_left = 0
        tf_mit.margin_top = 0
        p = tf_mit.paragraphs[0]
        r = p.add_run()
        r.text = row_info["mit_title"] + "\n"
        r.font.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = DARK_TEAL
        r = p.add_run()
        r.text = row_info["mit_sub"]
        r.font.size = Pt(7.0)
        r.font.color.rgb = GRAY_TEXT

    # =========================================================================
    # SECTION 3: FEASIBILITY HIGHLIGHT BANNER (Bottom width = 11.733")
    # =========================================================================
    banner = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.80), Inches(6.16), Inches(11.733), Inches(0.48))
    banner.fill.solid()
    banner.fill.fore_color.rgb = LIGHT_TEAL
    banner.line.color.rgb = BORDER_TEAL
    banner.line.width = Pt(1.5)
    
    tf_ban = banner.text_frame
    tf_ban.margin_top = Inches(0.08)
    tf_ban.word_wrap = False
    p = tf_ban.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER

    r = p.add_run()
    r.text = "FEASIBILITY PRINCIPLE:  "
    r.font.name = "Georgia"
    r.font.bold = True
    r.font.size = Pt(9.8)
    r.font.color.rgb = NAVY

    r = p.add_run()
    r.text = "Existing Frontend"
    r.font.bold = True
    r.font.size = Pt(9.8)
    r.font.color.rgb = DARK_TEAL

    r = p.add_run()
    r.text = "  +  "
    r.font.bold = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = SLATE_LIGHT

    r = p.add_run()
    r.text = "Modular AI Architecture"
    r.font.bold = True
    r.font.size = Pt(9.8)
    r.font.color.rgb = DARK_TEAL

    r = p.add_run()
    r.text = "  +  "
    r.font.bold = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = SLATE_LIGHT

    r = p.add_run()
    r.text = "Phased Deployment"
    r.font.bold = True
    r.font.size = Pt(9.8)
    r.font.color.rgb = DARK_TEAL

    r = p.add_run()
    r.text = "  =  "
    r.font.bold = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = SLATE_LIGHT

    r = p.add_run()
    r.text = "Practical Implementation"
    r.font.bold = True
    r.font.size = Pt(10.2)
    r.font.color.rgb = NAVY

    prs.save(ppt_path)
    print("Successfully rebuilt Slide 4 with strict heading compliance in:", ppt_path)

if __name__ == "__main__":
    update_slide4()
