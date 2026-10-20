import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

pptx_path = "SIH 2.0/SIH2025-IDEA-Presentation-Format.pptx"
prs = Presentation(pptx_path)

slide2 = prs.slides[1]

# Palette - Minimalist Editorial
COLOR_NAVY_DARK = RGBColor(0x0B, 0x13, 0x2B)
COLOR_NAVY_TITLE = RGBColor(0x1F, 0x4E, 0x79)
COLOR_SLATE_DARK = RGBColor(0x1E, 0x29, 0x3B)
COLOR_SLATE_MUTED = RGBColor(0x47, 0x55, 0x69)

COLOR_TEAL_DARK = RGBColor(0x0D, 0x94, 0x88)
COLOR_TEAL_ACCENT = RGBColor(0x14, 0xB8, 0xA6)
COLOR_TEAL_TINT = RGBColor(0xF0, 0xFD, 0xFA)
COLOR_TEAL_BORDER = RGBColor(0x0D, 0x94, 0x88)

COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_OFFWHITE = RGBColor(0xF8, 0xF9, 0xFA)
COLOR_CARD_BORDER = RGBColor(0xCB, 0xD5, 0xE1)
COLOR_SUBTLE_BORDER = RGBColor(0xE2, 0xE8, 0xF0)

# 1. Clean slide2 to only the 6 official base template shapes
base_names = {'Rectangle 8', 'Title 1', 'Slide Number Placeholder 5', 'Footer Placeholder 6', 'Oval 9', 'Picture 2'}
to_remove = [s for s in slide2.shapes if s.name not in base_names]
for s in to_remove:
    sp = s._element
    sp.getparent().remove(sp)

# 2. Update Oval 9 (Team Badge)
for s in slide2.shapes:
    if s.name == "Oval 9":
        s.left = Inches(0.80)
        s.top = Inches(0.30)
        s.width = Inches(1.70)
        s.height = Inches(0.55)
        s.fill.solid()
        s.fill.fore_color.rgb = COLOR_OFFWHITE
        s.line.color.rgb = COLOR_CARD_BORDER
        s.line.width = Pt(1)
        if s.has_text_frame:
            tf = s.text_frame
            tf.clear()
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            p = tf.paragraphs[0]
            p.text = "TeamChargers"
            p.alignment = PP_ALIGN.CENTER
            p.font.name = "Arial"
            p.font.size = Pt(11)
            p.font.bold = True
            p.font.color.rgb = COLOR_NAVY_DARK

# 3. Update Title 1 (Heading + Project Title)
for s in slide2.shapes:
    if s.name == "Title 1":
        s.left = Inches(2.65)
        s.top = Inches(0.16)
        s.width = Inches(7.95)
        s.height = Inches(0.95)
        if s.has_text_frame:
            tf = s.text_frame
            tf.word_wrap = True
            tf.clear()
            # Mandatory Template Heading: IDEA TITLE
            p1 = tf.paragraphs[0]
            p1.alignment = PP_ALIGN.CENTER
            r1 = p1.add_run()
            r1.text = "IDEA TITLE: "
            r1.font.name = "Garamond"
            r1.font.size = Pt(22)
            r1.font.bold = True
            r1.font.color.rgb = COLOR_NAVY_TITLE
            
            r2 = p1.add_run()
            r2.text = "SkillSaathi"
            r2.font.name = "Arial"
            r2.font.size = Pt(22)
            r2.font.bold = True
            r2.font.color.rgb = COLOR_NAVY_DARK
            p1.space_after = Pt(2)
            
            # Subtitle
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.CENTER
            r3 = p2.add_run()
            r3.text = "Beyond Career Advice, Towards Family Consensus"
            r3.font.name = "Arial"
            r3.font.size = Pt(14)
            r3.font.italic = True
            r3.font.color.rgb = COLOR_TEAL_DARK

# ==============================================================================
# SECTION A: ONE-LINE DESCRIPTION
# ==============================================================================
tb_desc = slide2.shapes.add_textbox(Inches(0.80), Inches(1.30), Inches(11.733), Inches(0.45))
tf_desc = tb_desc.text_frame
tf_desc.word_wrap = True
tf_desc.margin_left = tf_desc.margin_right = tf_desc.margin_top = tf_desc.margin_bottom = 0
p_desc = tf_desc.paragraphs[0]
p_desc.alignment = PP_ALIGN.CENTER
r_desc = p_desc.add_run()
r_desc.text = "An AI-powered platform helping students and parents make informed vocational career decisions together."
r_desc.font.name = "Arial"
r_desc.font.size = Pt(15.5)
r_desc.font.color.rgb = COLOR_SLATE_DARK

# ==============================================================================
# SECTION B: MAIN VISUAL — FAMILY CONSENSUS INTELLIGENCE
# ONE Large Horizontal Process Diagram: 3 Dominant Stages
# STUDENT + PARENT → AI COUNSELLING → FAMILY DECISION
# ==============================================================================
stage_w = Inches(3.35)
arrow_w = Inches(0.84)
stage_h = Inches(2.35)
diagram_y = Inches(2.10)
start_x = Inches(0.80)

stages_data = [
    {
        "title": "STUDENT + PARENT",
        "label": "Aspirations & Concerns",
        "bg": COLOR_WHITE,
        "border": COLOR_CARD_BORDER,
        "border_width": Pt(1),
        "title_color": COLOR_NAVY_DARK,
        "label_color": COLOR_SLATE_MUTED
    },
    {
        "title": "AI COUNSELLING",
        "label": "Verified Data & Personalized Guidance",
        "bg": COLOR_TEAL_TINT,
        "border": COLOR_TEAL_BORDER,
        "border_width": Pt(1.5),
        "title_color": COLOR_NAVY_DARK,
        "label_color": COLOR_TEAL_DARK
    },
    {
        "title": "FAMILY DECISION",
        "label": "Confidence & Career Clarity",
        "bg": COLOR_WHITE,
        "border": COLOR_CARD_BORDER,
        "border_width": Pt(1),
        "title_color": COLOR_NAVY_DARK,
        "label_color": COLOR_NAVY_DARK
    }
]

for idx, st in enumerate(stages_data):
    sx = start_x + idx * (stage_w + arrow_w)
    
    # Stage Box
    box = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, sx, diagram_y, stage_w, stage_h)
    box.fill.solid()
    box.fill.fore_color.rgb = st["bg"]
    box.line.color.rgb = st["border"]
    box.line.width = st["border_width"]
    
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.20)
    tf.margin_right = Inches(0.20)
    tf.margin_top = Inches(0.60)
    tf.margin_bottom = Inches(0.20)
    
    # Stage Title (Large & Bold)
    p0 = tf.paragraphs[0]
    p0.alignment = PP_ALIGN.CENTER
    p0.text = st["title"]
    p0.font.name = "Arial"
    p0.font.size = Pt(21)
    p0.font.bold = True
    p0.font.color.rgb = st["title_color"]
    p0.space_after = Pt(14)
    
    # Supporting Label
    p1 = tf.add_paragraph()
    p1.alignment = PP_ALIGN.CENTER
    p1.text = st["label"]
    p1.font.name = "Arial"
    p1.font.size = Pt(14)
    p1.font.bold = True
    p1.font.color.rgb = st["label_color"]
    
    # Large Connector Arrow
    if idx < 2:
        ax = sx + stage_w
        tb_arr = slide2.shapes.add_textbox(ax, diagram_y + Inches(0.85), arrow_w, Inches(0.65))
        tf_a = tb_arr.text_frame
        tf_a.margin_left = tf_a.margin_right = tf_a.margin_top = tf_a.margin_bottom = 0
        pa = tf_a.paragraphs[0]
        pa.text = "→"
        pa.alignment = PP_ALIGN.CENTER
        pa.font.name = "Arial"
        pa.font.size = Pt(32)
        pa.font.bold = True
        pa.font.color.rgb = COLOR_TEAL_ACCENT

# ==============================================================================
# SECTION C: THREE INNOVATION HIGHLIGHTS
# 01 — Evidence-Based Guidance | 02 — Family Confidence Tracking | 03 — Human Counsellor Support
# ==============================================================================
card_w = Inches(3.64)
gap_w = Inches(0.40)
innov_y = Inches(4.95)
innov_h = Inches(1.50)

innovations = [
    ("01 — Evidence-Based Guidance", "Verified earnings, placements & NSQF pathways."),
    ("02 — Family Confidence Tracking", "Measures parental concerns before and after counselling."),
    ("03 — Human Counsellor Support", "Escalates unresolved concerns to real counsellors.")
]

for idx, (head, body) in enumerate(innovations):
    ix = start_x + idx * (card_w + gap_w)
    
    ibox = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, ix, innov_y, card_w, innov_h)
    ibox.fill.solid()
    ibox.fill.fore_color.rgb = COLOR_OFFWHITE
    ibox.line.color.rgb = COLOR_SUBTLE_BORDER
    ibox.line.width = Pt(1)
    
    tf_i = ibox.text_frame
    tf_i.word_wrap = True
    tf_i.margin_left = Inches(0.24)
    tf_i.margin_right = Inches(0.24)
    tf_i.margin_top = Inches(0.28)
    tf_i.margin_bottom = Inches(0.20)
    
    # Highlight Title
    pi0 = tf_i.paragraphs[0]
    pi0.text = head
    pi0.font.name = "Arial"
    pi0.font.size = Pt(13.5)
    pi0.font.bold = True
    pi0.font.color.rgb = COLOR_NAVY_DARK
    pi0.space_after = Pt(8)
    
    # Highlight Text
    pi1 = tf_i.add_paragraph()
    pi1.text = body
    pi1.font.name = "Arial"
    pi1.font.size = Pt(11.5)
    pi1.font.color.rgb = COLOR_SLATE_MUTED

# Save updated presentation
prs.save(pptx_path)
print("Successfully saved minimal Slide 2 to:", pptx_path)
