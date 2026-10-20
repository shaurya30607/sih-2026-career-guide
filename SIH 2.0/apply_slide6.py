import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def update_slide6():
    ppt_path = "SIH 2.0/SIH2025-IDEA-Presentation-Format.pptx"
    prs = Presentation(ppt_path)
    
    NAVY = RGBColor(11, 19, 43)          # #0B132B
    TEAL = RGBColor(13, 148, 136)        # #0D9488
    DARK_TEAL = RGBColor(15, 118, 110)   # #0F766E
    LIGHT_TEAL = RGBColor(240, 253, 250) # #F0FDFA
    BORDER_TEAL = RGBColor(45, 212, 191) # #2DD4BF
    WHITE = RGBColor(255, 255, 255)
    GRAY_TEXT = RGBColor(71, 85, 105)    # #475569
    SLATE_LIGHT = RGBColor(100, 116, 139)# #64748B
    CARD_BG = RGBColor(248, 250, 252)    # #F8FAFC
    BORDER_LIGHT = RGBColor(226, 232, 240)# #E2E8F0
    LINK_BLUE = RGBColor(2, 132, 199)    # #0284C7
    PURPLE_TEXT = RGBColor(79, 70, 229)  # #4F46E5

    # =========================================================================
    # REBUILD SLIDE 6 (Index 5) - STRICT SIH COMPLIANCE & HIGH READABILITY
    # =========================================================================
    slide6 = prs.slides[5]
    preserved_names = {"Rectangle 9", "Slide Number Placeholder 5", "Footer Placeholder 6", "Picture 2"}
    for s in list(slide6.shapes):
        if s.name not in preserved_names:
            sp = s._element
            sp.getparent().remove(sp)

    # 1. Team Name Badge (Standard placement)
    team_badge6 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.80), Inches(0.32), Inches(1.65), Inches(0.42))
    team_badge6.fill.solid()
    team_badge6.fill.fore_color.rgb = WHITE
    team_badge6.line.color.rgb = BORDER_TEAL
    team_badge6.line.width = Pt(1.5)
    tf = team_badge6.text_frame
    tf.word_wrap = False
    tf.margin_top = Inches(0.06)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "TeamChargers"
    r.font.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = NAVY

    # 2. Official Slide Title - STRICTLY "RESEARCH AND REFERENCES"
    title_box6 = slide6.shapes.add_textbox(Inches(2.65), Inches(0.28), Inches(7.50), Inches(0.45))
    tf = title_box6.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "RESEARCH AND REFERENCES"
    r.font.name = "Georgia"
    r.font.size = Pt(21)
    r.font.bold = True
    r.font.color.rgb = NAVY

    # 3. Sub-header Description (Constrained width = 8.50" to guarantee zero overlap with SIH logo)
    desc_box6 = slide6.shapes.add_textbox(Inches(0.80), Inches(0.80), Inches(8.50), Inches(0.36))
    tf = desc_box6.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_top = 0
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Foundational Policy & Academic Frameworks — "
    r.font.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = DARK_TEAL
    r = p.add_run()
    r.text = "Grounding SkillSaathi's vocational taxonomies, wage baselines, and family mediation models in verified standards."
    r.font.size = Pt(9.5)
    r.font.color.rgb = GRAY_TEXT

    # =========================================================================
    # SECTION 1: VERIFIED IEEE CITATIONS (Left Side -> width = 7.70")
    # =========================================================================
    sec1_lbl = slide6.shapes.add_textbox(Inches(0.80), Inches(1.18), Inches(7.70), Inches(0.26))
    p = sec1_lbl.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = "SECTION 1: OFFICIAL FRAMEWORKS & RESEARCH REFERENCES"
    r.font.size = Pt(10.0)
    r.font.bold = True
    r.font.color.rgb = NAVY
    r = p.add_run()
    r.text = "   [ IEEE Citation Format ]"
    r.font.size = Pt(8.5)
    r.font.italic = True
    r.font.color.rgb = TEAL

    ref_w = Inches(7.70)
    ref_h = Inches(0.66)
    ref_gap = Inches(0.08)
    ref_top_start = Inches(1.46)

    references = [
        {
            "num": "[1]",
            "author": "National Council for Vocational Education and Training (NCVET), ",
            "title": "“National Skills Qualification Framework (NSQF),” ",
            "venue": "MSDE, Govt. of India, 2023. [Online]. Available: ",
            "url": "https://ncvet.gov.in/national-skills-qualification-framework/",
            "accessed": " [Accessed: Oct. 19, 2026].",
            "tag": "NCrF-Aligned Qualification Levels & Progression"
        },
        {
            "num": "[2]",
            "author": "Ministry of Skill Development and Entrepreneurship (MSDE), ",
            "title": "“Schemes and National Skill Policy,” ",
            "venue": "Government of India. [Online]. Available: ",
            "url": "https://www.msde.gov.in/",
            "accessed": " [Accessed: Oct. 19, 2026].",
            "tag": "National Vocational Mandates & Scheme Norms"
        },
        {
            "num": "[3]",
            "author": "Directorate General of Training (DGT), ",
            "title": "“Craftsmen Training Scheme (CTS),” ",
            "venue": "MSDE, Govt. of India. [Online]. Available: ",
            "url": "https://dgt.gov.in/en/CTS",
            "accessed": " [Accessed: Oct. 19, 2026].",
            "tag": "ITI Vocational Trades, Curricula & Duration"
        },
        {
            "num": "[4]",
            "author": "Directorate General of Employment (DGE), ",
            "title": "“National Career Service (NCS) Project,” ",
            "venue": "Ministry of Labour and Employment, Govt. of India. [Online]. Available: ",
            "url": "https://www.ncs.gov.in/",
            "accessed": " [Accessed: Oct. 19, 2026].",
            "tag": "Career Profiles, Sector Demand & Wage Benchmarks"
        },
        {
            "num": "[5]",
            "author": "Directorate General of Training (DGT), ",
            "title": "“Bharat Skills Portal,” ",
            "venue": "MSDE, Govt. of India. [Online]. Available: ",
            "url": "https://bharatskill.gov.in/",
            "accessed": " [Accessed: Oct. 19, 2026].",
            "tag": "Trade Learning Repositories & QP/NOS Question Banks"
        },
        {
            "num": "[6]",
            "author": "J. Dietrich and B. Kracke, ",
            "title": "“Career-specific parental behaviors in adolescents’ development,” ",
            "venue": "Journal of Vocational Behavior, vol. 75, no. 2, pp. 109–119, Oct. 2009. doi: ",
            "url": "https://doi.org/10.1016/j.jvb.2009.03.005",
            "accessed": "",
            "tag": "Parental Influence & Family Consensus Research"
        }
    ]

    for idx, ref in enumerate(references):
        r_top = ref_top_start + idx * (ref_h + ref_gap)

        # Card container
        rcard = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.80), r_top, ref_w, ref_h)
        rcard.fill.solid()
        rcard.fill.fore_color.rgb = WHITE
        rcard.line.color.rgb = BORDER_LIGHT
        rcard.line.width = Pt(1.0)

        # Left Number Badge [1]..[6]
        badge = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.90), r_top + Inches(0.12), Inches(0.46), Inches(0.42))
        badge.fill.solid()
        badge.fill.fore_color.rgb = LIGHT_TEAL
        badge.line.color.rgb = BORDER_TEAL
        badge.line.width = Pt(1.0)
        tf_bg = badge.text_frame
        tf_bg.margin_left = 0
        tf_bg.margin_right = 0
        tf_bg.margin_top = Inches(0.08)
        p_bg = tf_bg.paragraphs[0]
        p_bg.alignment = PP_ALIGN.CENTER
        r = p_bg.add_run()
        r.text = ref["num"]
        r.font.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = DARK_TEAL

        # Reference Citation Content
        tb = slide6.shapes.add_textbox(Inches(1.46), r_top + Inches(0.05), Inches(6.92), Inches(0.56))
        tf_r = tb.text_frame
        tf_r.word_wrap = True
        tf_r.margin_left = 0
        tf_r.margin_top = 0

        # Line 1: Author & Title + Scope Tag (Readable: 9.8pt)
        p1 = tf_r.paragraphs[0]
        p1.space_after = Pt(1.5)
        r = p1.add_run()
        r.text = ref["author"]
        r.font.bold = True
        r.font.size = Pt(9.8)
        r.font.color.rgb = NAVY

        r = p1.add_run()
        r.text = ref["title"]
        r.font.bold = True
        r.font.size = Pt(9.8)
        r.font.color.rgb = DARK_TEAL

        r = p1.add_run()
        r.text = f"  ({ref['tag']})"
        r.font.size = Pt(7.5)
        r.font.italic = True
        r.font.color.rgb = SLATE_LIGHT

        # Line 2: Venue, Clickable URL, and Accessed Date (Readable: 8.8pt)
        p2 = tf_r.add_paragraph()
        r = p2.add_run()
        r.text = ref["venue"]
        r.font.size = Pt(8.8)
        r.font.color.rgb = GRAY_TEXT

        r = p2.add_run()
        r.text = ref["url"]
        r.font.size = Pt(8.8)
        r.font.underline = True
        r.font.color.rgb = LINK_BLUE
        r.hyperlink.address = ref["url"]

        if ref["accessed"]:
            r = p2.add_run()
            r.text = ref["accessed"]
            r.font.size = Pt(8.2)
            r.font.italic = True
            r.font.color.rgb = GRAY_TEXT

    # =========================================================================
    # SECTION 2: REFINED DATA GROUNDING ARCHITECTURE (Right Side -> width = 3.75")
    # =========================================================================
    sec2_lbl = slide6.shapes.add_textbox(Inches(8.78), Inches(1.18), Inches(3.75), Inches(0.26))
    p = sec2_lbl.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = "SECTION 2: DATA GROUNDING ARCHITECTURE"
    r.font.size = Pt(10.0)
    r.font.bold = True
    r.font.color.rgb = NAVY

    # Right Container Card
    right_box = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.78), Inches(1.46), Inches(3.75), Inches(4.36))
    right_box.fill.solid()
    right_box.fill.fore_color.rgb = CARD_BG
    right_box.line.color.rgb = BORDER_LIGHT
    right_box.line.width = Pt(1.0)

    # 2A. Layer 1: Official Repositories
    l1 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.92), Inches(1.58), Inches(3.47), Inches(0.92))
    l1.fill.solid()
    l1.fill.fore_color.rgb = WHITE
    l1.line.color.rgb = BORDER_LIGHT
    l1.line.width = Pt(1.0)

    # Database icon
    icon_db = "SIH 2.0/assets/database.svg.png"
    if os.path.exists(icon_db):
        slide6.shapes.add_picture(icon_db, Inches(9.02), Inches(1.74), Inches(0.42), Inches(0.55))

    tb_l1 = slide6.shapes.add_textbox(Inches(9.52), Inches(1.64), Inches(2.78), Inches(0.80))
    tf_l1 = tb_l1.text_frame
    tf_l1.word_wrap = True
    tf_l1.margin_left = 0
    tf_l1.margin_top = 0
    p = tf_l1.paragraphs[0]
    r = p.add_run()
    r.text = "OFFICIAL GOVERNMENT SOURCES\n"
    r.font.bold = True
    r.font.size = Pt(8.8)
    r.font.color.rgb = NAVY
    r = p.add_run()
    r.text = "NCVET • MSDE • DGT • NCS • Bharat Skills\n"
    r.font.bold = True
    r.font.size = Pt(7.8)
    r.font.color.rgb = DARK_TEAL
    r = p.add_run()
    r.text = "Target sources for verified qualification packs"
    r.font.size = Pt(7.2)
    r.font.color.rgb = GRAY_TEXT

    # Arrow 1
    arr1 = slide6.shapes.add_textbox(Inches(10.45), Inches(2.51), Inches(0.40), Inches(0.24))
    p = arr1.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "↓"
    r.font.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = TEAL

    # 2B. Layer 2: Verified Vocational Information
    l2 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.92), Inches(2.76), Inches(3.47), Inches(0.96))
    l2.fill.solid()
    l2.fill.fore_color.rgb = WHITE
    l2.line.color.rgb = BORDER_LIGHT
    l2.line.width = Pt(1.0)

    # Scale icon
    icon_scale = "SIH 2.0/assets/tradeoff_scale.svg.png"
    if os.path.exists(icon_scale):
        slide6.shapes.add_picture(icon_scale, Inches(9.02), Inches(2.94), Inches(0.44), Inches(0.55))

    tb_l2 = slide6.shapes.add_textbox(Inches(9.52), Inches(2.82), Inches(2.78), Inches(0.84))
    tf_l2 = tb_l2.text_frame
    tf_l2.word_wrap = True
    tf_l2.margin_left = 0
    tf_l2.margin_top = 0
    p = tf_l2.paragraphs[0]
    r = p.add_run()
    r.text = "VERIFIED VOCATIONAL STANDARDS\n"
    r.font.bold = True
    r.font.size = Pt(8.8)
    r.font.color.rgb = NAVY
    r = p.add_run()
    r.text = "• NCrF-Aligned Qualification Packs & Progression\n• CTS ITI Trade Curricula & Practical Hours\n• NCS Regional Wage Baselines & Indicators"
    r.font.size = Pt(7.3)
    r.font.color.rgb = GRAY_TEXT

    # Arrow 2
    arr2 = slide6.shapes.add_textbox(Inches(10.45), Inches(3.73), Inches(0.40), Inches(0.24))
    p = arr2.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "↓"
    r.font.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = TEAL

    # 2C. Layer 3: SkillSaathi Proposed AI Guidance
    l3 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.92), Inches(3.98), Inches(3.47), Inches(1.10))
    l3.fill.solid()
    l3.fill.fore_color.rgb = LIGHT_TEAL
    l3.line.color.rgb = BORDER_TEAL
    l3.line.width = Pt(1.2)

    # AI Brain icon
    icon_brain = "SIH 2.0/assets/ai_brain.svg.png"
    if os.path.exists(icon_brain):
        slide6.shapes.add_picture(icon_brain, Inches(9.02), Inches(4.24), Inches(0.44), Inches(0.44))

    tb_l3 = slide6.shapes.add_textbox(Inches(9.52), Inches(4.04), Inches(2.78), Inches(0.98))
    tf_l3 = tb_l3.text_frame
    tf_l3.word_wrap = True
    tf_l3.margin_left = 0
    tf_l3.margin_top = 0
    p = tf_l3.paragraphs[0]
    r = p.add_run()
    r.text = "SKILLSAATHI PROPOSED AI GUIDANCE\n"
    r.font.bold = True
    r.font.size = Pt(8.8)
    r.font.color.rgb = DARK_TEAL
    r = p.add_run()
    r.text = "• Evidence-grounded guidance using verified frameworks\n• Structured family trade-off & consensus analysis\n• Curated evaluation baseline datasets in prototype"
    r.font.size = Pt(7.3)
    r.font.color.rgb = NAVY

    # Right Box Disclaimer Note
    note_box = slide6.shapes.add_textbox(Inches(8.92), Inches(5.16), Inches(3.47), Inches(0.55))
    tf_n = note_box.text_frame
    tf_n.word_wrap = True
    tf_n.margin_left = 0
    tf_n.margin_top = 0
    p = tf_n.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Note: Target integration sources for production rollout.\nCurrent prototype utilizes structured evaluation datasets."
    r.font.italic = True
    r.font.size = Pt(7.0)
    r.font.color.rgb = SLATE_LIGHT

    # =========================================================================
    # SECTION 3: REFINED BOTTOM STATEMENT (Exact user statement)
    # =========================================================================
    banner = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.80), Inches(6.04), Inches(11.733), Inches(0.50))
    banner.fill.solid()
    banner.fill.fore_color.rgb = LIGHT_TEAL
    banner.line.color.rgb = BORDER_TEAL
    banner.line.width = Pt(1.5)
    
    tf_ban = banner.text_frame
    tf_ban.margin_top = Inches(0.11)
    tf_ban.word_wrap = False
    p = tf_ban.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER

    r = p.add_run()
    r.text = "“ SkillSaathi prioritizes verified government frameworks and evidence-grounded recommendations "
    r.font.name = "Georgia"
    r.font.bold = True
    r.font.size = Pt(9.8)
    r.font.color.rgb = NAVY

    r = p.add_run()
    r.text = "to improve the reliability of vocational career guidance. ”"
    r.font.name = "Georgia"
    r.font.bold = True
    r.font.size = Pt(9.8)
    r.font.color.rgb = DARK_TEAL

    prs.save(ppt_path)
    print("Successfully built refined Slide 6 with enhanced typography and academic rigor in:", ppt_path)

if __name__ == "__main__":
    update_slide6()
