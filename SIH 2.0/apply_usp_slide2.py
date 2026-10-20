import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

pptx_path = "SIH 2.0/SIH2025-IDEA-Presentation-Format.pptx"
prs = Presentation(pptx_path)

# Target Slide 2 ONLY
slide2 = prs.slides[1]

# Palette - Editorial Minimalist
COLOR_NAVY_DARK = RGBColor(0x0B, 0x13, 0x2B)
COLOR_NAVY_TITLE = RGBColor(0x1F, 0x4E, 0x79)
COLOR_SLATE_DARK = RGBColor(0x1E, 0x29, 0x3B)
COLOR_SLATE_MUTED = RGBColor(0x47, 0x55, 0x69)
COLOR_SLATE_LIGHT = RGBColor(0x64, 0x74, 0x8B)

COLOR_TEAL_DARK = RGBColor(0x0D, 0x94, 0x88)
COLOR_TEAL_ACCENT = RGBColor(0x14, 0xB8, 0xA6)
COLOR_TEAL_TINT = RGBColor(0xF0, 0xFD, 0xFA)
COLOR_TEAL_BORDER = RGBColor(0x0D, 0x94, 0x88)
COLOR_TEAL_LIGHT_BORDER = RGBColor(0x99, 0xF6, 0xE4)

COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_OFFWHITE = RGBColor(0xF8, 0xF9, 0xFA)
COLOR_CARD_BORDER = RGBColor(0xCB, 0xD5, 0xE1)
COLOR_SUBTLE_BORDER = RGBColor(0xE2, 0xE8, 0xF0)

# 1. Clean slide2 back to only the 6 official base template shapes
base_names = {'Rectangle 8', 'Title 1', 'Slide Number Placeholder 5', 'Footer Placeholder 6', 'Oval 9', 'Picture 2'}
to_remove = [s for s in slide2.shapes if s.name not in base_names]
for s in to_remove:
    sp = s._element
    sp.getparent().remove(sp)

# 2. Update Oval 9 (Team Badge)
for s in slide2.shapes:
    if s.name == "Oval 9":
        s.left = Inches(0.80)
        s.top = Inches(0.28)
        s.width = Inches(1.65)
        s.height = Inches(0.52)
        s.fill.solid()
        s.fill.fore_color.rgb = COLOR_OFFWHITE
        s.line.color.rgb = COLOR_CARD_BORDER
        s.line.width = Pt(1)
        if s.has_text_frame:
            tf = s.text_frame
            tf.clear()
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            p = tf.paragraphs[0]
            p.text = "TeamChargers"
            p.alignment = PP_ALIGN.CENTER
            p.font.name = "Arial"
            p.font.size = Pt(10.5)
            p.font.bold = True
            p.font.color.rgb = COLOR_NAVY_DARK

# 3. Update Title 1 (Respecting safe buffer before top-right SIH logo at X=10.76")
for s in slide2.shapes:
    if s.name == "Title 1":
        s.left = Inches(2.65)
        s.top = Inches(0.14)
        s.width = Inches(7.70)
        s.height = Inches(0.90)
        if s.has_text_frame:
            tf = s.text_frame
            tf.word_wrap = True
            tf.clear()
            # Mandatory Template Heading: IDEA TITLE
            p1 = tf.paragraphs[0]
            p1.alignment = PP_ALIGN.CENTER
            r1 = p1.add_run()
            r1.text = "IDEA TITLE: "
            r1.font.name = "Garamond"
            r1.font.size = Pt(21)
            r1.font.bold = True
            r1.font.color.rgb = COLOR_NAVY_TITLE
            
            r2 = p1.add_run()
            r2.text = "SkillSaathi"
            r2.font.name = "Arial"
            r2.font.size = Pt(21)
            r2.font.bold = True
            r2.font.color.rgb = COLOR_NAVY_DARK
            p1.space_after = Pt(2)
            
            # Tagline
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.CENTER
            r3 = p2.add_run()
            r3.text = "Beyond Career Advice. Towards Family Consensus."
            r3.font.name = "Arial"
            r3.font.size = Pt(13)
            r3.font.italic = True
            r3.font.color.rgb = COLOR_TEAL_DARK

# ==============================================================================
# ONE-LINE DESCRIPTION (Sub-Header, Single Clean Line)
# ==============================================================================
tb_desc = slide2.shapes.add_textbox(Inches(0.80), Inches(1.10), Inches(9.80), Inches(0.32))
tb_desc.name = "USP_OneLineDesc"
tf_desc = tb_desc.text_frame
tf_desc.word_wrap = True
tf_desc.margin_left = tf_desc.margin_right = tf_desc.margin_top = tf_desc.margin_bottom = 0
p_desc = tf_desc.paragraphs[0]
p_desc.alignment = PP_ALIGN.CENTER
r_desc = p_desc.add_run()
r_desc.text = "An AI-powered family counselling platform that bridges student aspirations and parental concerns through evidence-backed vocational guidance."
r_desc.font.name = "Arial"
r_desc.font.size = Pt(11)
r_desc.font.color.rgb = COLOR_SLATE_DARK

# ==============================================================================
# SECTION A: FAMILY DECISION INTELLIGENCE WORKFLOW (Compact Visual Flow)
# Top: 1.48", Height: 2.15"
# ==============================================================================
wf_top = Inches(1.48)
wf_h = Inches(2.15)

# Outer workflow canvas panel
wf_canvas = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.80), wf_top, Inches(11.733), wf_h)
wf_canvas.name = "USP_WF_Canvas"
wf_canvas.fill.solid()
wf_canvas.fill.fore_color.rgb = COLOR_WHITE
wf_canvas.line.color.rgb = COLOR_SUBTLE_BORDER
wf_canvas.line.width = Pt(1)

# Tag for workflow
tb_wftag = slide2.shapes.add_textbox(Inches(1.00), wf_top + Inches(0.06), Inches(11.333), Inches(0.24))
tb_wftag.name = "USP_WF_Tag"
tf_wft = tb_wftag.text_frame
tf_wft.margin_left = tf_wft.margin_right = tf_wft.margin_top = tf_wft.margin_bottom = 0
pwft = tf_wft.paragraphs[0]
rw1 = pwft.add_run()
rw1.text = "PROPOSED SOLUTION: FAMILY DECISION INTELLIGENCE WORKFLOW"
rw1.font.name = "Arial"
rw1.font.size = Pt(8)
rw1.font.bold = True
rw1.font.color.rgb = COLOR_SLATE_LIGHT

# --- STAGE 1: DUAL INPUTS (Converging) ---
# Student Node
c_student = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.95), wf_top + Inches(0.32), Inches(2.35), Inches(0.80))
c_student.name = "USP_Node_Student"
c_student.fill.solid()
c_student.fill.fore_color.rgb = COLOR_OFFWHITE
c_student.line.color.rgb = COLOR_CARD_BORDER
c_student.line.width = Pt(1)
tfs = c_student.text_frame
tfs.word_wrap = True
tfs.margin_left = tfs.margin_right = Inches(0.10)
tfs.margin_top = Inches(0.08)
tfs.margin_bottom = Inches(0.04)
ps1 = tfs.paragraphs[0]
ps1.text = "Student Aspirations"
ps1.font.name = "Arial"
ps1.font.size = Pt(11)
ps1.font.bold = True
ps1.font.color.rgb = COLOR_NAVY_DARK
ps2 = tfs.add_paragraph()
ps2.text = "Interests, Aptitude & Career Passion"
ps2.font.name = "Arial"
ps2.font.size = Pt(8.5)
ps2.font.color.rgb = COLOR_SLATE_MUTED

# Parent Node
c_parent = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.95), wf_top + Inches(1.20), Inches(2.35), Inches(0.80))
c_parent.name = "USP_Node_Parent"
c_parent.fill.solid()
c_parent.fill.fore_color.rgb = COLOR_OFFWHITE
c_parent.line.color.rgb = COLOR_CARD_BORDER
c_parent.line.width = Pt(1)
tfp = c_parent.text_frame
tfp.word_wrap = True
tfp.margin_left = tfp.margin_right = Inches(0.10)
tfp.margin_top = Inches(0.08)
tfp.margin_bottom = Inches(0.04)
pp1 = tfp.paragraphs[0]
pp1.text = "Parent Concerns"
pp1.font.name = "Arial"
pp1.font.size = Pt(11)
pp1.font.bold = True
pp1.font.color.rgb = COLOR_NAVY_DARK
pp2 = tfp.add_paragraph()
pp2.text = "Financial Security, ROI & Social Stigma"
pp2.font.name = "Arial"
pp2.font.size = Pt(8.5)
pp2.font.color.rgb = COLOR_SLATE_MUTED

# Converging arrows
tb_arr_s = slide2.shapes.add_textbox(Inches(3.35), wf_top + Inches(0.55), Inches(0.35), Inches(0.35))
tfa_s = tb_arr_s.text_frame
tfa_s.margin_left = tfa_s.margin_right = tfa_s.margin_top = tfa_s.margin_bottom = 0
pa_s = tfa_s.paragraphs[0]
pa_s.text = "↘"
pa_s.font.name = "Arial"
pa_s.font.size = Pt(16)
pa_s.font.bold = True
pa_s.font.color.rgb = COLOR_TEAL_ACCENT

tb_arr_p = slide2.shapes.add_textbox(Inches(3.35), wf_top + Inches(1.42), Inches(0.35), Inches(0.35))
tfa_p = tb_arr_p.text_frame
tfa_p.margin_left = tfa_p.margin_right = tfa_p.margin_top = tfa_p.margin_bottom = 0
pa_p = tfa_p.paragraphs[0]
pa_p.text = "↗"
pa_p.font.name = "Arial"
pa_p.font.size = Pt(16)
pa_p.font.bold = True
pa_p.font.color.rgb = COLOR_TEAL_ACCENT

# --- STAGE 2: AI COUNSELLING (Center Convergence Hub) ---
c_counsel = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(3.75), wf_top + Inches(0.32), Inches(2.65), Inches(1.68))
c_counsel.name = "USP_Node_Counsel"
c_counsel.fill.solid()
c_counsel.fill.fore_color.rgb = COLOR_TEAL_TINT
c_counsel.line.color.rgb = COLOR_TEAL_BORDER
c_counsel.line.width = Pt(1.5)
tfc = c_counsel.text_frame
tfc.word_wrap = True
tfc.margin_left = tfc.margin_right = Inches(0.12)
tfc.margin_top = Inches(0.14)
tfc.margin_bottom = Inches(0.08)

pc0 = tfc.paragraphs[0]
pc0.text = "AI COUNSELLING"
pc0.font.name = "Arial"
pc0.font.size = Pt(12)
pc0.font.bold = True
pc0.font.color.rgb = COLOR_NAVY_DARK
pc0.space_after = Pt(4)

pc1 = tfc.add_paragraph()
pc1.text = "Verified Data & Personalized Guidance"
pc1.font.name = "Arial"
pc1.font.size = Pt(9.5)
pc1.font.bold = True
pc1.font.color.rgb = COLOR_TEAL_DARK
pc1.space_after = Pt(4)

pc2 = tfc.add_paragraph()
pc2.text = "• Dual-persona joint counselling dialogue\n• Synthesizes student passion with parent risk factors"
pc2.font.name = "Arial"
pc2.font.size = Pt(8.5)
pc2.font.color.rgb = COLOR_SLATE_DARK

# Arrow 2
tb_arr2 = slide2.shapes.add_textbox(Inches(6.45), wf_top + Inches(1.00), Inches(0.35), Inches(0.35))
tfa2 = tb_arr2.text_frame
tfa2.margin_left = tfa2.margin_right = tfa2.margin_top = tfa2.margin_bottom = 0
pa2 = tfa2.paragraphs[0]
pa2.text = "→"
pa2.alignment = PP_ALIGN.CENTER
pa2.font.name = "Arial"
pa2.font.size = Pt(20)
pa2.font.bold = True
pa2.font.color.rgb = COLOR_TEAL_ACCENT

# --- STAGE 3: EVIDENCE-BASED CAREER COMPARISON ---
c_compare = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.85), wf_top + Inches(0.32), Inches(2.60), Inches(1.68))
c_compare.name = "USP_Node_Compare"
c_compare.fill.solid()
c_compare.fill.fore_color.rgb = COLOR_OFFWHITE
c_compare.line.color.rgb = COLOR_CARD_BORDER
c_compare.line.width = Pt(1)
tfm = c_compare.text_frame
tfm.word_wrap = True
tfm.margin_left = tfm.margin_right = Inches(0.12)
tfm.margin_top = Inches(0.14)
tfm.margin_bottom = Inches(0.08)

pm0 = tfm.paragraphs[0]
pm0.text = "CAREER COMPARISON"
pm0.font.name = "Arial"
pm0.font.size = Pt(12)
pm0.font.bold = True
pm0.font.color.rgb = COLOR_NAVY_DARK
pm0.space_after = Pt(4)

pm1 = tfm.add_paragraph()
pm1.text = "Evidence-Based Benchmarking"
pm1.font.name = "Arial"
pm1.font.size = Pt(9.5)
pm1.font.bold = True
pm1.font.color.rgb = COLOR_SLATE_DARK
pm1.space_after = Pt(4)

pm2 = tfm.add_paragraph()
pm2.text = "• Official NSQF/NCVET accredited pathways\n• Real-world placement & salary comparisons"
pm2.font.name = "Arial"
pm2.font.size = Pt(8.5)
pm2.font.color.rgb = COLOR_SLATE_MUTED

# Arrow 3
tb_arr3 = slide2.shapes.add_textbox(Inches(9.50), wf_top + Inches(1.00), Inches(0.35), Inches(0.35))
tfa3 = tb_arr3.text_frame
tfa3.margin_left = tfa3.margin_right = tfa3.margin_top = tfa3.margin_bottom = 0
pa3 = tfa3.paragraphs[0]
pa3.text = "→"
pa3.alignment = PP_ALIGN.CENTER
pa3.font.name = "Arial"
pa3.font.size = Pt(20)
pa3.font.bold = True
pa3.font.color.rgb = COLOR_TEAL_ACCENT

# --- STAGE 4: FAMILY DECISION ---
c_decision = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.90), wf_top + Inches(0.32), Inches(2.50), Inches(1.68))
c_decision.name = "USP_Node_Decision"
c_decision.fill.solid()
c_decision.fill.fore_color.rgb = COLOR_OFFWHITE
c_decision.line.color.rgb = COLOR_CARD_BORDER
c_decision.line.width = Pt(1)
tfd = c_decision.text_frame
tfd.word_wrap = True
tfd.margin_left = tfd.margin_right = Inches(0.12)
tfd.margin_top = Inches(0.14)
tfd.margin_bottom = Inches(0.08)

pd0 = tfd.paragraphs[0]
pd0.text = "FAMILY DECISION"
pd0.font.name = "Arial"
pd0.font.size = Pt(12)
pd0.font.bold = True
pd0.font.color.rgb = COLOR_NAVY_DARK
pd0.space_after = Pt(4)

pd1 = tfd.add_paragraph()
pd1.text = "Confidence & Career Clarity"
pd1.font.name = "Arial"
pd1.font.size = Pt(9.5)
pd1.font.bold = True
pd1.font.color.rgb = COLOR_NAVY_DARK
pd1.space_after = Pt(4)

pd2 = tfd.add_paragraph()
pd2.text = "• High-clarity, united vocational commitment\n• Measurable confidence assessment"
pd2.font.name = "Arial"
pd2.font.size = Pt(8.5)
pd2.font.color.rgb = COLOR_SLATE_MUTED

# ==============================================================================
# SECTION B: INNOVATION & UNIQUENESS — WHY SKILLSAATHI?
# Asymmetric 45:55 Layout (Top: 3.75", Height: 3.05")
# Left: Primary Hero USP (45%) | Right: 2 Supporting Innovations (55%)
# ==============================================================================
sec_y = Inches(3.75)

# Section Title
tb_sec = slide2.shapes.add_textbox(Inches(0.80), sec_y, Inches(11.733), Inches(0.30))
tb_sec.name = "USP_Sec_Title"
tfs_t = tb_sec.text_frame
tfs_t.margin_left = tfs_t.margin_right = tfs_t.margin_top = tfs_t.margin_bottom = 0
ps_t = tfs_t.paragraphs[0]
rs1 = ps_t.add_run()
rs1.text = "INNOVATION & UNIQUENESS — WHY SKILLSAATHI?"
rs1.font.name = "Arial"
rs1.font.size = Pt(11)
rs1.font.bold = True
rs1.font.color.rgb = COLOR_NAVY_DARK

# Left Panel (45% Width): Hero USP - Family Consensus Intelligence
left_x = Inches(0.80)
left_w = Inches(5.15)
bot_y = Inches(4.08)
bot_h = Inches(2.70)

# Card Container (Shape with no text frame conflict)
hero_card = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, left_x, bot_y, left_w, bot_h)
hero_card.name = "USP_Hero_Card"
hero_card.fill.solid()
hero_card.fill.fore_color.rgb = COLOR_TEAL_TINT
hero_card.line.color.rgb = COLOR_TEAL_BORDER
hero_card.line.width = Pt(1.5)

# Dedicated Text Box for Top Text (Explicit height prevents overlap with flow box below)
tb_hero_text = slide2.shapes.add_textbox(left_x + Inches(0.18), bot_y + Inches(0.10), left_w - Inches(0.36), Inches(1.42))
tb_hero_text.name = "USP_Hero_TextBox"
tfh = tb_hero_text.text_frame
tfh.word_wrap = True
tfh.margin_left = tfh.margin_right = tfh.margin_top = tfh.margin_bottom = 0

# Badge
ph0 = tfh.paragraphs[0]
ph0.text = "PRIMARY USP — CORE INNOVATION"
ph0.font.name = "Arial"
ph0.font.size = Pt(8.5)
ph0.font.bold = True
ph0.font.color.rgb = COLOR_TEAL_DARK
ph0.space_after = Pt(2)

# Title
ph1 = tfh.add_paragraph()
ph1.text = "Family Consensus Intelligence"
ph1.font.name = "Arial"
ph1.font.size = Pt(15)
ph1.font.bold = True
ph1.font.color.rgb = COLOR_NAVY_DARK
ph1.space_after = Pt(4)

# Supporting Quote
ph2 = tfh.add_paragraph()
ph2.text = "\"We don't just recommend careers. We identify what prevents families from agreeing—and help them make informed decisions.\""
ph2.font.name = "Arial"
ph2.font.size = Pt(9.5)
ph2.font.italic = True
ph2.font.color.rgb = COLOR_SLATE_DARK

# Mini Visual Workflow: Before → AI Intervention → After (Positioned cleanly with clear gap)
flow_box = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, left_x + Inches(0.15), bot_y + Inches(1.65), left_w - Inches(0.30), Inches(0.92))
flow_box.name = "USP_Hero_FlowBox"
flow_box.fill.solid()
flow_box.fill.fore_color.rgb = COLOR_WHITE
flow_box.line.color.rgb = COLOR_TEAL_LIGHT_BORDER
flow_box.line.width = Pt(1)

tff = flow_box.text_frame
tff.word_wrap = True
tff.margin_left = tff.margin_right = Inches(0.10)
tff.margin_top = Inches(0.10)
tff.margin_bottom = Inches(0.04)

pf0 = tff.paragraphs[0]
rf_a = pf0.add_run()
rf_a.text = "Before Counselling"
rf_a.font.name = "Arial"
rf_a.font.size = Pt(8)
rf_a.font.bold = True
rf_a.font.color.rgb = COLOR_SLATE_MUTED

rf_arr1 = pf0.add_run()
rf_arr1.text = "   →   "
rf_arr1.font.name = "Arial"
rf_arr1.font.size = Pt(8.5)
rf_arr1.font.bold = True
rf_arr1.font.color.rgb = COLOR_TEAL_ACCENT

rf_b = pf0.add_run()
rf_b.text = "AI Intervention"
rf_b.font.name = "Arial"
rf_b.font.size = Pt(8)
rf_b.font.bold = True
rf_b.font.color.rgb = COLOR_TEAL_DARK

rf_arr2 = pf0.add_run()
rf_arr2.text = "   →   "
rf_arr2.font.name = "Arial"
rf_arr2.font.size = Pt(8.5)
rf_arr2.font.bold = True
rf_arr2.font.color.rgb = COLOR_TEAL_ACCENT

rf_c = pf0.add_run()
rf_c.text = "After Counselling"
rf_c.font.name = "Arial"
rf_c.font.size = Pt(8)
rf_c.font.bold = True
rf_c.font.color.rgb = COLOR_NAVY_DARK
pf0.space_after = Pt(3)

pf1 = tff.add_paragraph()
rf_sub = pf1.add_run()
rf_sub.text = "Parental Concerns   →   Evidence-Based Discussion   →   Confidence Reassessment"
rf_sub.font.name = "Arial"
rf_sub.font.size = Pt(7.5)
rf_sub.font.color.rgb = COLOR_SLATE_MUTED

# Right Panel (55% Width): Two Supporting Innovations
right_x = Inches(6.20)
right_w = Inches(6.33)

# Supporting Innovation 1
sup1 = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, right_x, bot_y, right_w, Inches(1.28))
sup1.name = "USP_Sup1_Card"
sup1.fill.solid()
sup1.fill.fore_color.rgb = COLOR_OFFWHITE
sup1.line.color.rgb = COLOR_CARD_BORDER
sup1.line.width = Pt(1)

tfs1 = sup1.text_frame
tfs1.word_wrap = True
tfs1.margin_left = tfs1.margin_right = Inches(0.18)
tfs1.margin_top = Inches(0.12)
tfs1.margin_bottom = Inches(0.08)

ps1_0 = tfs1.paragraphs[0]
ps1_0.text = "02. Evidence-Backed Decision Engine"
ps1_0.font.name = "Arial"
ps1_0.font.size = Pt(12)
ps1_0.font.bold = True
ps1_0.font.color.rgb = COLOR_NAVY_DARK
ps1_0.space_after = Pt(2)

ps1_1 = tfs1.add_paragraph()
ps1_1.text = "Compares vocational options using verified outcome benchmarks rather than generic AI advice:"
ps1_1.font.name = "Arial"
ps1_1.font.size = Pt(9.5)
ps1_1.font.color.rgb = COLOR_SLATE_MUTED
ps1_1.space_after = Pt(4)

ps1_2 = tfs1.add_paragraph()
r_p1 = ps1_2.add_run()
r_p1.text = "• Real Earnings     • Placement Outcomes     • NSQF Progression Pathways"
r_p1.font.name = "Arial"
r_p1.font.size = Pt(9.5)
r_p1.font.bold = True
r_p1.font.color.rgb = COLOR_TEAL_DARK

# Supporting Innovation 2
sup2 = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, right_x, bot_y + Inches(1.42), right_w, Inches(1.28))
sup2.name = "USP_Sup2_Card"
sup2.fill.solid()
sup2.fill.fore_color.rgb = COLOR_OFFWHITE
sup2.line.color.rgb = COLOR_CARD_BORDER
sup2.line.width = Pt(1)

tfs2 = sup2.text_frame
tfs2.word_wrap = True
tfs2.margin_left = tfs2.margin_right = Inches(0.18)
tfs2.margin_top = Inches(0.12)
tfs2.margin_bottom = Inches(0.08)

ps2_0 = tfs2.paragraphs[0]
ps2_0.text = "03. Resistance-to-Action Intelligence"
ps2_0.font.name = "Arial"
ps2_0.font.size = Pt(12)
ps2_0.font.bold = True
ps2_0.font.color.rgb = COLOR_NAVY_DARK
ps2_0.space_after = Pt(2)

ps2_1 = tfs2.add_paragraph()
ps2_1.text = "Converts unresolved family hesitation into structured escalation and systemic insights:"
ps2_1.font.name = "Arial"
ps2_1.font.size = Pt(9.5)
ps2_1.font.color.rgb = COLOR_SLATE_MUTED
ps2_1.space_after = Pt(4)

ps2_2 = tfs2.add_paragraph()
r_p2 = ps2_2.add_run()
r_p2.text = "• Human Counsellor Escalation     • Administrator Insights & Heatmaps"
r_p2.font.name = "Arial"
r_p2.font.size = Pt(9.5)
r_p2.font.bold = True
r_p2.font.color.rgb = COLOR_NAVY_DARK

# Save presentation
prs.save(pptx_path)
print("Successfully refined Slide 2 with dedicated text box in:", pptx_path)
