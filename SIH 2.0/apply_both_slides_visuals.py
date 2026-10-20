import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def update_all():
    ppt_path = "SIH 2.0/SIH2025-IDEA-Presentation-Format.pptx"
    prs = Presentation(ppt_path)
    
    NAVY = RGBColor(11, 19, 43)          # #0B132B
    TEAL = RGBColor(13, 148, 136)        # #0D9488
    DARK_TEAL = RGBColor(15, 118, 110)   # #0F766E
    LIGHT_TEAL = RGBColor(240, 253, 250) # #F0FDFA
    BORDER_TEAL = RGBColor(45, 212, 191) # #2DD4BF
    WHITE = RGBColor(255, 255, 255)
    GRAY_TEXT = RGBColor(71, 85, 105)    # #475569
    SLATE_LIGHT = RGBColor(100, 116, 139)# #64748B
    CARD_BG = RGBColor(248, 250, 252)    # #F8FAFC
    BORDER_GRAY = RGBColor(203, 213, 225)# #CBD5E1
    BORDER_LIGHT = RGBColor(226, 232, 240)# #E2E8F0
    TEXT_MUTED = RGBColor(100, 116, 139)
    AMBER_BG = RGBColor(254, 243, 199)   # #FEF3C7
    AMBER_TEXT = RGBColor(146, 64, 14)   # #92400E
    BLUE_BG = RGBColor(224, 242, 254)    # #E0F2FE
    BLUE_TEXT = RGBColor(7, 89, 133)     # #075985
    GREEN_BG = RGBColor(220, 252, 231)   # #DCFCE7
    GREEN_TEXT = RGBColor(22, 101, 52)   # #166534
    ROSE_BG = RGBColor(255, 228, 230)    # #FFE4E6
    ROSE_TEXT = RGBColor(159, 18, 57)    # #9F1239

    # =========================================================================
    # REBUILD SLIDE 2 WITH NON-OVERLAPPING ILLUSTRATION BADGES
    # =========================================================================
    slide2 = prs.slides[1]
    preserved_names = {"Rectangle 9", "Slide Number Placeholder 5", "Footer Placeholder 6", "Picture 2"}
    for s in list(slide2.shapes):
        if s.name not in preserved_names:
            sp = s._element
            sp.getparent().remove(sp)

    # 1. Team Name Badge
    team_badge2 = slide2.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.80), Inches(0.35), Inches(1.65), Inches(0.42)
    )
    team_badge2.name = "Team_Badge"
    team_badge2.fill.solid()
    team_badge2.fill.fore_color.rgb = WHITE
    team_badge2.line.color.rgb = BORDER_TEAL
    team_badge2.line.width = Pt(1.5)
    tf_tb = team_badge2.text_frame
    tf_tb.word_wrap = False
    tf_tb.margin_left = Inches(0.05)
    tf_tb.margin_right = Inches(0.05)
    tf_tb.margin_top = Inches(0.06)
    tf_tb.margin_bottom = Inches(0.06)
    p_tb = tf_tb.paragraphs[0]
    p_tb.alignment = PP_ALIGN.CENTER
    r_tb = p_tb.add_run()
    r_tb.text = "TeamChargers"
    r_tb.font.name = "Calibri"
    r_tb.font.size = Pt(12)
    r_tb.font.bold = True
    r_tb.font.color.rgb = NAVY

    # 2. Slide Title
    title_box2 = slide2.shapes.add_textbox(Inches(2.65), Inches(0.28), Inches(7.70), Inches(0.55))
    tf_t2 = title_box2.text_frame
    tf_t2.word_wrap = True
    tf_t2.margin_left = 0
    tf_t2.margin_top = 0
    p_t2 = tf_t2.paragraphs[0]
    p_t2.alignment = PP_ALIGN.LEFT
    r = p_t2.add_run()
    r.text = "IDEA TITLE: "
    r.font.name = "Georgia"
    r.font.size = Pt(21)
    r.font.bold = True
    r.font.color.rgb = NAVY
    r = p_t2.add_run()
    r.text = "SkillSaathi  "
    r.font.name = "Calibri"
    r.font.size = Pt(21)
    r.font.bold = True
    r.font.color.rgb = NAVY
    r = p_t2.add_run()
    r.text = "— Beyond Career Advice. Towards Family Consensus."
    r.font.name = "Calibri"
    r.font.size = Pt(13)
    r.font.italic = True
    r.font.color.rgb = TEAL

    # 3. Sub-header Description
    desc_box2 = slide2.shapes.add_textbox(Inches(0.80), Inches(0.88), Inches(9.60), Inches(0.30))
    tf_d2 = desc_box2.text_frame
    tf_d2.margin_left = 0
    tf_d2.margin_top = 0
    p = tf_d2.paragraphs[0]
    r = p.add_run()
    r.text = "An AI-powered family counselling platform that bridges student aspirations and parental concerns through evidence-backed vocational guidance."
    r.font.name = "Calibri"
    r.font.size = Pt(10.5)
    r.font.color.rgb = GRAY_TEXT

    # 4. Top Workflow Container
    wf_container = slide2.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.80), Inches(1.26), Inches(11.733), Inches(1.88)
    )
    wf_container.fill.solid()
    wf_container.fill.fore_color.rgb = CARD_BG
    wf_container.line.color.rgb = BORDER_LIGHT
    wf_container.line.width = Pt(1.0)

    # Workflow Label (Left)
    wf_lbl = slide2.shapes.add_textbox(Inches(0.95), Inches(1.30), Inches(6.00), Inches(0.24))
    r = wf_lbl.text_frame.paragraphs[0].add_run()
    r.text = "PROPOSED SOLUTION: FAMILY DECISION INTELLIGENCE WORKFLOW"
    r.font.name = "Calibri"
    r.font.size = Pt(8.5)
    r.font.bold = True
    r.font.color.rgb = SLATE_LIGHT

    # Accessibility Indicator (Right)
    acc_box = slide2.shapes.add_textbox(Inches(7.10), Inches(1.30), Inches(5.25), Inches(0.24))
    acc_p = acc_box.text_frame.paragraphs[0]
    acc_p.alignment = PP_ALIGN.RIGHT
    r = acc_p.add_run()
    r.text = "🌐 English + Regional Language  |  Simple, Family-Friendly Guidance"
    r.font.name = "Calibri"
    r.font.size = Pt(8.5)
    r.font.bold = True
    r.font.color.rgb = DARK_TEAL

    # Dual Input Boxes (Student & Parent)
    s_box = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.95), Inches(1.58), Inches(2.15), Inches(0.68))
    s_box.fill.solid()
    s_box.fill.fore_color.rgb = WHITE
    s_box.line.color.rgb = BORDER_GRAY
    tf = s_box.text_frame
    tf.margin_left = Inches(0.10)
    tf.margin_top = Inches(0.08)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Student Aspirations"
    r.font.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = NAVY
    p2 = tf.add_paragraph()
    r = p2.add_run()
    r.text = "Interests, Aptitude & Ambitions"
    r.font.size = Pt(8.5)
    r.font.color.rgb = TEXT_MUTED

    p_box = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.95), Inches(2.36), Inches(2.15), Inches(0.68))
    p_box.fill.solid()
    p_box.fill.fore_color.rgb = WHITE
    p_box.line.color.rgb = BORDER_GRAY
    tf = p_box.text_frame
    tf.margin_left = Inches(0.10)
    tf.margin_top = Inches(0.08)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Parent Concerns"
    r.font.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = NAVY
    p2 = tf.add_paragraph()
    r = p2.add_run()
    r.text = "Financial Security, ROI & Stability"
    r.font.size = Pt(8.5)
    r.font.color.rgb = TEXT_MUTED

    # Converging arrows
    arr1 = slide2.shapes.add_textbox(Inches(3.12), Inches(1.72), Inches(0.28), Inches(0.35))
    r = arr1.text_frame.paragraphs[0].add_run()
    r.text = "↘"
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = TEAL

    arr2 = slide2.shapes.add_textbox(Inches(3.12), Inches(2.50), Inches(0.28), Inches(0.35))
    r = arr2.text_frame.paragraphs[0].add_run()
    r.text = "↗"
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = TEAL

    # AI Counselling Box
    ai_box = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.42), Inches(1.58), Inches(2.55), Inches(1.46))
    ai_box.fill.solid()
    ai_box.fill.fore_color.rgb = WHITE
    ai_box.line.color.rgb = TEAL
    ai_box.line.width = Pt(1.75)
    tf = ai_box.text_frame
    tf.margin_left = Inches(0.12)
    tf.margin_top = Inches(0.10)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "AI FAMILY COUNSELLING"
    r.font.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = NAVY
    p2 = tf.add_paragraph()
    r = p2.add_run()
    r.text = "Dual-Persona Dialogue Mediation"
    r.font.bold = True
    r.font.size = Pt(9.0)
    r.font.color.rgb = TEAL
    p3 = tf.add_paragraph()
    r = p3.add_run()
    r.text = "• Identifies hidden parental objections\n• Synthesizes passion with practical ROI"
    r.font.size = Pt(8.5)
    r.font.color.rgb = GRAY_TEXT

    arr3 = slide2.shapes.add_textbox(Inches(6.02), Inches(2.15), Inches(0.28), Inches(0.35))
    r = arr3.text_frame.paragraphs[0].add_run()
    r.text = "→"
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = TEAL

    # Career Comparison Box
    cc_box = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.35), Inches(1.58), Inches(2.65), Inches(1.46))
    cc_box.fill.solid()
    cc_box.fill.fore_color.rgb = WHITE
    cc_box.line.color.rgb = BORDER_GRAY
    tf = cc_box.text_frame
    tf.margin_left = Inches(0.12)
    tf.margin_top = Inches(0.10)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "CAREER COMPARISON"
    r.font.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = NAVY
    p2 = tf.add_paragraph()
    r = p2.add_run()
    r.text = "Evidence-Based Benchmarking"
    r.font.bold = True
    r.font.size = Pt(9.0)
    r.font.color.rgb = DARK_TEAL
    p3 = tf.add_paragraph()
    r = p3.add_run()
    r.text = "• Official NSQF/NCVET accredited paths\n• Real-world placement & salary data"
    r.font.size = Pt(8.5)
    r.font.color.rgb = GRAY_TEXT

    arr4 = slide2.shapes.add_textbox(Inches(9.05), Inches(2.15), Inches(0.28), Inches(0.35))
    r = arr4.text_frame.paragraphs[0].add_run()
    r.text = "→"
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = TEAL

    # Informed Decision Box
    fd_box = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.38), Inches(1.58), Inches(2.95), Inches(1.46))
    fd_box.fill.solid()
    fd_box.fill.fore_color.rgb = WHITE
    fd_box.line.color.rgb = BORDER_GRAY
    tf = fd_box.text_frame
    tf.margin_left = Inches(0.12)
    tf.margin_top = Inches(0.10)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "INFORMED FAMILY DECISION"
    r.font.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = NAVY
    p2 = tf.add_paragraph()
    r = p2.add_run()
    r.text = "Consensus & Career Clarity"
    r.font.bold = True
    r.font.size = Pt(9.0)
    r.font.color.rgb = DARK_TEAL
    p3 = tf.add_paragraph()
    r = p3.add_run()
    r.text = "• Aligned student-parent commitment\n• Objective confidence change measured"
    r.font.size = Pt(8.5)
    r.font.color.rgb = GRAY_TEXT

    # 5. Bottom 2x2 Innovation Grid
    grid_lbl = slide2.shapes.add_textbox(Inches(0.80), Inches(3.22), Inches(11.733), Inches(0.28))
    r = grid_lbl.text_frame.paragraphs[0].add_run()
    r.text = "INNOVATION & UNIQUENESS — WHY SKILLSAATHI?"
    r.font.name = "Calibri"
    r.font.size = Pt(11.5)
    r.font.bold = True
    r.font.color.rgb = NAVY

    col_w = Inches(5.72)
    card_h = Inches(1.64)
    c1_left = Inches(0.80)
    c2_left = Inches(6.81)
    r1_top = Inches(3.52)
    r2_top = Inches(5.26)

    diag_h = Inches(0.70)
    diag_w = col_w - Inches(0.24)
    row1_diag_y = r1_top + Inches(0.82)
    row1_inner_y = row1_diag_y + Inches(0.08)
    row2_diag_y = r2_top + Inches(0.82)
    row2_inner_y = row2_diag_y + Inches(0.08)
    chip_h = Inches(0.54)

    # CARD 01: FAMILY CONSENSUS INTELLIGENCE
    c1 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c1_left, r1_top, col_w, card_h)
    c1.fill.solid()
    c1.fill.fore_color.rgb = LIGHT_TEAL
    c1.line.color.rgb = TEAL
    c1.line.width = Pt(1.75)

    # Text Box width constrained to 4.75" so it never touches icon!
    c1_tb = slide2.shapes.add_textbox(c1_left + Inches(0.12), r1_top + Inches(0.07), Inches(4.75), Inches(0.70))
    tf1 = c1_tb.text_frame
    tf1.word_wrap = True
    tf1.margin_left = 0
    tf1.margin_top = 0
    p = tf1.paragraphs[0]
    r = p.add_run()
    r.text = "01. Family Consensus Intelligence "
    r.font.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = NAVY
    r = p.add_run()
    r.text = "★ FLAGSHIP USP"
    r.font.bold = True
    r.font.size = Pt(9.0)
    r.font.color.rgb = TEAL
    p2 = tf1.add_paragraph()
    r = p2.add_run()
    r.text = "Identifies student-parent disagreements, addresses parental objections, and tracks changes in family confidence."
    r.font.size = Pt(8.5)
    r.font.color.rgb = GRAY_TEXT

    # Icon 1 (Transparent PNG badge, no overlap)
    pic1 = slide2.shapes.add_picture(
        "SIH 2.0/assets/family_convergence.svg.png",
        c1_left + Inches(4.96), r1_top + Inches(0.10), Inches(0.62), Inches(0.50)
    )

    diag1_bg = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c1_left + Inches(0.12), row1_diag_y, diag_w, diag_h)
    diag1_bg.fill.solid()
    diag1_bg.fill.fore_color.rgb = WHITE
    diag1_bg.line.color.rgb = BORDER_TEAL

    ch1 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c1_left + Inches(0.20), row1_inner_y, Inches(1.50), chip_h)
    ch1.fill.solid()
    ch1.fill.fore_color.rgb = AMBER_BG
    ch1.line.color.rgb = AMBER_BG
    tf = ch1.text_frame
    tf.margin_top = Inches(0.06)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Disagreement Detection\n"
    r.font.size = Pt(7.8)
    r.font.bold = True
    r.font.color.rgb = AMBER_TEXT
    r = p.add_run()
    r.text = "Student vs Parent View"
    r.font.size = Pt(7.0)
    r.font.color.rgb = AMBER_TEXT

    arr = slide2.shapes.add_textbox(c1_left + Inches(1.72), row1_inner_y + Inches(0.08), Inches(0.24), Inches(0.30))
    r = arr.text_frame.paragraphs[0].add_run()
    r.text = "→"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = TEAL

    ch2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c1_left + Inches(1.98), row1_inner_y, Inches(1.68), chip_h)
    ch2.fill.solid()
    ch2.fill.fore_color.rgb = LIGHT_TEAL
    ch2.line.color.rgb = BORDER_TEAL
    tf = ch2.text_frame
    tf.margin_top = Inches(0.06)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "AI Evidence Mediation\n"
    r.font.size = Pt(7.8)
    r.font.bold = True
    r.font.color.rgb = DARK_TEAL
    r = p.add_run()
    r.text = "Targeted Objection Fact-Check"
    r.font.size = Pt(7.0)
    r.font.color.rgb = DARK_TEAL

    arr = slide2.shapes.add_textbox(c1_left + Inches(3.68), row1_inner_y + Inches(0.08), Inches(0.24), Inches(0.30))
    r = arr.text_frame.paragraphs[0].add_run()
    r.text = "→"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = TEAL

    ch3 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c1_left + Inches(3.94), row1_inner_y, Inches(1.66), chip_h)
    ch3.fill.solid()
    ch3.fill.fore_color.rgb = GREEN_BG
    ch3.line.color.rgb = GREEN_BG
    tf = ch3.text_frame
    tf.margin_top = Inches(0.06)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Consensus & Agreement\n"
    r.font.size = Pt(7.8)
    r.font.bold = True
    r.font.color.rgb = GREEN_TEXT
    r = p.add_run()
    r.text = "Confidence Change Measured"
    r.font.size = Pt(7.0)
    r.font.color.rgb = GREEN_TEXT

    # CARD 02: HYPERLOCAL CAREER REALITY ENGINE
    c2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2_left, r1_top, col_w, card_h)
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = BORDER_GRAY

    c2_tb = slide2.shapes.add_textbox(c2_left + Inches(0.12), r1_top + Inches(0.07), Inches(4.75), Inches(0.70))
    tf2 = c2_tb.text_frame
    tf2.word_wrap = True
    tf2.margin_left = 0
    tf2.margin_top = 0
    p = tf2.paragraphs[0]
    r = p.add_run()
    r.text = "02. Hyperlocal Career Reality Engine"
    r.font.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = NAVY
    p2 = tf2.add_paragraph()
    r = p2.add_run()
    r.text = "Connects vocational choices to available regional earnings, placement outcomes, and accredited training pathways."
    r.font.size = Pt(8.5)
    r.font.color.rgb = GRAY_TEXT

    # Icon 2: India Map Outline
    pic2 = slide2.shapes.add_picture(
        "SIH 2.0/assets/india_map.svg.png",
        c2_left + Inches(5.02), r1_top + Inches(0.08), Inches(0.55), Inches(0.60)
    )

    diag2_bg = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2_left + Inches(0.12), row1_diag_y, diag_w, diag_h)
    diag2_bg.fill.solid()
    diag2_bg.fill.fore_color.rgb = WHITE
    diag2_bg.line.color.rgb = BORDER_LIGHT

    loc_node = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2_left + Inches(0.20), row1_inner_y, Inches(1.30), chip_h)
    loc_node.fill.solid()
    loc_node.fill.fore_color.rgb = BLUE_BG
    loc_node.line.color.rgb = BLUE_BG
    tf = loc_node.text_frame
    tf.margin_top = Inches(0.06)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "📍 District Cluster\n"
    r.font.size = Pt(7.8)
    r.font.bold = True
    r.font.color.rgb = BLUE_TEXT
    r = p.add_run()
    r.text = "Local ITI & Hubs"
    r.font.size = Pt(7.0)
    r.font.color.rgb = BLUE_TEXT

    arr = slide2.shapes.add_textbox(c2_left + Inches(1.52), row1_inner_y + Inches(0.08), Inches(0.22), Inches(0.30))
    r = arr.text_frame.paragraphs[0].add_run()
    r.text = "→"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = TEAL

    met1 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2_left + Inches(1.78), row1_inner_y, Inches(1.18), chip_h)
    met1.fill.solid()
    met1.fill.fore_color.rgb = CARD_BG
    met1.line.color.rgb = BORDER_GRAY
    tf = met1.text_frame
    tf.margin_top = Inches(0.06)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "₹ Regional Wages\n"
    r.font.size = Pt(7.8)
    r.font.bold = True
    r.font.color.rgb = NAVY
    r = p.add_run()
    r.text = "Trade Wage Data"
    r.font.size = Pt(6.8)
    r.font.color.rgb = TEXT_MUTED

    met2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2_left + Inches(3.04), row1_inner_y, Inches(1.20), chip_h)
    met2.fill.solid()
    met2.fill.fore_color.rgb = CARD_BG
    met2.line.color.rgb = BORDER_GRAY
    tf = met2.text_frame
    tf.margin_top = Inches(0.06)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Placement Trends\n"
    r.font.size = Pt(7.8)
    r.font.bold = True
    r.font.color.rgb = DARK_TEAL
    r = p.add_run()
    r.text = "Verified Datasets"
    r.font.size = Pt(6.8)
    r.font.color.rgb = TEXT_MUTED

    met3 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2_left + Inches(4.32), row1_inner_y, Inches(1.28), chip_h)
    met3.fill.solid()
    met3.fill.fore_color.rgb = CARD_BG
    met3.line.color.rgb = BORDER_GRAY
    tf = met3.text_frame
    tf.margin_top = Inches(0.06)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "NSQF Pathway\n"
    r.font.size = Pt(7.8)
    r.font.bold = True
    r.font.color.rgb = NAVY
    r = p.add_run()
    r.text = "ITI → Poly → B.Voc"
    r.font.size = Pt(6.8)
    r.font.color.rgb = TEXT_MUTED

    # CARD 03: EXPLAINABLE CAREER TRADE-OFF ENGINE
    c3 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c1_left, r2_top, col_w, card_h)
    c3.fill.solid()
    c3.fill.fore_color.rgb = CARD_BG
    c3.line.color.rgb = BORDER_GRAY

    c3_tb = slide2.shapes.add_textbox(c1_left + Inches(0.12), r2_top + Inches(0.07), Inches(4.75), Inches(0.70))
    tf3 = c3_tb.text_frame
    tf3.word_wrap = True
    tf3.margin_left = 0
    tf3.margin_top = 0
    p = tf3.paragraphs[0]
    r = p.add_run()
    r.text = "03. Explainable Career Trade-off Engine"
    r.font.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = NAVY
    p2 = tf3.add_paragraph()
    r = p2.add_run()
    r.text = "Compares career options based on family priorities, training costs, earnings, and NSQF progression."
    r.font.size = Pt(8.5)
    r.font.color.rgb = GRAY_TEXT

    # Icon 3: Trade-off Scale
    pic3 = slide2.shapes.add_picture(
        "SIH 2.0/assets/tradeoff_scale.svg.png",
        c1_left + Inches(4.98), r2_top + Inches(0.10), Inches(0.60), Inches(0.50)
    )

    diag3_bg = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c1_left + Inches(0.12), row2_diag_y, diag_w, diag_h)
    diag3_bg.fill.solid()
    diag3_bg.fill.fore_color.rgb = WHITE
    diag3_bg.line.color.rgb = BORDER_LIGHT

    p_w = Inches(1.26)
    p1 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c1_left + Inches(0.20), row2_inner_y, p_w, chip_h)
    p1.fill.solid()
    p1.fill.fore_color.rgb = CARD_BG
    p1.line.color.rgb = BORDER_GRAY
    tf = p1.text_frame
    tf.margin_top = Inches(0.06)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Family Priorities\n"
    r.font.size = Pt(7.8)
    r.font.bold = True
    r.font.color.rgb = NAVY
    r = p.add_run()
    r.text = "Budget & Values"
    r.font.size = Pt(6.8)
    r.font.color.rgb = TEXT_MUTED

    p2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c1_left + Inches(1.54), row2_inner_y, p_w, chip_h)
    p2.fill.solid()
    p2.fill.fore_color.rgb = CARD_BG
    p2.line.color.rgb = BORDER_GRAY
    tf = p2.text_frame
    tf.margin_top = Inches(0.06)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Training Cost\n"
    r.font.size = Pt(7.8)
    r.font.bold = True
    r.font.color.rgb = DARK_TEAL
    r = p.add_run()
    r.text = "Fees vs Subsidy"
    r.font.size = Pt(6.8)
    r.font.color.rgb = TEXT_MUTED

    p3 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c1_left + Inches(2.88), row2_inner_y, p_w, chip_h)
    p3.fill.solid()
    p3.fill.fore_color.rgb = CARD_BG
    p3.line.color.rgb = BORDER_GRAY
    tf = p3.text_frame
    tf.margin_top = Inches(0.06)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Time-to-Earn\n"
    r.font.size = Pt(7.8)
    r.font.bold = True
    r.font.color.rgb = NAVY
    r = p.add_run()
    r.text = "1–2 Yrs vs 4+ Yrs"
    r.font.size = Pt(6.8)
    r.font.color.rgb = TEXT_MUTED

    p4 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c1_left + Inches(4.22), row2_inner_y, Inches(1.38), chip_h)
    p4.fill.solid()
    p4.fill.fore_color.rgb = CARD_BG
    p4.line.color.rgb = BORDER_GRAY
    tf = p4.text_frame
    tf.margin_top = Inches(0.06)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "NSQF Progression\n"
    r.font.size = Pt(7.8)
    r.font.bold = True
    r.font.color.rgb = DARK_TEAL
    r = p.add_run()
    r.text = "Degree Equivalency"
    r.font.size = Pt(6.8)
    r.font.color.rgb = TEXT_MUTED

    # CARD 04: PARENTAL RESISTANCE INTELLIGENCE
    c4 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2_left, r2_top, col_w, card_h)
    c4.fill.solid()
    c4.fill.fore_color.rgb = CARD_BG
    c4.line.color.rgb = BORDER_GRAY

    c4_tb = slide2.shapes.add_textbox(c2_left + Inches(0.12), r2_top + Inches(0.07), Inches(4.75), Inches(0.70))
    tf4 = c4_tb.text_frame
    tf4.word_wrap = True
    tf4.margin_left = 0
    tf4.margin_top = 0
    p = tf4.paragraphs[0]
    r = p.add_run()
    r.text = "04. Parental Resistance Intelligence"
    r.font.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = NAVY
    p2 = tf4.add_paragraph()
    r = p2.add_run()
    r.text = "Transforms counselling concerns into aggregated regional insights and escalation pathways for scheme administrators."
    r.font.size = Pt(8.5)
    r.font.color.rgb = GRAY_TEXT

    # Icon 4: Regional Heatmap
    pic4 = slide2.shapes.add_picture(
        "SIH 2.0/assets/resistance_heatmap.svg.png",
        c2_left + Inches(5.04), r2_top + Inches(0.10), Inches(0.55), Inches(0.52)
    )

    diag4_bg = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2_left + Inches(0.12), row2_diag_y, diag_w, diag_h)
    diag4_bg.fill.solid()
    diag4_bg.fill.fore_color.rgb = WHITE
    diag4_bg.line.color.rgb = BORDER_LIGHT

    in_res = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2_left + Inches(0.20), row2_inner_y, Inches(1.60), chip_h)
    in_res.fill.solid()
    in_res.fill.fore_color.rgb = ROSE_BG
    in_res.line.color.rgb = ROSE_BG
    tf = in_res.text_frame
    tf.margin_top = Inches(0.06)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Unresolved Hesitation\n"
    r.font.size = Pt(7.8)
    r.font.bold = True
    r.font.color.rgb = ROSE_TEXT
    r = p.add_run()
    r.text = "Stigma / Financial Anxiety"
    r.font.size = Pt(6.8)
    r.font.color.rgb = ROSE_TEXT

    arr = slide2.shapes.add_textbox(c2_left + Inches(1.84), row2_inner_y + Inches(0.08), Inches(0.22), Inches(0.30))
    r = arr.text_frame.paragraphs[0].add_run()
    r.text = "→"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = TEAL

    out_a = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2_left + Inches(2.10), row2_inner_y, Inches(1.68), chip_h)
    out_a.fill.solid()
    out_a.fill.fore_color.rgb = BLUE_BG
    out_a.line.color.rgb = BLUE_BG
    tf = out_a.text_frame
    tf.margin_top = Inches(0.06)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Human Counsellor Routing\n"
    r.font.size = Pt(7.8)
    r.font.bold = True
    r.font.color.rgb = BLUE_TEXT
    r = p.add_run()
    r.text = "1-on-1 Escalation Hand-off"
    r.font.size = Pt(6.8)
    r.font.color.rgb = BLUE_TEXT

    arr = slide2.shapes.add_textbox(c2_left + Inches(3.82), row2_inner_y + Inches(0.08), Inches(0.18), Inches(0.30))
    r = arr.text_frame.paragraphs[0].add_run()
    r.text = "+"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = TEAL

    out_b = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2_left + Inches(4.04), row2_inner_y, Inches(1.56), chip_h)
    out_b.fill.solid()
    out_b.fill.fore_color.rgb = LIGHT_TEAL
    out_b.line.color.rgb = BORDER_TEAL
    tf = out_b.text_frame
    tf.margin_top = Inches(0.06)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Admin Heatmaps\n"
    r.font.size = Pt(7.8)
    r.font.bold = True
    r.font.color.rgb = DARK_TEAL
    r = p.add_run()
    r.text = "District Scheme Insights"
    r.font.size = Pt(6.8)
    r.font.color.rgb = DARK_TEAL

    # =========================================================================
    # REBUILD SLIDE 3 WITH AUTHENTIC LOGOS & CLEAN SPACING
    # =========================================================================
    slide3 = prs.slides[2]
    for s in list(slide3.shapes):
        if s.name not in preserved_names:
            sp = s._element
            sp.getparent().remove(sp)

    # 1. Team Name Badge
    team_badge3 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.80), Inches(0.32), Inches(1.65), Inches(0.42))
    team_badge3.fill.solid()
    team_badge3.fill.fore_color.rgb = WHITE
    team_badge3.line.color.rgb = BORDER_TEAL
    team_badge3.line.width = Pt(1.5)
    tf = team_badge3.text_frame
    tf.word_wrap = False
    tf.margin_top = Inches(0.06)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "TeamChargers"
    r.font.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = NAVY

    # 2. Slide Title
    title_box3 = slide3.shapes.add_textbox(Inches(2.65), Inches(0.28), Inches(8.00), Inches(0.45))
    tf = title_box3.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "TECHNICAL APPROACH: "
    r.font.name = "Georgia"
    r.font.size = Pt(20)
    r.font.bold = True
    r.font.color.rgb = NAVY
    r = p.add_run()
    r.text = "System Architecture & Implementation"
    r.font.name = "Calibri"
    r.font.size = Pt(18)
    r.font.bold = True
    r.font.color.rgb = TEAL

    # 3. Sub-header Description
    desc_box3 = slide3.shapes.add_textbox(Inches(0.80), Inches(0.82), Inches(9.80), Inches(0.28))
    tf = desc_box3.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "An end-to-end technical pipeline translating joint family dialogue into grounded, evidence-backed vocational pathways."
    r.font.size = Pt(10.5)
    r.font.color.rgb = GRAY_TEXT

    # 4. Main 4-Stage Architecture
    card_w = Inches(2.62)
    stage_card_h = Inches(2.55)
    stage_top = Inches(1.24)

    stages = [
        {
            "num": "STAGE 01",
            "title": "User Interface",
            "status": "React Frontend [Active]",
            "status_color": DARK_TEAL,
            "items": ["React + TypeScript", "Student & Parent Portals", "Multilingual Interaction"],
            "highlight": False,
            "left": Inches(0.80)
        },
        {
            "num": "STAGE 02",
            "title": "Backend & Data",
            "status": "FastAPI Services [Planned]",
            "status_color": SLATE_LIGHT,
            "items": ["FastAPI Architecture", "Family Profiles & Sessions", "Vocational Outcome Datasets"],
            "highlight": False,
            "left": Inches(3.84)
        },
        {
            "num": "STAGE 03",
            "title": "AI Decision Engine",
            "status": "Core Innovation [Proposed]",
            "status_color": TEAL,
            "items": ["Evidence-Grounded Counselling", "Disagreement Detection", "Career Comparison & Confidence Tracking"],
            "highlight": True,
            "left": Inches(6.88)
        },
        {
            "num": "STAGE 04",
            "title": "Outcomes & Analytics",
            "status": "Action & Governance [Planned]",
            "status_color": SLATE_LIGHT,
            "items": ["Family Decision Summary", "Human Counsellor Escalation", "Administrator Dashboard"],
            "highlight": False,
            "left": Inches(9.92)
        }
    ]

    for stage in stages:
        c_left = stage["left"]
        box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_left, stage_top, card_w, stage_card_h)
        box.fill.solid()
        if stage["highlight"]:
            box.fill.fore_color.rgb = LIGHT_TEAL
            box.line.color.rgb = TEAL
            box.line.width = Pt(2.0)
        else:
            box.fill.fore_color.rgb = WHITE
            box.line.color.rgb = BORDER_GRAY
            box.line.width = Pt(1.25)

        tf = box.text_frame
        tf.margin_left = Inches(0.14)
        tf.margin_right = Inches(0.14)
        tf.margin_top = Inches(0.15)
        tf.margin_bottom = Inches(0.12)
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        r = p0.add_run()
        r.text = stage["num"] + "  •  "
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = stage["status_color"]
        r = p0.add_run()
        r.text = stage["status"]
        r.font.size = Pt(8.0)
        r.font.bold = True
        r.font.color.rgb = stage["status_color"]

        p_title = tf.add_paragraph()
        r = p_title.add_run()
        r.text = stage["title"]
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = NAVY

        p_div = tf.add_paragraph()
        p_div.add_run().font.size = Pt(4)

        for item in stage["items"]:
            p_item = tf.add_paragraph()
            p_item.space_before = Pt(4)
            r = p_item.add_run()
            r.text = "•  "
            r.font.size = Pt(10)
            r.font.bold = True
            r.font.color.rgb = TEAL if stage["highlight"] else DARK_TEAL
            r = p_item.add_run()
            r.text = item
            r.font.size = Pt(10)
            r.font.bold = (stage["highlight"] and "Evidence" in item)
            r.font.color.rgb = NAVY if stage["highlight"] else GRAY_TEXT

    # Connecting Arrows
    for ax in [Inches(3.46), Inches(6.50), Inches(9.54)]:
        arr = slide3.shapes.add_textbox(ax, Inches(2.25), Inches(0.34), Inches(0.40))
        p = arr.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = "➔"
        r.font.size = Pt(18)
        r.font.bold = True
        r.font.color.rgb = TEAL

    # Stage 1 Tech Logos: React & TypeScript (Top-Right of Stage 1 Card)
    slide3.shapes.add_picture("SIH 2.0/assets/react.svg.png", Inches(0.80 + 2.62 - 0.78), Inches(1.24 + 0.14), Inches(0.32), Inches(0.32))
    slide3.shapes.add_picture("SIH 2.0/assets/typescript.svg.png", Inches(0.80 + 2.62 - 0.42), Inches(1.24 + 0.14), Inches(0.30), Inches(0.30))

    # Stage 2 Tech Logos: Python & FastAPI (Top-Right of Stage 2 Card)
    slide3.shapes.add_picture("SIH 2.0/assets/python.svg.png", Inches(3.84 + 2.62 - 0.78), Inches(1.24 + 0.14), Inches(0.30), Inches(0.30))
    slide3.shapes.add_picture("SIH 2.0/assets/fastapi.svg.png", Inches(3.84 + 2.62 - 0.42), Inches(1.24 + 0.14), Inches(0.30), Inches(0.30))

    # Stage 3 AI Brain Icon (Top-Right of Stage 3 Card)
    slide3.shapes.add_picture("SIH 2.0/assets/ai_brain.svg.png", Inches(6.88 + 2.62 - 0.50), Inches(1.24 + 0.14), Inches(0.36), Inches(0.36))

    # Stage 4 Analytics Icon (Top-Right of Stage 4 Card)
    slide3.shapes.add_picture("SIH 2.0/assets/analytics.svg.png", Inches(9.92 + 2.62 - 0.50), Inches(1.24 + 0.14), Inches(0.36), Inches(0.36))

    # Data Grounding Engine Box Underneath Stage 3
    db_left = Inches(6.30)
    db_w = Inches(3.78)
    db_top = Inches(4.04)
    db_h = Inches(0.68)

    arr_up = slide3.shapes.add_textbox(Inches(8.05), Inches(3.78), Inches(0.30), Inches(0.28))
    p = arr_up.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "⬆"
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = TEAL

    db_box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, db_left, db_top, db_w, db_h)
    db_box.fill.solid()
    db_box.fill.fore_color.rgb = WHITE
    db_box.line.color.rgb = TEAL
    db_box.line.width = Pt(1.5)

    # Database Icon on left of the box
    slide3.shapes.add_picture("SIH 2.0/assets/database.svg.png", db_left + Inches(0.14), db_top + Inches(0.14), Inches(0.32), Inches(0.40))

    tf = db_box.text_frame
    tf.margin_left = Inches(0.52) # Padded so text never clips icon!
    tf.margin_top = Inches(0.06)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "DATA GROUNDING ENGINE\n"
    r.font.bold = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = DARK_TEAL
    p2 = tf.add_paragraph()
    r = p2.add_run()
    r.text = "Vocational Data: Earnings • Placements • NSQF Pathways"
    r.font.bold = True
    r.font.size = Pt(9.0)
    r.font.color.rgb = NAVY

    flow_lbl = slide3.shapes.add_textbox(Inches(3.84), Inches(4.18), Inches(2.35), Inches(0.40))
    p = flow_lbl.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run()
    r.text = "Official Benchmarks  ➔\n(NCVET & NCS Data)"
    r.font.size = Pt(8.0)
    r.font.italic = True
    r.font.color.rgb = SLATE_LIGHT

    # 5. Technology Stack Container (Bottom Strip)
    tech_container = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.80), Inches(4.96), Inches(11.733), Inches(1.82))
    tech_container.fill.solid()
    tech_container.fill.fore_color.rgb = CARD_BG
    tech_container.line.color.rgb = BORDER_LIGHT

    t_lbl = slide3.shapes.add_textbox(Inches(0.95), Inches(5.02), Inches(8.00), Inches(0.24))
    r = t_lbl.text_frame.paragraphs[0].add_run()
    r.text = "TECHNOLOGY STACK & IMPLEMENTATION SCOPE"
    r.font.bold = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = SLATE_LIGHT

    t_col_w = Inches(2.78)
    t_col_gap = Inches(0.15)
    t_base_x = Inches(0.92)
    t_col_top = Inches(5.30)
    t_col_h = Inches(1.36)

    tech_categories = [
        {
            "category": "FRONTEND", "tag": "[ACTIVE IN CODE]", "tag_color": DARK_TEAL,
            "headline": "React, TypeScript, Vite",
            "desc": "Modular component library, strict TypeScript contracts, responsive family UI & active mock services."
        },
        {
            "category": "BACKEND", "tag": "[PLANNED ARCHITECTURE]", "tag_color": SLATE_LIGHT,
            "headline": "FastAPI (Planned)",
            "desc": "Pre-architected REST API client (client.ts), typed Pydantic envelopes, family session & profile data models."
        },
        {
            "category": "AI & DATA", "tag": "[PROPOSED LOGIC]", "tag_color": TEAL,
            "headline": "Grounded AI + Vocational Data",
            "desc": "Evidence-grounded decision engine bound to official NCVET qualification packs and NCS wage benchmarks."
        },
        {
            "category": "OUTPUTS", "tag": "[PLANNED SYSTEM]", "tag_color": SLATE_LIGHT,
            "headline": "Counselling, Referrals, Analytics",
            "desc": "Family decision summaries, automated human counsellor escalation routing, and administrative district analytics."
        }
    ]

    for i, cat in enumerate(tech_categories):
        cx = t_base_x + i * (t_col_w + t_col_gap)
        cbox = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, t_col_top, t_col_w, t_col_h)
        cbox.fill.solid()
        cbox.fill.fore_color.rgb = WHITE
        cbox.line.color.rgb = BORDER_GRAY

        tf = cbox.text_frame
        tf.margin_left = Inches(0.10)
        tf.margin_right = Inches(0.10)
        tf.margin_top = Inches(0.08)
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        r = p1.add_run()
        r.text = cat["category"] + " "
        r.font.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = NAVY
        r = p1.add_run()
        r.text = cat["tag"]
        r.font.bold = True
        r.font.size = Pt(7.5)
        r.font.color.rgb = cat["tag_color"]

        p2 = tf.add_paragraph()
        r = p2.add_run()
        r.text = cat["headline"]
        r.font.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = DARK_TEAL if "AI" in cat["category"] else NAVY

        p3 = tf.add_paragraph()
        r = p3.add_run()
        r.text = cat["desc"]
        r.font.size = Pt(7.8)
        r.font.color.rgb = GRAY_TEXT

    # Vite Logo in Column 1 (Frontend)
    slide3.shapes.add_picture("SIH 2.0/assets/vite.svg.png", Inches(0.92 + 2.78 - 0.40), Inches(5.38), Inches(0.28), Inches(0.28))

    prs.save(ppt_path)
    print("Successfully rebuilt Slides 2 & 3 with transparent visuals & logos in:", ppt_path)

if __name__ == "__main__":
    update_all()
