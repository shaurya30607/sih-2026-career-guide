import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def apply_slide_3():
    ppt_path = "SIH 2.0/SIH2025-IDEA-Presentation-Format.pptx"
    prs = Presentation(ppt_path)
    slide3 = prs.slides[2] # 0-indexed 2 is Slide 3
    
    # Preserve template essentials
    preserved_names = {"Rectangle 9", "Slide Number Placeholder 5", "Footer Placeholder 6", "Picture 2"}
    
    shapes_to_remove = [s for s in slide3.shapes if s.name not in preserved_names]
    for s in shapes_to_remove:
        sp = s._element
        sp.getparent().remove(sp)
        
    # Color palette
    NAVY = RGBColor(11, 19, 43)          # #0B132B
    TEAL = RGBColor(13, 148, 136)        # #0D9488
    DARK_TEAL = RGBColor(15, 118, 110)   # #0F766E
    LIGHT_TEAL = RGBColor(240, 253, 250) # #F0FDFA
    BORDER_TEAL = RGBColor(45, 212, 191) # #2DD4BF
    WHITE = RGBColor(255, 255, 255)
    GRAY_TEXT = RGBColor(71, 85, 105)    # #475569
    SLATE_LIGHT = RGBColor(100, 116, 139)# #64748B
    CARD_BG = RGBColor(248, 250, 252)    # #F8FAFC
    BORDER_GRAY = RGBColor(203, 213, 225)# #CBD5E1
    BORDER_LIGHT = RGBColor(226, 232, 240)# #E2E8F0
    TEXT_MUTED = RGBColor(100, 116, 139)
    AMBER_BG = RGBColor(254, 243, 199)   # #FEF3C7
    AMBER_TEXT = RGBColor(146, 64, 14)   # #92400E
    BLUE_BG = RGBColor(224, 242, 254)    # #E0F2FE
    BLUE_TEXT = RGBColor(7, 89, 133)     # #075985
    GREEN_BG = RGBColor(220, 252, 231)   # #DCFCE7
    GREEN_TEXT = RGBColor(22, 101, 52)   # #166534

    # 1. Team Name Badge
    team_badge = slide3.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.80), Inches(0.32), Inches(1.65), Inches(0.42)
    )
    team_badge.name = "Team_Badge"
    team_badge.fill.solid()
    team_badge.fill.fore_color.rgb = WHITE
    team_badge.line.color.rgb = BORDER_TEAL
    team_badge.line.width = Pt(1.5)
    tf_tb = team_badge.text_frame
    tf_tb.word_wrap = False
    tf_tb.margin_left = Inches(0.05)
    tf_tb.margin_right = Inches(0.05)
    tf_tb.margin_top = Inches(0.06)
    tf_tb.margin_bottom = Inches(0.06)
    p_tb = tf_tb.paragraphs[0]
    p_tb.alignment = PP_ALIGN.CENTER
    r_tb = p_tb.add_run()
    r_tb.text = "TeamChargers"
    r_tb.font.name = "Calibri"
    r_tb.font.size = Pt(12)
    r_tb.font.bold = True
    r_tb.font.color.rgb = NAVY

    # 2. Slide Title (Single crisp line, 8.0" wide)
    title_box = slide3.shapes.add_textbox(Inches(2.65), Inches(0.28), Inches(8.00), Inches(0.45))
    title_box.name = "Slide_Title_Box"
    tf_title = title_box.text_frame
    tf_title.word_wrap = False
    tf_title.margin_left = Inches(0)
    tf_title.margin_top = Inches(0)
    p_title = tf_title.paragraphs[0]
    p_title.alignment = PP_ALIGN.LEFT
    
    r_t1 = p_title.add_run()
    r_t1.text = "TECHNICAL APPROACH: "
    r_t1.font.name = "Georgia"
    r_t1.font.size = Pt(20)
    r_t1.font.bold = True
    r_t1.font.color.rgb = NAVY

    r_t2 = p_title.add_run()
    r_t2.text = "End-to-End System Architecture"
    r_t2.font.name = "Calibri"
    r_t2.font.size = Pt(18)
    r_t2.font.bold = True
    r_t2.font.color.rgb = TEAL

    # 3. Sub-header Description
    desc_box = slide3.shapes.add_textbox(Inches(0.80), Inches(0.82), Inches(9.80), Inches(0.28))
    desc_box.name = "Description_Box"
    tf_desc = desc_box.text_frame
    tf_desc.word_wrap = True
    tf_desc.margin_left = Inches(0)
    tf_desc.margin_top = Inches(0)
    p_desc = tf_desc.paragraphs[0]
    r_desc = p_desc.add_run()
    r_desc.text = "A multi-tier decision-support pipeline integrating vernacular voice inputs, grounded RAG intelligence, and administrative analytics."
    r_desc.font.name = "Calibri"
    r_desc.font.size = Pt(10)
    r_desc.font.color.rgb = GRAY_TEXT

    # =========================================================================
    # MAIN ARCHITECTURE DIAGRAM (Horizontal 4-Tier Pipeline)
    # Total Width: 11.733" (0.80" to 12.533")
    # Top: 1.22" to 5.20" (Height = 3.98")
    # =========================================================================
    tier_top = Inches(1.22)
    tier_h = Inches(3.98)
    
    # -------------------------------------------------------------------------
    # TIER 1: CLIENT & INTERACTION LAYER
    # -------------------------------------------------------------------------
    t1_left = Inches(0.80)
    t1_w = Inches(2.60)
    t1_box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, t1_left, tier_top, t1_w, tier_h)
    t1_box.fill.solid()
    t1_box.fill.fore_color.rgb = CARD_BG
    t1_box.line.color.rgb = BORDER_GRAY
    t1_box.line.width = Pt(1.0)
    
    # Tier 1 Header (Generous top breathing room)
    t1_hdr = slide3.shapes.add_textbox(t1_left + Inches(0.12), tier_top + Inches(0.16), t1_w - Inches(0.24), Inches(0.50))
    t1_tf = t1_hdr.text_frame
    t1_tf.margin_left = 0
    t1_tf.margin_top = 0
    p_t1 = t1_tf.paragraphs[0]
    r_t1_1 = p_t1.add_run()
    r_t1_1.text = "TIER 1: INTERACTION\n"
    r_t1_1.font.size = Pt(8.0)
    r_t1_1.font.bold = True
    r_t1_1.font.color.rgb = SLATE_LIGHT
    r_t1_2 = p_t1.add_run()
    r_t1_2.text = "Dual-Persona Client Layer"
    r_t1_2.font.size = Pt(10)
    r_t1_2.font.bold = True
    r_t1_2.font.color.rgb = NAVY

    # Inside Tier 1: 3 Modular Blocks
    b1_1 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, t1_left + Inches(0.12), tier_top + Inches(0.72), t1_w - Inches(0.24), Inches(0.94))
    b1_1.fill.solid()
    b1_1.fill.fore_color.rgb = WHITE
    b1_1.line.color.rgb = BORDER_LIGHT
    tf_b1_1 = b1_1.text_frame
    tf_b1_1.margin_left = Inches(0.08)
    tf_b1_1.margin_top = Inches(0.06)
    p_b1_1 = tf_b1_1.paragraphs[0]
    r = p_b1_1.add_run()
    r.text = "Student Web Portal\n"
    r.font.size = Pt(9.0)
    r.font.bold = True
    r.font.color.rgb = NAVY
    r_sub = p_b1_1.add_run()
    r_sub.text = "• React 18, Vite & TypeScript\n• Aptitude & Career Explorer\n• Interactive NSQF pathway maps"
    r_sub.font.size = Pt(7.5)
    r_sub.font.color.rgb = GRAY_TEXT

    b1_2 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, t1_left + Inches(0.12), tier_top + Inches(1.74), t1_w - Inches(0.24), Inches(0.94))
    b1_2.fill.solid()
    b1_2.fill.fore_color.rgb = WHITE
    b1_2.line.color.rgb = BORDER_LIGHT
    tf_b1_2 = b1_2.text_frame
    tf_b1_2.margin_left = Inches(0.08)
    tf_b1_2.margin_top = Inches(0.06)
    p_b1_2 = tf_b1_2.paragraphs[0]
    r = p_b1_2.add_run()
    r.text = "Parent Assisted Portal\n"
    r.font.size = Pt(9.0)
    r.font.bold = True
    r.font.color.rgb = NAVY
    r_sub = p_b1_2.add_run()
    r_sub.text = "• Simplified vernacular UI\n• Financial ROI & Security focus\n• Audio-first concern submission"
    r_sub.font.size = Pt(7.5)
    r_sub.font.color.rgb = GRAY_TEXT

    b1_3 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, t1_left + Inches(0.12), tier_top + Inches(2.76), t1_w - Inches(0.24), Inches(1.10))
    b1_3.fill.solid()
    b1_3.fill.fore_color.rgb = BLUE_BG
    b1_3.line.color.rgb = BORDER_TEAL
    tf_b1_3 = b1_3.text_frame
    tf_b1_3.margin_left = Inches(0.08)
    tf_b1_3.margin_top = Inches(0.06)
    p_b1_3 = tf_b1_3.paragraphs[0]
    r = p_b1_3.add_run()
    r.text = "Vernacular Language Gateway\n"
    r.font.size = Pt(9.0)
    r.font.bold = True
    r.font.color.rgb = BLUE_TEXT
    r_sub = p_b1_3.add_run()
    r_sub.text = "• Bhashini API & Whisper AI\n• Real-time speech-to-text (12+ Indic languages)\n• Eliminates literacy & tech barriers"
    r_sub.font.size = Pt(7.5)
    r_sub.font.color.rgb = BLUE_TEXT

    # Pipeline Arrow 1 -> 2
    arr_t12 = slide3.shapes.add_textbox(Inches(3.42), Inches(2.80), Inches(0.22), Inches(0.35))
    arr_t12.text_frame.margin_left = 0
    arr_t12.text_frame.margin_top = 0
    p_a1 = arr_t12.text_frame.paragraphs[0]
    r_a1 = p_a1.add_run()
    r_a1.text = "→"
    r_a1.font.size = Pt(18)
    r_a1.font.bold = True
    r_a1.font.color.rgb = TEAL

    # -------------------------------------------------------------------------
    # TIER 2: API & MEDIATION GATEWAY
    # -------------------------------------------------------------------------
    t2_left = Inches(3.66)
    t2_w = Inches(2.62)
    t2_box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, t2_left, tier_top, t2_w, tier_h)
    t2_box.fill.solid()
    t2_box.fill.fore_color.rgb = CARD_BG
    t2_box.line.color.rgb = BORDER_GRAY
    t2_box.line.width = Pt(1.0)

    # Tier 2 Header
    t2_hdr = slide3.shapes.add_textbox(t2_left + Inches(0.12), tier_top + Inches(0.16), t2_w - Inches(0.24), Inches(0.50))
    t2_tf = t2_hdr.text_frame
    t2_tf.margin_left = 0
    t2_tf.margin_top = 0
    p_t2 = t2_tf.paragraphs[0]
    r_t2_1 = p_t2.add_run()
    r_t2_1.text = "TIER 2: ORCHESTRATION\n"
    r_t2_1.font.size = Pt(8.0)
    r_t2_1.font.bold = True
    r_t2_1.font.color.rgb = SLATE_LIGHT
    r_t2_2 = p_t2.add_run()
    r_t2_2.text = "FastAPI Backend & Mediation"
    r_t2_2.font.size = Pt(10)
    r_t2_2.font.bold = True
    r_t2_2.font.color.rgb = NAVY

    b2_1 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, t2_left + Inches(0.12), tier_top + Inches(0.72), t2_w - Inches(0.24), Inches(0.94))
    b2_1.fill.solid()
    b2_1.fill.fore_color.rgb = WHITE
    b2_1.line.color.rgb = BORDER_LIGHT
    tf_b2_1 = b2_1.text_frame
    tf_b2_1.margin_left = Inches(0.08)
    tf_b2_1.margin_top = Inches(0.06)
    p_b2_1 = tf_b2_1.paragraphs[0]
    r = p_b2_1.add_run()
    r.text = "FastAPI Async Core\n"
    r.font.size = Pt(9.0)
    r.font.bold = True
    r.font.color.rgb = NAVY
    r_sub = p_b2_1.add_run()
    r_sub.text = "• Python 3.11 asynchronous API\n• Pydantic type-safe contracts\n• WebSockets for live joint counselling"
    r_sub.font.size = Pt(7.5)
    r_sub.font.color.rgb = GRAY_TEXT

    b2_2 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, t2_left + Inches(0.12), tier_top + Inches(1.74), t2_w - Inches(0.24), Inches(1.06))
    b2_2.fill.solid()
    b2_2.fill.fore_color.rgb = WHITE
    b2_2.line.color.rgb = BORDER_LIGHT
    tf_b2_2 = b2_2.text_frame
    tf_b2_2.margin_left = Inches(0.08)
    tf_b2_2.margin_top = Inches(0.06)
    p_b2_2 = tf_b2_2.paragraphs[0]
    r = p_b2_2.add_run()
    r.text = "Joint Context Manager\n"
    r.font.size = Pt(9.0)
    r.font.bold = True
    r.font.color.rgb = NAVY
    r_sub = p_b2_2.add_run()
    r_sub.text = "• Maintains dual conversational threads\n• Disagreement & sentiment parser\n• Privacy-first session tokenization\n• Redis distributed session store"
    r_sub.font.size = Pt(7.5)
    r_sub.font.color.rgb = GRAY_TEXT

    b2_3 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, t2_left + Inches(0.12), tier_top + Inches(2.88), t2_w - Inches(0.24), Inches(0.98))
    b2_3.fill.solid()
    b2_3.fill.fore_color.rgb = WHITE
    b2_3.line.color.rgb = BORDER_LIGHT
    tf_b2_3 = b2_3.text_frame
    tf_b2_3.margin_left = Inches(0.08)
    tf_b2_3.margin_top = Inches(0.06)
    p_b2_3 = tf_b2_3.paragraphs[0]
    r = p_b2_3.add_run()
    r.text = "Security & DPDP Shield\n"
    r.font.size = Pt(9.0)
    r.font.bold = True
    r.font.color.rgb = NAVY
    r_sub = p_b2_3.add_run()
    r_sub.text = "• Anonymizes PII before AI processing\n• Explicit parental consent guardrails\n• JWT authorization & RBAC"
    r_sub.font.size = Pt(7.5)
    r_sub.font.color.rgb = GRAY_TEXT

    # Pipeline Arrow 2 -> 3
    arr_t23 = slide3.shapes.add_textbox(Inches(6.30), Inches(2.80), Inches(0.22), Inches(0.35))
    arr_t23.text_frame.margin_left = 0
    arr_t23.text_frame.margin_top = 0
    p_a2 = arr_t23.text_frame.paragraphs[0]
    r_a2 = p_a2.add_run()
    r_a2.text = "→"
    r_a2.font.size = Pt(18)
    r_a2.font.bold = True
    r_a2.font.color.rgb = TEAL

    # -------------------------------------------------------------------------
    # TIER 3: INTELLIGENCE & GROUNDED RAG CORE (Highlighted in Teal)
    # -------------------------------------------------------------------------
    t3_left = Inches(6.54)
    t3_w = Inches(3.20)
    t3_box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, t3_left, tier_top, t3_w, tier_h)
    t3_box.fill.solid()
    t3_box.fill.fore_color.rgb = LIGHT_TEAL
    t3_box.line.color.rgb = TEAL
    t3_box.line.width = Pt(1.75)

    # Tier 3 Header
    t3_hdr = slide3.shapes.add_textbox(t3_left + Inches(0.12), tier_top + Inches(0.16), t3_w - Inches(0.24), Inches(0.50))
    t3_tf = t3_hdr.text_frame
    t3_tf.margin_left = 0
    t3_tf.margin_top = 0
    p_t3 = t3_tf.paragraphs[0]
    r_t3_1 = p_t3.add_run()
    r_t3_1.text = "TIER 3: INTELLIGENCE CORE ★\n"
    r_t3_1.font.size = Pt(8.0)
    r_t3_1.font.bold = True
    r_t3_1.font.color.rgb = TEAL
    r_t3_2 = p_t3.add_run()
    r_t3_2.text = "Consensus Intelligence & RAG Engine"
    r_t3_2.font.size = Pt(10)
    r_t3_2.font.bold = True
    r_t3_2.font.color.rgb = NAVY

    b3_1 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, t3_left + Inches(0.12), tier_top + Inches(0.72), t3_w - Inches(0.24), Inches(1.12))
    b3_1.fill.solid()
    b3_1.fill.fore_color.rgb = WHITE
    b3_1.line.color.rgb = BORDER_TEAL
    tf_b3_1 = b3_1.text_frame
    tf_b3_1.margin_left = Inches(0.08)
    tf_b3_1.margin_top = Inches(0.06)
    p_b3_1 = tf_b3_1.paragraphs[0]
    r = p_b3_1.add_run()
    r.text = "Family Consensus & Dialogue Agent\n"
    r.font.size = Pt(9.0)
    r.font.bold = True
    r.font.color.rgb = DARK_TEAL
    r_sub = p_b3_1.add_run()
    r_sub.text = "• LangChain / LangGraph Agentic Pipeline\n• Identifies root parental objections (pay, stability, stigma)\n• Injects evidence-based counter-perspectives\n• Measures pre/post confidence progression"
    r_sub.font.size = Pt(7.5)
    r_sub.font.color.rgb = GRAY_TEXT

    b3_2 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, t3_left + Inches(0.12), tier_top + Inches(1.92), t3_w - Inches(0.24), Inches(1.08))
    b3_2.fill.solid()
    b3_2.fill.fore_color.rgb = WHITE
    b3_2.line.color.rgb = BORDER_TEAL
    tf_b3_2 = b3_2.text_frame
    tf_b3_2.margin_left = Inches(0.08)
    tf_b3_2.margin_top = Inches(0.06)
    p_b3_2 = tf_b3_2.paragraphs[0]
    r = p_b3_2.add_run()
    r.text = "Verified RAG Knowledge Engine\n"
    r.font.size = Pt(9.0)
    r.font.bold = True
    r.font.color.rgb = DARK_TEAL
    r_sub = p_b3_2.add_run()
    r_sub.text = "• Vector DB (Chroma / Pinecone) with hybrid search\n• Grounded in official NCVET & NSQF curriculum\n• National Career Service (NCS) wage benchmarks\n• Strict zero-hallucination constraint"
    r_sub.font.size = Pt(7.5)
    r_sub.font.color.rgb = GRAY_TEXT

    b3_3 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, t3_left + Inches(0.12), tier_top + Inches(3.08), t3_w - Inches(0.24), Inches(0.78))
    b3_3.fill.solid()
    b3_3.fill.fore_color.rgb = WHITE
    b3_3.line.color.rgb = BORDER_TEAL
    tf_b3_3 = b3_3.text_frame
    tf_b3_3.margin_left = Inches(0.08)
    tf_b3_3.margin_top = Inches(0.06)
    p_b3_3 = tf_b3_3.paragraphs[0]
    r = p_b3_3.add_run()
    r.text = "Explainable Trade-off Engine\n"
    r.font.size = Pt(9.0)
    r.font.bold = True
    r.font.color.rgb = DARK_TEAL
    r_sub = p_b3_3.add_run()
    r_sub.text = "• Computes Training Cost vs Time-to-Earn ROI\n• Visualizes NSQF upward degree pathways"
    r_sub.font.size = Pt(7.5)
    r_sub.font.color.rgb = GRAY_TEXT

    # Pipeline Arrow 3 -> 4
    arr_t34 = slide3.shapes.add_textbox(Inches(9.76), Inches(2.80), Inches(0.22), Inches(0.35))
    arr_t34.text_frame.margin_left = 0
    arr_t34.text_frame.margin_top = 0
    p_a3 = arr_t34.text_frame.paragraphs[0]
    r_a3 = p_a3.add_run()
    r_a3.text = "→"
    r_a3.font.size = Pt(18)
    r_a3.font.bold = True
    r_a3.font.color.rgb = TEAL

    # -------------------------------------------------------------------------
    # TIER 4: OUTCOME, ESCALATION & ANALYTICS LAYER
    # -------------------------------------------------------------------------
    t4_left = Inches(10.00)
    t4_w = Inches(2.53)
    t4_box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, t4_left, tier_top, t4_w, tier_h)
    t4_box.fill.solid()
    t4_box.fill.fore_color.rgb = CARD_BG
    t4_box.line.color.rgb = BORDER_GRAY
    t4_box.line.width = Pt(1.0)

    # Tier 4 Header
    t4_hdr = slide3.shapes.add_textbox(t4_left + Inches(0.12), tier_top + Inches(0.16), t4_w - Inches(0.24), Inches(0.50))
    t4_tf = t4_hdr.text_frame
    t4_tf.margin_left = 0
    t4_tf.margin_top = 0
    p_t4 = t4_tf.paragraphs[0]
    r_t4_1 = p_t4.add_run()
    r_t4_1.text = "TIER 4: ACTION & INSIGHTS\n"
    r_t4_1.font.size = Pt(8.0)
    r_t4_1.font.bold = True
    r_t4_1.font.color.rgb = SLATE_LIGHT
    r_t4_2 = p_t4.add_run()
    r_t4_2.text = "Outcomes & Administration"
    r_t4_2.font.size = Pt(10)
    r_t4_2.font.bold = True
    r_t4_2.font.color.rgb = NAVY

    b4_1 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, t4_left + Inches(0.12), tier_top + Inches(0.72), t4_w - Inches(0.24), Inches(0.94))
    b4_1.fill.solid()
    b4_1.fill.fore_color.rgb = GREEN_BG
    b4_1.line.color.rgb = GREEN_BG
    tf_b4_1 = b4_1.text_frame
    tf_b4_1.margin_left = Inches(0.08)
    tf_b4_1.margin_top = Inches(0.06)
    p_b4_1 = tf_b4_1.paragraphs[0]
    r = p_b4_1.add_run()
    r.text = "Family Consensus Summary\n"
    r.font.size = Pt(9.0)
    r.font.bold = True
    r.font.color.rgb = GREEN_TEXT
    r_sub = p_b4_1.add_run()
    r_sub.text = "• Downloadable Career Roadmaps\n• Aligned student-parent consensus score\n• Direct admission & scholarship links"
    r_sub.font.size = Pt(7.5)
    r_sub.font.color.rgb = GREEN_TEXT

    b4_2 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, t4_left + Inches(0.12), tier_top + Inches(1.74), t4_w - Inches(0.24), Inches(0.94))
    b4_2.fill.solid()
    b4_2.fill.fore_color.rgb = WHITE
    b4_2.line.color.rgb = BORDER_LIGHT
    tf_b4_2 = b4_2.text_frame
    tf_b4_2.margin_left = Inches(0.08)
    tf_b4_2.margin_top = Inches(0.06)
    p_b4_2 = tf_b4_2.paragraphs[0]
    r = p_b4_2.add_run()
    r.text = "Counsellor Escalation\n"
    r.font.size = Pt(9.0)
    r.font.bold = True
    r.font.color.rgb = NAVY
    r_sub = p_b4_2.add_run()
    r_sub.text = "• Triggered by persistent impasse\n• Pre-packaged conflict brief sent to human counselor for 1-on-1 session"
    r_sub.font.size = Pt(7.5)
    r_sub.font.color.rgb = GRAY_TEXT

    b4_3 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, t4_left + Inches(0.12), tier_top + Inches(2.76), t4_w - Inches(0.24), Inches(1.10))
    b4_3.fill.solid()
    b4_3.fill.fore_color.rgb = WHITE
    b4_3.line.color.rgb = BORDER_LIGHT
    tf_b4_3 = b4_3.text_frame
    tf_b4_3.margin_left = Inches(0.08)
    tf_b4_3.margin_top = Inches(0.06)
    p_b4_3 = tf_b4_3.paragraphs[0]
    r = p_b4_3.add_run()
    r.text = "Admin Intelligence Portal\n"
    r.font.size = Pt(9.0)
    r.font.bold = True
    r.font.color.rgb = NAVY
    r_sub = p_b4_3.add_run()
    r_sub.text = "• Aggregated regional objection maps\n• Identifies state/district skill hesitations\n• Policy insights for MSDE & NSDC"
    r_sub.font.size = Pt(7.5)
    r_sub.font.color.rgb = GRAY_TEXT

    # =========================================================================
    # BOTTOM SECTION: TECHNOLOGY STACK & METHODOLOGY SUMMARY STRIP
    # Top: 5.34", Height: 1.48", Left: 0.80", Width: 11.733"
    # =========================================================================
    bottom_container = slide3.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.80), Inches(5.34), Inches(11.733), Inches(1.48)
    )
    bottom_container.name = "Bottom_Tech_Container"
    bottom_container.fill.solid()
    bottom_container.fill.fore_color.rgb = WHITE
    bottom_container.line.color.rgb = BORDER_LIGHT
    bottom_container.line.width = Pt(1.0)

    # 4 Tech Category Columns
    cat_w = Inches(2.80)
    cat_gap = Inches(0.15)
    cat_left_base = Inches(0.90)

    categories = [
        {
            "title": "FRONTEND & VOICE",
            "items": "• React 18, Vite, TypeScript\n• Vanilla CSS Design Tokens\n• Bhashini Indic Speech API\n• Progressive Web App (PWA)"
        },
        {
            "title": "BACKEND & AGENTS",
            "items": "• Python 3.11, FastAPI Async\n• LangChain / LangGraph Multi-Agent\n• Redis Dialogue Session Store\n• PostgreSQL Relational DB"
        },
        {
            "title": "GROUNDING & DATA",
            "items": "• Chroma / Pinecone Vector DB\n• Official NCVET Qualifications\n• National Career Service (NCS)\n• Regional ITI & NSTI Directory"
        },
        {
            "title": "SECURITY & DEPLOYMENT",
            "items": "• Docker & Kubernetes Containers\n• MeitY Empanelled Cloud (AWS/NIC)\n• DPDP Act Consent Architecture\n• End-to-End JWT Auth & RBAC"
        }
    ]

    for i, cat in enumerate(categories):
        cx = cat_left_base + i * (cat_w + cat_gap)
        cbox = slide3.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            cx, Inches(5.42), cat_w, Inches(1.32)
        )
        cbox.fill.solid()
        cbox.fill.fore_color.rgb = CARD_BG
        cbox.line.color.rgb = BORDER_GRAY
        cbox.line.width = Pt(0.75)
        
        tf_c = cbox.text_frame
        tf_c.margin_left = Inches(0.08)
        tf_c.margin_top = Inches(0.06)
        tf_c.word_wrap = True
        
        p_c1 = tf_c.paragraphs[0]
        r_c1 = p_c1.add_run()
        r_c1.text = cat["title"] + "\n"
        r_c1.font.name = "Calibri"
        r_c1.font.size = Pt(8.5)
        r_c1.font.bold = True
        r_c1.font.color.rgb = DARK_TEAL
        
        p_c2 = tf_c.add_paragraph()
        r_c2 = p_c2.add_run()
        r_c2.text = cat["items"]
        r_c2.font.name = "Calibri"
        r_c2.font.size = Pt(7.5)
        r_c2.font.color.rgb = GRAY_TEXT

    prs.save(ppt_path)
    print("Successfully built perfected Slide 3 in:", ppt_path)

if __name__ == "__main__":
    apply_slide_3()
