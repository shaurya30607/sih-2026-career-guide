import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

pptx_path = "SIH 2.0/SIH2025-IDEA-Presentation-Format.pptx"
prs = Presentation(pptx_path)

slide2 = prs.slides[1]

# Palette - Editorial Minimalist
COLOR_NAVY_DARK = RGBColor(0x0B, 0x13, 0x2B)
COLOR_NAVY_TITLE = RGBColor(0x1F, 0x4E, 0x79)
COLOR_SLATE_DARK = RGBColor(0x1E, 0x29, 0x3B)
COLOR_SLATE_MUTED = RGBColor(0x47, 0x55, 0x69)
COLOR_SLATE_LIGHT = RGBColor(0x64, 0x74, 0x8B)

COLOR_TEAL_DARK = RGBColor(0x0D, 0x94, 0x88)
COLOR_TEAL_ACCENT = RGBColor(0x14, 0xB8, 0xA6)
COLOR_TEAL_TINT = RGBColor(0xF0, 0xFD, 0xFA)
COLOR_TEAL_BORDER = RGBColor(0x99, 0xF6, 0xE4)

COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_OFFWHITE = RGBColor(0xF8, 0xF9, 0xFA)
COLOR_CARD_BORDER = RGBColor(0xE2, 0xE8, 0xF0)
COLOR_LINE_BORDER = RGBColor(0xCB, 0xD5, 0xE1)

# 1. Update Oval 9 (Team Name Badge at top-left)
for s in slide2.shapes:
    if s.name.startswith("Oval"):
        s.left = Inches(0.70)
        s.top = Inches(0.30)
        s.width = Inches(1.75)
        s.height = Inches(0.56)
        s.fill.solid()
        s.fill.fore_color.rgb = COLOR_OFFWHITE
        s.line.color.rgb = COLOR_LINE_BORDER
        s.line.width = Pt(1)
        if s.has_text_frame:
            tf = s.text_frame
            tf.clear()
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            p = tf.paragraphs[0]
            p.text = "TeamChargers"
            p.alignment = PP_ALIGN.CENTER
            p.font.name = "Arial"
            p.font.size = Pt(10.5)
            p.font.bold = True
            p.font.color.rgb = COLOR_NAVY_DARK

# 2. Update Title 1
for s in slide2.shapes:
    if s.name == "Title 1":
        s.left = Inches(2.55)
        s.top = Inches(0.12)
        s.width = Inches(8.05)
        s.height = Inches(1.05)
        if s.has_text_frame:
            tf = s.text_frame
            tf.word_wrap = True
            tf.clear()
            # Mandatory Template Heading + Project
            p1 = tf.paragraphs[0]
            r1 = p1.add_run()
            r1.text = "IDEA TITLE: "
            r1.font.name = "Garamond"
            r1.font.size = Pt(24)
            r1.font.bold = True
            r1.font.color.rgb = COLOR_NAVY_TITLE
            
            r2 = p1.add_run()
            r2.text = "SkillSaathi — Proposed Solution"
            r2.font.name = "Arial"
            r2.font.size = Pt(22)
            r2.font.bold = True
            r2.font.color.rgb = COLOR_NAVY_DARK
            p1.space_after = Pt(2)
            
            # Headline
            p2 = tf.add_paragraph()
            r3 = p2.add_run()
            r3.text = "\"Beyond Career Advice. Towards Family Consensus.\""
            r3.font.name = "Arial"
            r3.font.size = Pt(14)
            r3.font.italic = True
            r3.font.color.rgb = COLOR_TEAL_DARK

# 3. Remove previous content shapes added to slide2 if re-running
shapes_to_remove = [s for s in slide2.shapes if s.name.startswith("Dynamic_") or s.name == "TextBox 8"]
for s in shapes_to_remove:
    sp = s._element
    sp.getparent().remove(sp)

# ==============================================================================
# SECTION A: NARRATIVE FLOW (Problem → Solution → Outcome)
# Addressing "How it addresses the problem"
# Top: 1.35", Height: 0.95"
# ==============================================================================
flow_top = Inches(1.35)
flow_h = Inches(0.95)
box_w = Inches(3.64)
gap_w = Inches(0.40)
start_x = Inches(0.80)

steps_data = [
    ("THE PROBLEM", "Student Aspirations vs. Parent Anxiety", "Vocational stigma and lack of verified data create deep family friction.", COLOR_OFFWHITE, COLOR_CARD_BORDER, COLOR_SLATE_LIGHT),
    ("THE SOLUTION", "Dual-Persona Joint AI Counselling", "AI facilitates synchronized dialogue, grounding choices in NSQF & ROI data.", COLOR_TEAL_TINT, COLOR_TEAL_BORDER, COLOR_TEAL_DARK),
    ("THE OUTCOME", "High-Confidence Family Consensus", "Families overcome hesitation to make informed, united vocational decisions.", COLOR_OFFWHITE, COLOR_CARD_BORDER, COLOR_SLATE_LIGHT)
]

for i, (tag, title, desc, bg_col, border_col, tag_col) in enumerate(steps_data):
    bx = start_x + i * (box_w + gap_w)
    
    # Card Background
    card = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, bx, flow_top, box_w, flow_h)
    card.name = f"Dynamic_FlowCard_{i}"
    card.fill.solid()
    card.fill.fore_color.rgb = bg_col
    card.line.color.rgb = border_col
    card.line.width = Pt(1)
    
    # Text
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.08)
    tf.margin_bottom = Inches(0.06)
    
    p0 = tf.paragraphs[0]
    p0.text = tag
    p0.font.name = "Arial"
    p0.font.size = Pt(8.5)
    p0.font.bold = True
    p0.font.color.rgb = tag_col
    p0.space_after = Pt(2)
    
    p1 = tf.add_paragraph()
    p1.text = title
    p1.font.name = "Arial"
    p1.font.size = Pt(11)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_NAVY_DARK
    p1.space_after = Pt(2)
    
    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.name = "Arial"
    p2.font.size = Pt(9)
    p2.font.color.rgb = COLOR_SLATE_MUTED

    # Arrow separator between cards
    if i < 2:
        arrow_x = bx + box_w + Inches(0.10)
        tb_arr = slide2.shapes.add_textbox(arrow_x, flow_top + Inches(0.25), gap_w - Inches(0.20), Inches(0.40))
        tb_arr.name = f"Dynamic_FlowArrow_{i}"
        tf_a = tb_arr.text_frame
        tf_a.margin_left = tf_a.margin_right = tf_a.margin_top = tf_a.margin_bottom = 0
        pa = tf_a.paragraphs[0]
        pa.text = "→"
        pa.alignment = PP_ALIGN.CENTER
        pa.font.name = "Arial"
        pa.font.size = Pt(18)
        pa.font.bold = True
        pa.font.color.rgb = COLOR_TEAL_ACCENT

# ==============================================================================
# SECTION B: CENTRAL USP — FAMILY CONSENSUS INTELLIGENCE DIAGRAM
# Dominant visual centerpiece (Top: 2.45", Height: 2.40")
# Addressing "Detailed explanation of the proposed solution" & "Innovation & uniqueness"
# ==============================================================================
usp_frame_top = Inches(2.45)
usp_frame_h = Inches(2.40)
usp_frame_w = Inches(11.733)

# Container Box
usp_frame = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.80), usp_frame_top, usp_frame_w, usp_frame_h)
usp_frame.name = "Dynamic_USPFrame"
usp_frame.fill.solid()
usp_frame.fill.fore_color.rgb = COLOR_WHITE
usp_frame.line.color.rgb = COLOR_LINE_BORDER
usp_frame.line.width = Pt(1)

# Left accent indicator line
usp_accent = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.80), usp_frame_top, Inches(0.08), usp_frame_h)
usp_accent.name = "Dynamic_USPAccent"
usp_accent.fill.solid()
usp_accent.fill.fore_color.rgb = COLOR_TEAL_DARK
usp_accent.line.fill.background()

# Section Label
tb_usp_lbl = slide2.shapes.add_textbox(Inches(1.05), usp_frame_top + Inches(0.12), Inches(11.00), Inches(0.35))
tb_usp_lbl.name = "Dynamic_USPLbl"
tf_usp_lbl = tb_usp_lbl.text_frame
tf_usp_lbl.margin_left = tf_usp_lbl.margin_right = tf_usp_lbl.margin_top = tf_usp_lbl.margin_bottom = 0
pu_lbl = tf_usp_lbl.paragraphs[0]
r_u1 = pu_lbl.add_run()
r_u1.text = "CORE USP ARCHITECTURE: "
r_u1.font.name = "Arial"
r_u1.font.size = Pt(9.5)
r_u1.font.bold = True
r_u1.font.color.rgb = COLOR_TEAL_DARK

r_u2 = pu_lbl.add_run()
r_u2.text = "FAMILY CONSENSUS INTELLIGENCE™ PIPELINE"
r_u2.font.name = "Arial"
r_u2.font.size = Pt(9.5)
r_u2.font.bold = True
r_u2.font.color.rgb = COLOR_NAVY_DARK

# 4 Horizontal Pipeline Stages
pipeline_top = usp_frame_top + Inches(0.50)
stage_w = Inches(2.60)
stage_gap = Inches(0.40)
stage_start_x = Inches(1.05)
stage_h = Inches(1.68)

pipeline_stages = [
    ("STAGE 1", "Dual-Persona Intake", "Student Aspirations + Parent Concerns", 
     "• Maps student aptitude & passions\n• Identifies parental financial & ROI worries\n• Separates emotional bias from data", COLOR_OFFWHITE, COLOR_CARD_BORDER),
    ("STAGE 2", "Evidence-Backed AI", "NSQF & Economic Benchmarking", 
     "• Official NSQF/NCVET accredited pathways\n• Real-world salary & placement outcomes\n• Trade-specific career progression maps", COLOR_OFFWHITE, COLOR_CARD_BORDER),
    ("STAGE 3", "Consensus Engine", "Confidence & Objection Scoring", 
     "• AI detects misalignments & objections\n• Objective trade-off comparison matrix\n• Quantified family readiness index", COLOR_TEAL_TINT, COLOR_TEAL_BORDER),
    ("STAGE 4", "Aligned Decision", "Informed Actionable Roadmap", 
     "• Unified student-parent agreement\n• Concrete vocational enrollment steps\n• Seamless human counsellor escalation", COLOR_OFFWHITE, COLOR_CARD_BORDER)
]

for idx, (st_num, st_name, st_lead, st_bullets, bg_col, b_col) in enumerate(pipeline_stages):
    sx = stage_start_x + idx * (stage_w + stage_gap)
    
    stage_box = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, sx, pipeline_top, stage_w, stage_h)
    stage_box.name = f"Dynamic_PipelineStage_{idx}"
    stage_box.fill.solid()
    stage_box.fill.fore_color.rgb = bg_col
    stage_box.line.color.rgb = b_col
    stage_box.line.width = Pt(1)
    
    tf_s = stage_box.text_frame
    tf_s.word_wrap = True
    tf_s.margin_left = Inches(0.12)
    tf_s.margin_right = Inches(0.12)
    tf_s.margin_top = Inches(0.08)
    tf_s.margin_bottom = Inches(0.06)
    
    ps0 = tf_s.paragraphs[0]
    ps0.text = st_num
    ps0.font.name = "Arial"
    ps0.font.size = Pt(8)
    ps0.font.bold = True
    ps0.font.color.rgb = COLOR_TEAL_DARK if idx == 2 else COLOR_SLATE_LIGHT
    ps0.space_after = Pt(1)
    
    ps1 = tf_s.add_paragraph()
    ps1.text = st_name
    ps1.font.name = "Arial"
    ps1.font.size = Pt(11)
    ps1.font.bold = True
    ps1.font.color.rgb = COLOR_NAVY_DARK
    ps1.space_after = Pt(2)
    
    ps2 = tf_s.add_paragraph()
    ps2.text = st_lead
    ps2.font.name = "Arial"
    ps2.font.size = Pt(8.5)
    ps2.font.bold = True
    ps2.font.color.rgb = COLOR_TEAL_DARK if idx == 2 else COLOR_SLATE_DARK
    ps2.space_after = Pt(4)
    
    ps3 = tf_s.add_paragraph()
    ps3.text = st_bullets
    ps3.font.name = "Arial"
    ps3.font.size = Pt(8.5)
    ps3.font.color.rgb = COLOR_SLATE_MUTED

    # Connecting Arrow between stages
    if idx < 3:
        arr_x = sx + stage_w + Inches(0.08)
        tb_p_arr = slide2.shapes.add_textbox(arr_x, pipeline_top + Inches(0.60), stage_gap - Inches(0.16), Inches(0.40))
        tb_p_arr.name = f"Dynamic_PipelineArrow_{idx}"
        tf_pa = tb_p_arr.text_frame
        tf_pa.margin_left = tf_pa.margin_right = tf_pa.margin_top = tf_pa.margin_bottom = 0
        ppa = tf_pa.paragraphs[0]
        ppa.text = "→"
        ppa.alignment = PP_ALIGN.CENTER
        ppa.font.name = "Arial"
        ppa.font.size = Pt(16)
        ppa.font.bold = True
        ppa.font.color.rgb = COLOR_TEAL_ACCENT

# ==============================================================================
# SECTION C: SUPPORTING INNOVATIONS (What Makes It Different from Generic Chatbots)
# Top: 4.98", Height: 1.68"
# Answering "Innovation and uniqueness of the solution"
# ==============================================================================
innov_top = Inches(4.98)
innov_h = Inches(1.68)
innov_w = Inches(3.64)
innov_gap = Inches(0.40)

innovations = [
    ("1. VERIFIED CAREER INSIGHTS",
     "Evidence-Based Outcomes Over Advice",
     "Replaces generic career advice with verified industry data: salary percentiles, placement ratios, NSQF credit mobility, and real employer demand across trades."),
    
    ("2. MULTILINGUAL ACCESSIBILITY",
     "Designed for Low-Literacy Families",
     "Vernacular voice and text interactions in regional languages, enabling non-English speaking parents to participate equally in their child's career decisions."),
    
    ("3. HUMAN & ADMIN ESCALATION",
     "Safety Net & Resistance Analytics",
     "Seamless handoff to certified human career counsellors when high resistance is flagged; provides MSDE administrators systemic resistance heatmaps.")
]

for idx, (head, sub, detail) in enumerate(innovations):
    ix = start_x + idx * (innov_w + innov_gap)
    
    ibox = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, ix, innov_top, innov_w, innov_h)
    ibox.name = f"Dynamic_InnovCard_{idx}"
    ibox.fill.solid()
    ibox.fill.fore_color.rgb = COLOR_OFFWHITE
    ibox.line.color.rgb = COLOR_CARD_BORDER
    ibox.line.width = Pt(1)
    
    tf_i = ibox.text_frame
    tf_i.word_wrap = True
    tf_i.margin_left = Inches(0.14)
    tf_i.margin_right = Inches(0.14)
    tf_i.margin_top = Inches(0.08)
    tf_i.margin_bottom = Inches(0.06)
    
    pi0 = tf_i.paragraphs[0]
    pi0.text = head
    pi0.font.name = "Arial"
    pi0.font.size = Pt(8.5)
    pi0.font.bold = True
    pi0.font.color.rgb = COLOR_TEAL_DARK
    pi0.space_after = Pt(2)
    
    pi1 = tf_i.add_paragraph()
    pi1.text = sub
    pi1.font.name = "Arial"
    pi1.font.size = Pt(11)
    pi1.font.bold = True
    pi1.font.color.rgb = COLOR_NAVY_DARK
    pi1.space_after = Pt(3)
    
    pi2 = tf_i.add_paragraph()
    pi2.text = detail
    pi2.font.name = "Arial"
    pi2.font.size = Pt(8.5)
    pi2.font.color.rgb = COLOR_SLATE_MUTED

# Save presentation
prs.save(pptx_path)
print("Saved Slide 2 to:", pptx_path)
