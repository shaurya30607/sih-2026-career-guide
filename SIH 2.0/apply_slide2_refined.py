import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def apply_slide_2():
    ppt_path = "SIH 2.0/SIH2025-IDEA-Presentation-Format.pptx"
    prs = Presentation(ppt_path)
    slide2 = prs.slides[1] # 0-indexed 1 is Slide 2
    
    preserved_names = {"Rectangle 9", "Slide Number Placeholder 5", "Footer Placeholder 6", "Picture 2"}
    
    shapes_to_remove = [s for s in slide2.shapes if s.name not in preserved_names]
    for s in shapes_to_remove:
        sp = s._element
        sp.getparent().remove(sp)
        
    NAVY = RGBColor(11, 19, 43)          # #0B132B
    TEAL = RGBColor(13, 148, 136)        # #0D9488
    DARK_TEAL = RGBColor(15, 118, 110)   # #0F766E
    LIGHT_TEAL = RGBColor(240, 253, 250) # #F0FDFA
    BORDER_TEAL = RGBColor(45, 212, 191) # #2DD4BF
    WHITE = RGBColor(255, 255, 255)
    GRAY_TEXT = RGBColor(71, 85, 105)    # #475569
    SLATE_LIGHT = RGBColor(100, 116, 139)# #64748B
    CARD_BG = RGBColor(248, 250, 252)    # #F8FAFC
    BORDER_GRAY = RGBColor(203, 213, 225)# #CBD5E1
    BORDER_LIGHT = RGBColor(226, 232, 240)# #E2E8F0
    TEXT_MUTED = RGBColor(100, 116, 139)
    AMBER_BG = RGBColor(254, 243, 199)   # #FEF3C7
    AMBER_TEXT = RGBColor(146, 64, 14)   # #92400E
    BLUE_BG = RGBColor(224, 242, 254)    # #E0F2FE
    BLUE_TEXT = RGBColor(7, 89, 133)     # #075985
    GREEN_BG = RGBColor(220, 252, 231)   # #DCFCE7
    GREEN_TEXT = RGBColor(22, 101, 52)   # #166534
    ROSE_BG = RGBColor(255, 228, 230)    # #FFE4E6
    ROSE_TEXT = RGBColor(159, 18, 57)    # #9F1239

    # 1. Team Name Badge
    team_badge = slide2.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.80), Inches(0.35), Inches(1.65), Inches(0.42)
    )
    team_badge.name = "Team_Badge"
    team_badge.fill.solid()
    team_badge.fill.fore_color.rgb = WHITE
    team_badge.line.color.rgb = BORDER_TEAL
    team_badge.line.width = Pt(1.5)
    tf_tb = team_badge.text_frame
    tf_tb.word_wrap = False
    tf_tb.margin_left = Inches(0.05)
    tf_tb.margin_right = Inches(0.05)
    tf_tb.margin_top = Inches(0.06)
    tf_tb.margin_bottom = Inches(0.06)
    p_tb = tf_tb.paragraphs[0]
    p_tb.alignment = PP_ALIGN.CENTER
    r_tb = p_tb.add_run()
    r_tb.text = "TeamChargers"
    r_tb.font.name = "Calibri"
    r_tb.font.size = Pt(12)
    r_tb.font.bold = True
    r_tb.font.color.rgb = NAVY

    # 2. Slide Title
    title_box = slide2.shapes.add_textbox(Inches(2.65), Inches(0.28), Inches(7.70), Inches(0.55))
    title_box.name = "Slide_Title_Box"
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = Inches(0)
    tf_title.margin_top = Inches(0)
    p_title = tf_title.paragraphs[0]
    p_title.alignment = PP_ALIGN.LEFT
    
    r_t1 = p_title.add_run()
    r_t1.text = "IDEA TITLE: "
    r_t1.font.name = "Georgia"
    r_t1.font.size = Pt(21)
    r_t1.font.bold = True
    r_t1.font.color.rgb = NAVY

    r_t2 = p_title.add_run()
    r_t2.text = "SkillSaathi  "
    r_t2.font.name = "Calibri"
    r_t2.font.size = Pt(21)
    r_t2.font.bold = True
    r_t2.font.color.rgb = NAVY

    r_t3 = p_title.add_run()
    r_t3.text = "— Beyond Career Advice. Towards Family Consensus."
    r_t3.font.name = "Calibri"
    r_t3.font.size = Pt(13)
    r_t3.font.italic = True
    r_t3.font.bold = False
    r_t3.font.color.rgb = TEAL

    # 3. Sub-header Description
    desc_box = slide2.shapes.add_textbox(Inches(0.80), Inches(0.88), Inches(9.60), Inches(0.30))
    desc_box.name = "Description_Box"
    tf_desc = desc_box.text_frame
    tf_desc.word_wrap = True
    tf_desc.margin_left = Inches(0)
    tf_desc.margin_top = Inches(0)
    p_desc = tf_desc.paragraphs[0]
    r_desc = p_desc.add_run()
    r_desc.text = "An AI-powered family counselling platform that bridges student aspirations and parental concerns through evidence-backed vocational guidance."
    r_desc.font.name = "Calibri"
    r_desc.font.size = Pt(10.5)
    r_desc.font.color.rgb = GRAY_TEXT

    # =========================================================================
    # TOP SECTION: PROPOSED SOLUTION WORKFLOW (Compact ~32% Height)
    # =========================================================================
    wf_container = slide2.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.80), Inches(1.26), Inches(11.733), Inches(1.88)
    )
    wf_container.name = "Workflow_Container"
    wf_container.fill.solid()
    wf_container.fill.fore_color.rgb = CARD_BG
    wf_container.line.color.rgb = BORDER_LIGHT
    wf_container.line.width = Pt(1.0)

    # Workflow Section Label (Left)
    wf_lbl_box = slide2.shapes.add_textbox(Inches(0.95), Inches(1.30), Inches(6.00), Inches(0.24))
    wf_lbl_tf = wf_lbl_box.text_frame
    wf_lbl_tf.margin_left = Inches(0)
    wf_lbl_tf.margin_top = Inches(0)
    p_wf_lbl = wf_lbl_tf.paragraphs[0]
    r_wf_lbl = p_wf_lbl.add_run()
    r_wf_lbl.text = "PROPOSED SOLUTION: FAMILY DECISION INTELLIGENCE WORKFLOW"
    r_wf_lbl.font.name = "Calibri"
    r_wf_lbl.font.size = Pt(8.5)
    r_wf_lbl.font.bold = True
    r_wf_lbl.font.color.rgb = SLATE_LIGHT

    # Accessibility Indicator (Right) - Factual Addition
    access_box = slide2.shapes.add_textbox(Inches(7.10), Inches(1.30), Inches(5.25), Inches(0.24))
    tf_acc = access_box.text_frame
    tf_acc.margin_left = Inches(0)
    tf_acc.margin_top = Inches(0)
    p_acc = tf_acc.paragraphs[0]
    p_acc.alignment = PP_ALIGN.RIGHT
    r_acc = p_acc.add_run()
    r_acc.text = "🌐 English + Regional Language  |  Simple, Family-Friendly Guidance"
    r_acc.font.name = "Calibri"
    r_acc.font.size = Pt(8.5)
    r_acc.font.bold = True
    r_acc.font.color.rgb = DARK_TEAL

    # Dual Input Boxes (Student & Parent)
    # Student Box
    s_box = slide2.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.95), Inches(1.58), Inches(2.15), Inches(0.68)
    )
    s_box.fill.solid()
    s_box.fill.fore_color.rgb = WHITE
    s_box.line.color.rgb = BORDER_GRAY
    s_box.line.width = Pt(1.0)
    tf_sb = s_box.text_frame
    tf_sb.margin_left = Inches(0.10)
    tf_sb.margin_top = Inches(0.08)
    tf_sb.word_wrap = True
    p_sb1 = tf_sb.paragraphs[0]
    r_sb1 = p_sb1.add_run()
    r_sb1.text = "Student Aspirations"
    r_sb1.font.name = "Calibri"
    r_sb1.font.size = Pt(10.5)
    r_sb1.font.bold = True
    r_sb1.font.color.rgb = NAVY
    p_sb2 = tf_sb.add_paragraph()
    r_sb2 = p_sb2.add_run()
    r_sb2.text = "Interests, Aptitude & Ambitions"
    r_sb2.font.name = "Calibri"
    r_sb2.font.size = Pt(8.5)
    r_sb2.font.color.rgb = TEXT_MUTED

    # Parent Box
    p_box = slide2.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.95), Inches(2.36), Inches(2.15), Inches(0.68)
    )
    p_box.fill.solid()
    p_box.fill.fore_color.rgb = WHITE
    p_box.line.color.rgb = BORDER_GRAY
    p_box.line.width = Pt(1.0)
    tf_pb = p_box.text_frame
    tf_pb.margin_left = Inches(0.10)
    tf_pb.margin_top = Inches(0.08)
    tf_pb.word_wrap = True
    p_pb1 = tf_pb.paragraphs[0]
    r_pb1 = p_pb1.add_run()
    r_pb1.text = "Parent Concerns"
    r_pb1.font.name = "Calibri"
    r_pb1.font.size = Pt(10.5)
    r_pb1.font.bold = True
    r_pb1.font.color.rgb = NAVY
    p_pb2 = tf_pb.add_paragraph()
    r_pb2 = p_pb2.add_run()
    r_pb2.text = "Financial Security, ROI & Stability"
    r_pb2.font.name = "Calibri"
    r_pb2.font.size = Pt(8.5)
    r_pb2.font.color.rgb = TEXT_MUTED

    # Converging arrows
    arr1 = slide2.shapes.add_textbox(Inches(3.12), Inches(1.72), Inches(0.28), Inches(0.35))
    arr1.text_frame.margin_left = 0
    arr1.text_frame.margin_top = 0
    p_arr1 = arr1.text_frame.paragraphs[0]
    r_arr1 = p_arr1.add_run()
    r_arr1.text = "↘"
    r_arr1.font.size = Pt(15)
    r_arr1.font.bold = True
    r_arr1.font.color.rgb = TEAL

    arr2 = slide2.shapes.add_textbox(Inches(3.12), Inches(2.50), Inches(0.28), Inches(0.35))
    arr2.text_frame.margin_left = 0
    arr2.text_frame.margin_top = 0
    p_arr2 = arr2.text_frame.paragraphs[0]
    r_arr2 = p_arr2.add_run()
    r_arr2.text = "↗"
    r_arr2.font.size = Pt(15)
    r_arr2.font.bold = True
    r_arr2.font.color.rgb = TEAL

    # AI Counselling Box
    ai_box = slide2.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(3.42), Inches(1.58), Inches(2.55), Inches(1.46)
    )
    ai_box.fill.solid()
    ai_box.fill.fore_color.rgb = WHITE
    ai_box.line.color.rgb = TEAL
    ai_box.line.width = Pt(1.75)
    tf_ai = ai_box.text_frame
    tf_ai.margin_left = Inches(0.12)
    tf_ai.margin_top = Inches(0.10)
    tf_ai.word_wrap = True
    p_ai1 = tf_ai.paragraphs[0]
    r_ai1 = p_ai1.add_run()
    r_ai1.text = "AI FAMILY COUNSELLING"
    r_ai1.font.name = "Calibri"
    r_ai1.font.size = Pt(10.5)
    r_ai1.font.bold = True
    r_ai1.font.color.rgb = NAVY
    
    p_ai2 = tf_ai.add_paragraph()
    r_ai2 = p_ai2.add_run()
    r_ai2.text = "Dual-Persona Dialogue Mediation"
    r_ai2.font.name = "Calibri"
    r_ai2.font.size = Pt(9.0)
    r_ai2.font.bold = True
    r_ai2.font.color.rgb = TEAL

    p_ai3 = tf_ai.add_paragraph()
    r_ai3 = p_ai3.add_run()
    r_ai3.text = "• Identifies hidden parental objections\n• Synthesizes passion with practical ROI"
    r_ai3.font.name = "Calibri"
    r_ai3.font.size = Pt(8.5)
    r_ai3.font.color.rgb = GRAY_TEXT

    # Arrow to Career Comparison
    arr3 = slide2.shapes.add_textbox(Inches(6.02), Inches(2.15), Inches(0.28), Inches(0.35))
    arr3.text_frame.margin_left = 0
    arr3.text_frame.margin_top = 0
    p_arr3 = arr3.text_frame.paragraphs[0]
    r_arr3 = p_arr3.add_run()
    r_arr3.text = "→"
    r_arr3.font.size = Pt(16)
    r_arr3.font.bold = True
    r_arr3.font.color.rgb = TEAL

    # Career Comparison Box
    cc_box = slide2.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(6.35), Inches(1.58), Inches(2.65), Inches(1.46)
    )
    cc_box.fill.solid()
    cc_box.fill.fore_color.rgb = WHITE
    cc_box.line.color.rgb = BORDER_GRAY
    cc_box.line.width = Pt(1.0)
    tf_cc = cc_box.text_frame
    tf_cc.margin_left = Inches(0.12)
    tf_cc.margin_top = Inches(0.10)
    tf_cc.word_wrap = True
    p_cc1 = tf_cc.paragraphs[0]
    r_cc1 = p_cc1.add_run()
    r_cc1.text = "CAREER COMPARISON"
    r_cc1.font.name = "Calibri"
    r_cc1.font.size = Pt(10.5)
    r_cc1.font.bold = True
    r_cc1.font.color.rgb = NAVY

    p_cc2 = tf_cc.add_paragraph()
    r_cc2 = p_cc2.add_run()
    r_cc2.text = "Evidence-Based Benchmarking"
    r_cc2.font.name = "Calibri"
    r_cc2.font.size = Pt(9.0)
    r_cc2.font.bold = True
    r_cc2.font.color.rgb = DARK_TEAL

    p_cc3 = tf_cc.add_paragraph()
    r_cc3 = p_cc3.add_run()
    r_cc3.text = "• Official NSQF/NCVET accredited paths\n• Real-world placement & salary data"
    r_cc3.font.name = "Calibri"
    r_cc3.font.size = Pt(8.5)
    r_cc3.font.color.rgb = GRAY_TEXT

    # Arrow to Decision
    arr4 = slide2.shapes.add_textbox(Inches(9.05), Inches(2.15), Inches(0.28), Inches(0.35))
    arr4.text_frame.margin_left = 0
    arr4.text_frame.margin_top = 0
    p_arr4 = arr4.text_frame.paragraphs[0]
    r_arr4 = p_arr4.add_run()
    r_arr4.text = "→"
    r_arr4.font.size = Pt(16)
    r_arr4.font.bold = True
    r_arr4.font.color.rgb = TEAL

    # Informed Decision Box (Corrected neutral confidence label)
    fd_box = slide2.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(9.38), Inches(1.58), Inches(2.95), Inches(1.46)
    )
    fd_box.fill.solid()
    fd_box.fill.fore_color.rgb = WHITE
    fd_box.line.color.rgb = BORDER_GRAY
    fd_box.line.width = Pt(1.0)
    tf_fd = fd_box.text_frame
    tf_fd.margin_left = Inches(0.12)
    tf_fd.margin_top = Inches(0.10)
    tf_fd.word_wrap = True
    p_fd1 = tf_fd.paragraphs[0]
    r_fd1 = p_fd1.add_run()
    r_fd1.text = "INFORMED FAMILY DECISION"
    r_fd1.font.name = "Calibri"
    r_fd1.font.size = Pt(10.5)
    r_fd1.font.bold = True
    r_fd1.font.color.rgb = NAVY

    p_fd2 = tf_fd.add_paragraph()
    r_fd2 = p_fd2.add_run()
    r_fd2.text = "Consensus & Career Clarity"
    r_fd2.font.name = "Calibri"
    r_fd2.font.size = Pt(9.0)
    r_fd2.font.bold = True
    r_fd2.font.color.rgb = DARK_TEAL

    p_fd3 = tf_fd.add_paragraph()
    r_fd3 = p_fd3.add_run()
    r_fd3.text = "• Aligned student-parent commitment\n• Objective confidence change measured"
    r_fd3.font.name = "Calibri"
    r_fd3.font.size = Pt(8.5)
    r_fd3.font.color.rgb = GRAY_TEXT

    # =========================================================================
    # BOTTOM SECTION: 2x2 INNOVATION GRID (65% area, ~3.65" total height)
    # =========================================================================
    grid_lbl_box = slide2.shapes.add_textbox(Inches(0.80), Inches(3.22), Inches(11.733), Inches(0.28))
    tf_grid_lbl = grid_lbl_box.text_frame
    tf_grid_lbl.margin_left = Inches(0)
    tf_grid_lbl.margin_top = Inches(0)
    p_grid_lbl = tf_grid_lbl.paragraphs[0]
    r_grid_lbl = p_grid_lbl.add_run()
    r_grid_lbl.text = "INNOVATION & UNIQUENESS — WHY SKILLSAATHI?"
    r_grid_lbl.font.name = "Calibri"
    r_grid_lbl.font.size = Pt(11.5)
    r_grid_lbl.font.bold = True
    r_grid_lbl.font.color.rgb = NAVY

    # Grid Dimensions
    col_w = Inches(5.72)
    card_h = Inches(1.64)
    c1_left = Inches(0.80)
    c2_left = Inches(6.81)
    r1_top = Inches(3.52)
    r2_top = Inches(5.26)

    diag_h = Inches(0.70)
    diag_w = col_w - Inches(0.24)
    
    row1_diag_y = r1_top + Inches(0.82)
    row1_inner_y = row1_diag_y + Inches(0.08)
    
    row2_diag_y = r2_top + Inches(0.82)
    row2_inner_y = row2_diag_y + Inches(0.08)
    
    chip_h = Inches(0.54)

    # -------------------------------------------------------------------------
    # CARD 01: FAMILY CONSENSUS INTELLIGENCE (FLAGSHIP USP) - Top Left
    # -------------------------------------------------------------------------
    c1 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c1_left, r1_top, col_w, card_h)
    c1.name = "Card_01_USP"
    c1.fill.solid()
    c1.fill.fore_color.rgb = LIGHT_TEAL
    c1.line.color.rgb = TEAL
    c1.line.width = Pt(1.75)

    c1_tb = slide2.shapes.add_textbox(c1_left + Inches(0.12), r1_top + Inches(0.07), col_w - Inches(0.24), Inches(0.70))
    tf_c1 = c1_tb.text_frame
    tf_c1.margin_left = Inches(0)
    tf_c1.margin_top = Inches(0)
    tf_c1.word_wrap = True
    
    p_c1_h = tf_c1.paragraphs[0]
    r_c1_badge = p_c1_h.add_run()
    r_c1_badge.text = "01. Family Consensus Intelligence "
    r_c1_badge.font.name = "Calibri"
    r_c1_badge.font.size = Pt(11)
    r_c1_badge.font.bold = True
    r_c1_badge.font.color.rgb = NAVY

    r_c1_star = p_c1_h.add_run()
    r_c1_star.text = "★ FLAGSHIP USP"
    r_c1_star.font.name = "Calibri"
    r_c1_star.font.size = Pt(9.0)
    r_c1_star.font.bold = True
    r_c1_star.font.color.rgb = TEAL

    p_c1_d = tf_c1.add_paragraph()
    r_c1_d = p_c1_d.add_run()
    r_c1_d.text = "Identifies student-parent disagreements, addresses parental objections, and tracks changes in family confidence."
    r_c1_d.font.name = "Calibri"
    r_c1_d.font.size = Pt(8.5)
    r_c1_d.font.color.rgb = GRAY_TEXT

    diag1_bg = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c1_left + Inches(0.12), row1_diag_y, diag_w, diag_h)
    diag1_bg.fill.solid()
    diag1_bg.fill.fore_color.rgb = WHITE
    diag1_bg.line.color.rgb = BORDER_TEAL
    diag1_bg.line.width = Pt(1.0)

    # Chip 1: Disagreement Detection
    ch1 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c1_left + Inches(0.20), row1_inner_y, Inches(1.50), chip_h)
    ch1.fill.solid()
    ch1.fill.fore_color.rgb = AMBER_BG
    ch1.line.color.rgb = AMBER_BG
    tf_ch1 = ch1.text_frame
    tf_ch1.margin_left = Inches(0.02)
    tf_ch1.margin_top = Inches(0.06)
    p_ch1 = tf_ch1.paragraphs[0]
    p_ch1.alignment = PP_ALIGN.CENTER
    r_ch1_1 = p_ch1.add_run()
    r_ch1_1.text = "Disagreement Detection\n"
    r_ch1_1.font.size = Pt(7.8)
    r_ch1_1.font.bold = True
    r_ch1_1.font.color.rgb = AMBER_TEXT
    r_ch1_2 = p_ch1.add_run()
    r_ch1_2.text = "Student vs Parent View"
    r_ch1_2.font.size = Pt(7.0)
    r_ch1_2.font.color.rgb = AMBER_TEXT

    arr_c1_1 = slide2.shapes.add_textbox(c1_left + Inches(1.72), row1_inner_y + Inches(0.08), Inches(0.24), Inches(0.30))
    arr_c1_1.text_frame.margin_left = 0
    arr_c1_1.text_frame.margin_top = 0
    p_a1 = arr_c1_1.text_frame.paragraphs[0]
    r_a1 = p_a1.add_run()
    r_a1.text = "→"
    r_a1.font.size = Pt(13)
    r_a1.font.bold = True
    r_a1.font.color.rgb = TEAL

    # Chip 2: Evidence Mediation
    ch2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c1_left + Inches(1.98), row1_inner_y, Inches(1.68), chip_h)
    ch2.fill.solid()
    ch2.fill.fore_color.rgb = LIGHT_TEAL
    ch2.line.color.rgb = BORDER_TEAL
    tf_ch2 = ch2.text_frame
    tf_ch2.margin_left = Inches(0.02)
    tf_ch2.margin_top = Inches(0.06)
    p_ch2 = tf_ch2.paragraphs[0]
    p_ch2.alignment = PP_ALIGN.CENTER
    r_ch2_1 = p_ch2.add_run()
    r_ch2_1.text = "AI Evidence Mediation\n"
    r_ch2_1.font.size = Pt(7.8)
    r_ch2_1.font.bold = True
    r_ch2_1.font.color.rgb = DARK_TEAL
    r_ch2_2 = p_ch2.add_run()
    r_ch2_2.text = "Targeted Objection Fact-Check"
    r_ch2_2.font.size = Pt(7.0)
    r_ch2_2.font.color.rgb = DARK_TEAL

    arr_c1_2 = slide2.shapes.add_textbox(c1_left + Inches(3.68), row1_inner_y + Inches(0.08), Inches(0.24), Inches(0.30))
    arr_c1_2.text_frame.margin_left = 0
    arr_c1_2.text_frame.margin_top = 0
    p_a2 = arr_c1_2.text_frame.paragraphs[0]
    r_a2 = p_a2.add_run()
    r_a2.text = "→"
    r_a2.font.size = Pt(13)
    r_a2.font.bold = True
    r_a2.font.color.rgb = TEAL

    # Chip 3: Measured Consensus (Factual revision: Confidence Change Measured)
    ch3 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c1_left + Inches(3.94), row1_inner_y, Inches(1.66), chip_h)
    ch3.fill.solid()
    ch3.fill.fore_color.rgb = GREEN_BG
    ch3.line.color.rgb = GREEN_BG
    tf_ch3 = ch3.text_frame
    tf_ch3.margin_left = Inches(0.02)
    tf_ch3.margin_top = Inches(0.06)
    p_ch3 = tf_ch3.paragraphs[0]
    p_ch3.alignment = PP_ALIGN.CENTER
    r_ch3_1 = p_ch3.add_run()
    r_ch3_1.text = "Consensus & Agreement\n"
    r_ch3_1.font.size = Pt(7.8)
    r_ch3_1.font.bold = True
    r_ch3_1.font.color.rgb = GREEN_TEXT
    r_ch3_2 = p_ch3.add_run()
    r_ch3_2.text = "Confidence Change Measured"
    r_ch3_2.font.size = Pt(7.0)
    r_ch3_2.font.color.rgb = GREEN_TEXT

    # -------------------------------------------------------------------------
    # CARD 02: HYPERLOCAL CAREER REALITY ENGINE - Top Right
    # -------------------------------------------------------------------------
    c2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2_left, r1_top, col_w, card_h)
    c2.name = "Card_02_Hyperlocal"
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = BORDER_GRAY
    c2.line.width = Pt(1.0)

    c2_tb = slide2.shapes.add_textbox(c2_left + Inches(0.12), r1_top + Inches(0.07), col_w - Inches(0.24), Inches(0.70))
    tf_c2 = c2_tb.text_frame
    tf_c2.margin_left = Inches(0)
    tf_c2.margin_top = Inches(0)
    tf_c2.word_wrap = True
    
    p_c2_h = tf_c2.paragraphs[0]
    r_c2_h = p_c2_h.add_run()
    r_c2_h.text = "02. Hyperlocal Career Reality Engine"
    r_c2_h.font.name = "Calibri"
    r_c2_h.font.size = Pt(11)
    r_c2_h.font.bold = True
    r_c2_h.font.color.rgb = NAVY

    p_c2_d = tf_c2.add_paragraph()
    r_c2_d = p_c2_d.add_run()
    r_c2_d.text = "Connects vocational choices to available regional earnings, placement outcomes, and accredited training pathways."
    r_c2_d.font.name = "Calibri"
    r_c2_d.font.size = Pt(8.5)
    r_c2_d.font.color.rgb = GRAY_TEXT

    diag2_bg = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2_left + Inches(0.12), row1_diag_y, diag_w, diag_h)
    diag2_bg.fill.solid()
    diag2_bg.fill.fore_color.rgb = WHITE
    diag2_bg.line.color.rgb = BORDER_LIGHT
    diag2_bg.line.width = Pt(1.0)

    # Node: Location Pin
    loc_node = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2_left + Inches(0.20), row1_inner_y, Inches(1.30), chip_h)
    loc_node.fill.solid()
    loc_node.fill.fore_color.rgb = BLUE_BG
    loc_node.line.color.rgb = BLUE_BG
    tf_loc = loc_node.text_frame
    tf_loc.margin_left = 0
    tf_loc.margin_top = Inches(0.06)
    p_loc = tf_loc.paragraphs[0]
    p_loc.alignment = PP_ALIGN.CENTER
    r_loc1 = p_loc.add_run()
    r_loc1.text = "📍 District Cluster\n"
    r_loc1.font.size = Pt(7.8)
    r_loc1.font.bold = True
    r_loc1.font.color.rgb = BLUE_TEXT
    r_loc2 = p_loc.add_run()
    r_loc2.text = "Local ITI & Hubs"
    r_loc2.font.size = Pt(7.0)
    r_loc2.font.color.rgb = BLUE_TEXT

    arr_c2 = slide2.shapes.add_textbox(c2_left + Inches(1.52), row1_inner_y + Inches(0.08), Inches(0.22), Inches(0.30))
    arr_c2.text_frame.margin_left = 0
    arr_c2.text_frame.margin_top = 0
    p_ac2 = arr_c2.text_frame.paragraphs[0]
    r_ac2 = p_ac2.add_run()
    r_ac2.text = "→"
    r_ac2.font.size = Pt(13)
    r_ac2.font.bold = True
    r_ac2.font.color.rgb = TEAL

    # 3 Neutral Verified Dataset Metric Badges (Factual: replaces unsupported 88%)
    met1 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2_left + Inches(1.78), row1_inner_y, Inches(1.18), chip_h)
    met1.fill.solid()
    met1.fill.fore_color.rgb = CARD_BG
    met1.line.color.rgb = BORDER_GRAY
    tf_m1 = met1.text_frame
    tf_m1.margin_left = 0
    tf_m1.margin_top = Inches(0.06)
    p_m1 = tf_m1.paragraphs[0]
    p_m1.alignment = PP_ALIGN.CENTER
    r_m1_1 = p_m1.add_run()
    r_m1_1.text = "₹ Regional Wages\n"
    r_m1_1.font.size = Pt(7.8)
    r_m1_1.font.bold = True
    r_m1_1.font.color.rgb = NAVY
    r_m1_2 = p_m1.add_run()
    r_m1_2.text = "Trade Wage Data"
    r_m1_2.font.size = Pt(6.8)
    r_m1_2.font.color.rgb = TEXT_MUTED

    met2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2_left + Inches(3.04), row1_inner_y, Inches(1.20), chip_h)
    met2.fill.solid()
    met2.fill.fore_color.rgb = CARD_BG
    met2.line.color.rgb = BORDER_GRAY
    tf_m2 = met2.text_frame
    tf_m2.margin_left = 0
    tf_m2.margin_top = Inches(0.06)
    p_m2 = tf_m2.paragraphs[0]
    p_m2.alignment = PP_ALIGN.CENTER
    r_m2_1 = p_m2.add_run()
    r_m2_1.text = "Placement Trends\n"
    r_m2_1.font.size = Pt(7.8)
    r_m2_1.font.bold = True
    r_m2_1.font.color.rgb = DARK_TEAL
    r_m2_2 = p_m2.add_run()
    r_m2_2.text = "Verified Datasets"
    r_m2_2.font.size = Pt(6.8)
    r_m2_2.font.color.rgb = TEXT_MUTED

    met3 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2_left + Inches(4.32), row1_inner_y, Inches(1.28), chip_h)
    met3.fill.solid()
    met3.fill.fore_color.rgb = CARD_BG
    met3.line.color.rgb = BORDER_GRAY
    tf_m3 = met3.text_frame
    tf_m3.margin_left = 0
    tf_m3.margin_top = Inches(0.06)
    p_m3 = tf_m3.paragraphs[0]
    p_m3.alignment = PP_ALIGN.CENTER
    r_m3_1 = p_m3.add_run()
    r_m3_1.text = "NSQF Pathway\n"
    r_m3_1.font.size = Pt(7.8)
    r_m3_1.font.bold = True
    r_m3_1.font.color.rgb = NAVY
    r_m3_2 = p_m3.add_run()
    r_m3_2.text = "ITI → Poly → B.Voc"
    r_m3_2.font.size = Pt(6.8)
    r_m3_2.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------------------
    # CARD 03: EXPLAINABLE CAREER TRADE-OFF ENGINE - Bottom Left
    # -------------------------------------------------------------------------
    c3 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c1_left, r2_top, col_w, card_h)
    c3.name = "Card_03_Tradeoff"
    c3.fill.solid()
    c3.fill.fore_color.rgb = CARD_BG
    c3.line.color.rgb = BORDER_GRAY
    c3.line.width = Pt(1.0)

    c3_tb = slide2.shapes.add_textbox(c1_left + Inches(0.12), r2_top + Inches(0.07), col_w - Inches(0.24), Inches(0.70))
    tf_c3 = c3_tb.text_frame
    tf_c3.margin_left = Inches(0)
    tf_c3.margin_top = Inches(0)
    tf_c3.word_wrap = True
    
    p_c3_h = tf_c3.paragraphs[0]
    r_c3_h = p_c3_h.add_run()
    r_c3_h.text = "03. Explainable Career Trade-off Engine"
    r_c3_h.font.name = "Calibri"
    r_c3_h.font.size = Pt(11)
    r_c3_h.font.bold = True
    r_c3_h.font.color.rgb = NAVY

    p_c3_d = tf_c3.add_paragraph()
    r_c3_d = p_c3_d.add_run()
    r_c3_d.text = "Compares career options based on family priorities, training costs, earnings, and NSQF progression."
    r_c3_d.font.name = "Calibri"
    r_c3_d.font.size = Pt(8.5)
    r_c3_d.font.color.rgb = GRAY_TEXT

    diag3_bg = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c1_left + Inches(0.12), row2_diag_y, diag_w, diag_h)
    diag3_bg.fill.solid()
    diag3_bg.fill.fore_color.rgb = WHITE
    diag3_bg.line.color.rgb = BORDER_LIGHT
    diag3_bg.line.width = Pt(1.0)

    p_w = Inches(1.26)
    p1 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c1_left + Inches(0.20), row2_inner_y, p_w, chip_h)
    p1.fill.solid()
    p1.fill.fore_color.rgb = CARD_BG
    p1.line.color.rgb = BORDER_GRAY
    tf_p1 = p1.text_frame
    tf_p1.margin_left = 0
    tf_p1.margin_top = Inches(0.06)
    p_tp1 = tf_p1.paragraphs[0]
    p_tp1.alignment = PP_ALIGN.CENTER
    r_tp1_1 = p_tp1.add_run()
    r_tp1_1.text = "Family Priorities\n"
    r_tp1_1.font.size = Pt(7.8)
    r_tp1_1.font.bold = True
    r_tp1_1.font.color.rgb = NAVY
    r_tp1_2 = p_tp1.add_run()
    r_tp1_2.text = "Budget & Values"
    r_tp1_2.font.size = Pt(6.8)
    r_tp1_2.font.color.rgb = TEXT_MUTED

    p2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c1_left + Inches(1.54), row2_inner_y, p_w, chip_h)
    p2.fill.solid()
    p2.fill.fore_color.rgb = CARD_BG
    p2.line.color.rgb = BORDER_GRAY
    tf_p2 = p2.text_frame
    tf_p2.margin_left = 0
    tf_p2.margin_top = Inches(0.06)
    p_tp2 = tf_p2.paragraphs[0]
    p_tp2.alignment = PP_ALIGN.CENTER
    r_tp2_1 = p_tp2.add_run()
    r_tp2_1.text = "Training Cost\n"
    r_tp2_1.font.size = Pt(7.8)
    r_tp2_1.font.bold = True
    r_tp2_1.font.color.rgb = DARK_TEAL
    r_tp2_2 = p_tp2.add_run()
    r_tp2_2.text = "Fees vs Subsidy"
    r_tp2_2.font.size = Pt(6.8)
    r_tp2_2.font.color.rgb = TEXT_MUTED

    p3 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c1_left + Inches(2.88), row2_inner_y, p_w, chip_h)
    p3.fill.solid()
    p3.fill.fore_color.rgb = CARD_BG
    p3.line.color.rgb = BORDER_GRAY
    tf_p3 = p3.text_frame
    tf_p3.margin_left = 0
    tf_p3.margin_top = Inches(0.06)
    p_tp3 = tf_p3.paragraphs[0]
    p_tp3.alignment = PP_ALIGN.CENTER
    r_tp3_1 = p_tp3.add_run()
    r_tp3_1.text = "Time-to-Earn\n"
    r_tp3_1.font.size = Pt(7.8)
    r_tp3_1.font.bold = True
    r_tp3_1.font.color.rgb = NAVY
    r_tp3_2 = p_tp3.add_run()
    r_tp3_2.text = "1–2 Yrs vs 4+ Yrs"
    r_tp3_2.font.size = Pt(6.8)
    r_tp3_2.font.color.rgb = TEXT_MUTED

    p4 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c1_left + Inches(4.22), row2_inner_y, Inches(1.38), chip_h)
    p4.fill.solid()
    p4.fill.fore_color.rgb = CARD_BG
    p4.line.color.rgb = BORDER_GRAY
    tf_p4 = p4.text_frame
    tf_p4.margin_left = 0
    tf_p4.margin_top = Inches(0.06)
    p_tp4 = tf_p4.paragraphs[0]
    p_tp4.alignment = PP_ALIGN.CENTER
    r_tp4_1 = p_tp4.add_run()
    r_tp4_1.text = "NSQF Progression\n"
    r_tp4_1.font.size = Pt(7.8)
    r_tp4_1.font.bold = True
    r_tp4_1.font.color.rgb = DARK_TEAL
    r_tp4_2 = p_tp4.add_run()
    r_tp4_2.text = "Degree Equivalency"
    r_tp4_2.font.size = Pt(6.8)
    r_tp4_2.font.color.rgb = TEXT_MUTED

    # -------------------------------------------------------------------------
    # CARD 04: PARENTAL RESISTANCE INTELLIGENCE - Bottom Right
    # -------------------------------------------------------------------------
    c4 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2_left, r2_top, col_w, card_h)
    c4.name = "Card_04_Resistance"
    c4.fill.solid()
    c4.fill.fore_color.rgb = CARD_BG
    c4.line.color.rgb = BORDER_GRAY
    c4.line.width = Pt(1.0)

    c4_tb = slide2.shapes.add_textbox(c2_left + Inches(0.12), r2_top + Inches(0.07), col_w - Inches(0.24), Inches(0.70))
    tf_c4 = c4_tb.text_frame
    tf_c4.margin_left = Inches(0)
    tf_c4.margin_top = Inches(0)
    tf_c4.word_wrap = True
    
    p_c4_h = tf_c4.paragraphs[0]
    r_c4_h = p_c4_h.add_run()
    r_c4_h.text = "04. Parental Resistance Intelligence"
    r_c4_h.font.name = "Calibri"
    r_c4_h.font.size = Pt(11)
    r_c4_h.font.bold = True
    r_c4_h.font.color.rgb = NAVY

    p_c4_d = tf_c4.add_paragraph()
    r_c4_d = p_c4_d.add_run()
    r_c4_d.text = "Transforms counselling concerns into aggregated regional insights and escalation pathways for scheme administrators."
    r_c4_d.font.name = "Calibri"
    r_c4_d.font.size = Pt(8.5)
    r_c4_d.font.color.rgb = GRAY_TEXT

    diag4_bg = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2_left + Inches(0.12), row2_diag_y, diag_w, diag_h)
    diag4_bg.fill.solid()
    diag4_bg.fill.fore_color.rgb = WHITE
    diag4_bg.line.color.rgb = BORDER_LIGHT
    diag4_bg.line.width = Pt(1.0)

    in_res = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2_left + Inches(0.20), row2_inner_y, Inches(1.60), chip_h)
    in_res.fill.solid()
    in_res.fill.fore_color.rgb = ROSE_BG
    in_res.line.color.rgb = ROSE_BG
    tf_ir = in_res.text_frame
    tf_ir.margin_left = 0
    tf_ir.margin_top = Inches(0.06)
    p_ir = tf_ir.paragraphs[0]
    p_ir.alignment = PP_ALIGN.CENTER
    r_ir1 = p_ir.add_run()
    r_ir1.text = "Unresolved Hesitation\n"
    r_ir1.font.size = Pt(7.8)
    r_ir1.font.bold = True
    r_ir1.font.color.rgb = ROSE_TEXT
    r_ir2 = p_ir.add_run()
    r_ir2.text = "Stigma / Financial Anxiety"
    r_ir2.font.size = Pt(6.8)
    r_ir2.font.color.rgb = ROSE_TEXT

    arr_c4 = slide2.shapes.add_textbox(c2_left + Inches(1.84), row2_inner_y + Inches(0.08), Inches(0.22), Inches(0.30))
    arr_c4.text_frame.margin_left = 0
    arr_c4.text_frame.margin_top = 0
    p_ac4 = arr_c4.text_frame.paragraphs[0]
    r_ac4 = p_ac4.add_run()
    r_ac4.text = "→"
    r_ac4.font.size = Pt(13)
    r_ac4.font.bold = True
    r_ac4.font.color.rgb = TEAL

    out_a = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2_left + Inches(2.10), row2_inner_y, Inches(1.68), chip_h)
    out_a.fill.solid()
    out_a.fill.fore_color.rgb = BLUE_BG
    out_a.line.color.rgb = BLUE_BG
    tf_oa = out_a.text_frame
    tf_oa.margin_left = 0
    tf_oa.margin_top = Inches(0.06)
    p_oa = tf_oa.paragraphs[0]
    p_oa.alignment = PP_ALIGN.CENTER
    r_oa1 = p_oa.add_run()
    r_oa1.text = "Human Counsellor Routing\n"
    r_oa1.font.size = Pt(7.8)
    r_oa1.font.bold = True
    r_oa1.font.color.rgb = BLUE_TEXT
    r_oa2 = p_oa.add_run()
    r_oa2.text = "1-on-1 Escalation Hand-off"
    r_oa2.font.size = Pt(6.8)
    r_oa2.font.color.rgb = BLUE_TEXT

    arr_plus = slide2.shapes.add_textbox(c2_left + Inches(3.82), row2_inner_y + Inches(0.08), Inches(0.18), Inches(0.30))
    arr_plus.text_frame.margin_left = 0
    arr_plus.text_frame.margin_top = 0
    p_ap = arr_plus.text_frame.paragraphs[0]
    r_ap = p_ap.add_run()
    r_ap.text = "+"
    r_ap.font.size = Pt(13)
    r_ap.font.bold = True
    r_ap.font.color.rgb = TEAL

    out_b = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2_left + Inches(4.04), row2_inner_y, Inches(1.56), chip_h)
    out_b.fill.solid()
    out_b.fill.fore_color.rgb = LIGHT_TEAL
    out_b.line.color.rgb = BORDER_TEAL
    tf_ob = out_b.text_frame
    tf_ob.margin_left = 0
    tf_ob.margin_top = Inches(0.06)
    p_ob = tf_ob.paragraphs[0]
    p_ob.alignment = PP_ALIGN.CENTER
    r_ob1 = p_ob.add_run()
    r_ob1.text = "Admin Heatmaps\n"
    r_ob1.font.size = Pt(7.8)
    r_ob1.font.bold = True
    r_ob1.font.color.rgb = DARK_TEAL
    r_ob2 = p_ob.add_run()
    r_ob2.text = "District Scheme Insights"
    r_ob2.font.size = Pt(6.8)
    r_ob2.font.color.rgb = DARK_TEAL

    prs.save(ppt_path)
    print("Successfully applied factual corrections to Slide 2 in:", ppt_path)

if __name__ == "__main__":
    apply_slide_2()
